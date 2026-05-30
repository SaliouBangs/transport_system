from django.contrib import messages
from django.core.exceptions import ValidationError
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.contrib.auth.models import User
from django.db import transaction
from django.db.models import DecimalField, Q, Sum
from django.db.models.functions import Coalesce
from django.utils.dateparse import parse_date
from django.utils import timezone
from decimal import Decimal
from commandes.models import Commande
from operations.models import Operation
from utilisateurs.models import journaliser_action
from utilisateurs.permissions import get_active_supervision_entity, get_user_role, is_admin_user, role_required
from utilisateurs.constants import ROLE_COMMERCIAL, ROLE_RESPONSABLE_COMMERCIAL

from .forms import BanqueForm, ClientDestinationFormSet, ClientForm, EncaissementClientForm, VillePerequationForm
from .models import (
    Banque,
    Client,
    EncaissementClient,
    EncaissementClientAllocation,
    VillePerequation,
    commande_est_engagement,
    commande_est_facturee,
    commande_est_risque_potentiel,
    dernier_encaissement_sur_commande,
    latest_operation_for_commande,
    montant_total_commande,
    total_encaisse_sur_commande,
)


def _can_impute_client_avance(user):
    return is_admin_user(user) or get_user_role(user) in {"comptable", "comptable_sogefi"}


def _parse_decimal_input(value):
    raw = str(value or "").strip()
    if not raw:
        raise ValueError("empty")
    normalized = "".join(raw.split())
    has_comma = "," in normalized
    has_dot = "." in normalized
    if has_comma and has_dot:
        if normalized.rfind(",") > normalized.rfind("."):
            normalized = normalized.replace(".", "").replace(",", ".")
        else:
            normalized = normalized.replace(",", "")
    elif has_comma:
        normalized = normalized.replace(",", ".")
    return Decimal(normalized)


def _build_encaissement_allocations(request, client, type_encaissement, total_montant, encaissement_instance=None):
    if type_encaissement != "multi_commandes":
        return [], []

    cibles = request.POST.getlist("allocation_cible")
    montants = request.POST.getlist("allocation_montant")
    rows = []
    allocations = []
    total_allocations = Decimal("0.00")
    existing_allocations_commandes = {}
    existing_solde_initial = Decimal("0.00")
    if encaissement_instance and getattr(encaissement_instance, "pk", None):
        existing_allocations_commandes = {
            item.commande_id: item.montant_affecte or Decimal("0.00")
            for item in encaissement_instance.allocations.all()
            if item.cible_type == "commande" and item.commande_id
        }
        existing_solde_initial = sum(
            (
                item.montant_affecte or Decimal("0.00")
                for item in encaissement_instance.allocations.all()
                if item.cible_type == "solde_initial"
            ),
            Decimal("0.00"),
        )

    seen_targets = set()
    allocated_solde_initial = Decimal("0.00")
    for index, (cible, montant_raw) in enumerate(zip(cibles, montants), start=1):
        cible = (cible or "").strip()
        montant_raw = (montant_raw or "").strip()
        if not cible and not montant_raw:
            continue
        rows.append({"cible": cible, "montant": montant_raw})
        if not cible or not montant_raw:
            raise ValidationError(f"Ligne {index}: choisissez une affectation et un montant.")
        try:
            montant_value = _parse_decimal_input(montant_raw)
        except Exception:
            raise ValidationError(f"Ligne {index}: le montant affecte est invalide.")
        if montant_value <= 0:
            raise ValidationError(f"Ligne {index}: le montant affecte doit etre strictement positif.")

        if cible == "solde_initial":
            if cible in seen_targets:
                raise ValidationError("Le solde initial ne peut apparaitre qu'une seule fois dans la repartition.")
            solde_initial_restant = (client.solde_initial_restant or Decimal("0.00")) + existing_solde_initial
            if montant_value > solde_initial_restant:
                raise ValidationError(
                    f"Ligne {index}: le montant affecte depasse le solde initial restant ({solde_initial_restant})."
                )
            allocations.append({"cible_type": "solde_initial", "commande": None, "montant_affecte": montant_value})
            allocated_solde_initial += montant_value
        elif cible == "avance_client":
            if cible in seen_targets:
                raise ValidationError("L'avance client ne peut apparaitre qu'une seule fois dans la repartition.")
            allocations.append({"cible_type": "avance_client", "commande": None, "montant_affecte": montant_value})
        elif cible.startswith("commande:"):
            commande_id = cible.split(":", 1)[1].strip()
            commande = Commande.objects.filter(id=commande_id, client=client).first()
            if not commande:
                raise ValidationError(f"Ligne {index}: la commande selectionnee est introuvable pour ce client.")
            if not commande_est_facturee(commande):
                raise ValidationError(f"Ligne {index}: seule une commande livree peut etre reglee.")
            if commande.id in seen_targets:
                raise ValidationError("Une meme commande ne peut apparaitre qu'une seule fois dans la repartition.")
            montant_commande = montant_total_commande(commande)
            solde_commande = max(Decimal("0.00"), montant_commande - total_encaisse_sur_commande(commande))
            solde_commande += existing_allocations_commandes.get(commande.id, Decimal("0.00"))
            if montant_value > solde_commande:
                raise ValidationError(
                    f"Ligne {index}: le montant affecte depasse le solde restant de la commande ({solde_commande})."
                )
            allocations.append({"cible_type": "commande", "commande": commande, "montant_affecte": montant_value})
        else:
            raise ValidationError(f"Ligne {index}: l'affectation selectionnee est invalide.")

        seen_targets.add(cible)
        total_allocations += montant_value

    if not allocations:
        raise ValidationError("Ajoutez au moins une ligne d'affectation pour un paiement reparti.")
    if total_allocations != total_montant:
        raise ValidationError(
            f"Le total reparti ({total_allocations}) doit etre exactement egal au montant de l'encaissement ({total_montant})."
        )
    return allocations, rows


def _apply_date_range(queryset, field_name, date_debut=None, date_fin=None):
    if date_debut:
        queryset = queryset.filter(**{f"{field_name}__gte": date_debut})
    if date_fin:
        queryset = queryset.filter(**{f"{field_name}__lte": date_fin})
    return queryset


def _filtered_total_encaisse_sur_commande(commande, date_debut=None, date_fin=None):
    if not commande:
        return Decimal("0.00")

    direct_queryset = _apply_date_range(commande.encaissements_clients.all(), "date_encaissement", date_debut, date_fin)
    allocation_queryset = _apply_date_range(
        commande.encaissement_allocations.all(),
        "encaissement__date_encaissement",
        date_debut,
        date_fin,
    )
    total_direct = (
        direct_queryset.aggregate(
            total=Coalesce(
                Sum("montant"),
                Decimal("0.00"),
                output_field=DecimalField(max_digits=18, decimal_places=2),
            )
        ).get("total")
        or Decimal("0.00")
    )
    total_reparti = (
        allocation_queryset.aggregate(
            total=Coalesce(
                Sum("montant_affecte"),
                Decimal("0.00"),
                output_field=DecimalField(max_digits=18, decimal_places=2),
            )
        ).get("total")
        or Decimal("0.00")
    )
    return total_direct + total_reparti


def _filtered_dernier_encaissement_sur_commande(commande, date_debut=None, date_fin=None):
    if not commande:
        return None

    queryset = EncaissementClient.objects.filter(Q(commande=commande) | Q(allocations__commande=commande)).distinct()
    queryset = _apply_date_range(queryset, "date_encaissement", date_debut, date_fin)
    return queryset.order_by("-date_encaissement", "-id").first()


def _build_client_snapshot(client, date_debut=None, date_fin=None):
    commandes_queryset = _apply_date_range(
        Commande.objects.filter(client=client),
        "date_commande",
        date_debut,
        date_fin,
    ).prefetch_related("operations").order_by("-date_creation")

    encaissements_queryset = _apply_date_range(
        EncaissementClient.objects.filter(client=client),
        "date_encaissement",
        date_debut,
        date_fin,
    ).select_related("commande").prefetch_related("allocations__commande").order_by("-date_encaissement", "-id")

    commandes = []
    total_facture = Decimal("0.00")
    total_paye_commandes = Decimal("0.00")
    engagements = Decimal("0.00")
    risque_potentiel = Decimal("0.00")
    total_paye_solde_initial = Decimal("0.00")
    total_avances_recues = Decimal("0.00")
    total_avances_affectees = Decimal("0.00")

    commandes_liste = list(commandes_queryset)
    for commande in commandes_liste:
        montant_commande = montant_total_commande(commande)
        total_paye = _filtered_total_encaisse_sur_commande(commande, date_debut, date_fin)
        solde = max(Decimal("0.00"), montant_commande - total_paye)
        dernier_reglement = _filtered_dernier_encaissement_sur_commande(commande, date_debut, date_fin)
        latest_operation = latest_operation_for_commande(commande)
        est_facturee = commande_est_facturee(commande)
        est_engagement = commande_est_engagement(commande)
        est_risque_potentiel = commande_est_risque_potentiel(commande)
        total_paye_commandes += min(montant_commande, total_paye)

        if est_facturee:
            total_facture += montant_commande
        if est_engagement:
            engagements += montant_commande
        if est_risque_potentiel:
            risque_potentiel += montant_commande

        est_soldee = solde <= Decimal("0.00") and total_paye > Decimal("0.00")
        est_operation_facturee = bool(latest_operation and (latest_operation.numero_facture or latest_operation.date_facture))
        if latest_operation and latest_operation.date_bon_retour:
            etat_affiche = "Livree / retournee / payee / soldee" if est_soldee else "Livree / retournee / facturee" if est_operation_facturee else "Livree / retournee"
        elif latest_operation and latest_operation.etat_bon == "livre":
            etat_affiche = "Livree / payee / soldee" if est_soldee else "Livree / facturee" if est_operation_facturee else "Livree"
        elif latest_operation and latest_operation.etat_bon == "charge":
            etat_affiche = "Chargee / payee / soldee" if est_soldee else "Chargee / facturee" if est_operation_facturee else "Chargee"
        elif latest_operation and latest_operation.etat_bon == "liquide":
            etat_affiche = "Liquidee / payee / soldee" if est_soldee else "Liquidee / facturee" if est_operation_facturee else "Liquidee"
        elif latest_operation and latest_operation.etat_bon == "declare":
            etat_affiche = "Declaree / payee / soldee" if est_soldee else "Declaree / facturee" if est_operation_facturee else "Declaree"
        elif est_operation_facturee:
            etat_affiche = "Facturee / payee / soldee" if est_soldee else "Facturee"
        elif est_facturee:
            etat_affiche = "Livree"
        else:
            etat_affiche = commande.get_statut_display()

        commandes.append(
            {
                "id": commande.id,
                "reference": commande.reference_affichee,
                "label": f"{commande.reference_affichee} - solde {solde}",
                "montant": str(montant_commande),
                "total_paye": str(total_paye),
                "solde": str(solde),
                "statut": commande.get_statut_display(),
                "etat_affiche": etat_affiche,
                "soldee": est_soldee,
                "livree": est_facturee,
                "paiement_type": dernier_reglement.get_type_encaissement_display() if dernier_reglement else "",
                "paiement_mode": dernier_reglement.get_mode_paiement_display() if dernier_reglement else "",
                "paiement_reference": (dernier_reglement.reference or "") if dernier_reglement else "",
            }
        )

    encaissements = [
        {
            "date": encaissement.date_encaissement.strftime("%Y-%m-%d") if encaissement.date_encaissement else "",
            "montant": str(encaissement.montant or Decimal("0.00")),
            "type": encaissement.get_type_encaissement_display(),
            "reference": encaissement.reference or "-",
            "mode": encaissement.get_mode_paiement_display(),
            "commandes": encaissement.commandes_resume,
        }
        for encaissement in encaissements_queryset[:6]
    ]

    avances_disponibles = []
    for encaissement in encaissements_queryset.filter(type_encaissement="avance_client"):
        montant_affecte = (
            encaissement.allocations.filter(cible_type="commande").aggregate(
                total=Coalesce(
                    Sum("montant_affecte"),
                    Decimal("0.00"),
                    output_field=DecimalField(max_digits=18, decimal_places=2),
                )
            ).get("total")
            or Decimal("0.00")
        )
        disponible = max(Decimal("0.00"), (encaissement.montant or Decimal("0.00")) - montant_affecte)
        total_avances_recues += encaissement.montant or Decimal("0.00")
        total_avances_affectees += montant_affecte
        if disponible <= Decimal("0.00"):
            continue
        avances_disponibles.append(
            {
                "id": encaissement.id,
                "date": encaissement.date_encaissement.strftime("%Y-%m-%d") if encaissement.date_encaissement else "",
                "montant_disponible": str(disponible),
                "reference": encaissement.reference or "",
                "mode": encaissement.get_mode_paiement_display(),
            }
        )

    multi_solde_initial = (
        EncaissementClientAllocation.objects.filter(
            encaissement__in=encaissements_queryset,
            cible_type="solde_initial",
        ).aggregate(
            total=Coalesce(
                Sum("montant_affecte"),
                Decimal("0.00"),
                output_field=DecimalField(max_digits=18, decimal_places=2),
            )
        ).get("total")
        or Decimal("0.00")
    )
    multi_avances = (
        EncaissementClientAllocation.objects.filter(
            encaissement__in=encaissements_queryset,
            cible_type="avance_client",
        ).aggregate(
            total=Coalesce(
                Sum("montant_affecte"),
                Decimal("0.00"),
                output_field=DecimalField(max_digits=18, decimal_places=2),
            )
        ).get("total")
        or Decimal("0.00")
    )

    if multi_avances > Decimal("0.00"):
        total_avances_recues += multi_avances

    total_paye_solde_initial = (
        encaissements_queryset.filter(type_encaissement="solde_initial").aggregate(
            total=Coalesce(
                Sum("montant"),
                Decimal("0.00"),
                output_field=DecimalField(max_digits=18, decimal_places=2),
            )
        ).get("total")
        or Decimal("0.00")
    ) + multi_solde_initial
    paiements_anticipes = max(Decimal("0.00"), total_avances_recues - total_avances_affectees)
    encours_client = max(Decimal("0.00"), total_facture - total_paye_commandes)
    reste_a_encaisser = encours_client
    exposition_client_totale = max(
        Decimal("0.00"),
        (client.solde_initial or Decimal("0.00"))
        - total_paye_solde_initial
        + encours_client
        + engagements
        + risque_potentiel
        - paiements_anticipes,
    )
    total_paye_global = (
        encaissements_queryset.aggregate(
            total=Coalesce(
                Sum("montant"),
                Decimal("0.00"),
                output_field=DecimalField(max_digits=18, decimal_places=2),
            )
        ).get("total")
        or Decimal("0.00")
    )

    suggestions_avance = [
        {
            "commande_id": item["id"],
            "commande_reference": item["reference"],
            "solde": item["solde"],
        }
        for item in commandes
        if item["livree"] and not item["soldee"] and Decimal(item["solde"]) > Decimal("0.00")
    ]

    return {
        "commandes": commandes[:30],
        "encaissements": encaissements,
        "avances_disponibles": avances_disponibles,
        "suggestions_avance": suggestions_avance,
        "client_resume": {
            "solde_initial_restant": str(max(Decimal("0.00"), (client.solde_initial or Decimal("0.00")) - total_paye_solde_initial)),
            "encours_client": str(encours_client),
            "paiements_anticipes": str(paiements_anticipes),
            "engagements": str(engagements),
            "risque_potentiel": str(risque_potentiel),
            "reste_a_encaisser_reel": str(reste_a_encaisser),
            "total_facture": str(total_facture),
            "total_paye_commandes": str(total_paye_commandes),
            "total_paye_solde_initial": str(total_paye_solde_initial),
            "total_avances_recues": str(total_avances_recues),
            "total_avances_affectees": str(total_avances_affectees),
            "total_paye_global": str(total_paye_global),
            "exposition_client_totale": str(exposition_client_totale),
        },
    }


def liste_clients(request):
    query = request.GET.get("q", "").strip()
    risque_scope = request.GET.get("risque", "").strip() or "all"

    clients = Client.objects.select_related("prospect", "commercial").prefetch_related("destinations__ville_perequation", "encaissements")
    if get_user_role(request.user) == ROLE_COMMERCIAL and not is_admin_user(request.user):
        clients = clients.filter(commercial=request.user)
    if query:
        clients = clients.filter(
            Q(entreprise__icontains=query)
            | Q(nom__icontains=query)
            | Q(telephone__icontains=query)
            | Q(ville__icontains=query)
            | Q(adresse__icontains=query)
        )

    clients = list(clients)

    if risque_scope == "a_risque":
        clients = [client for client in clients if client.niveau_risque in {"alerte", "critique"}]
    elif risque_scope == "critique":
        clients = [client for client in clients if client.niveau_risque == "critique"]
    elif risque_scope == "alerte":
        clients = [client for client in clients if client.niveau_risque == "alerte"]
    elif risque_scope == "couvert":
        clients = [client for client in clients if client.niveau_risque == "ok"]
    elif risque_scope == "plus_risque":
        clients.sort(key=lambda client: client.risque_client or Decimal("0.00"), reverse=True)

    total_risque = sum((client.risque_client or Decimal("0.00") for client in clients), Decimal("0.00"))
    total_creance = sum((client.creance_client or Decimal("0.00") for client in clients), Decimal("0.00"))
    clients_critique = sum(1 for client in clients if client.niveau_risque == "critique")

    return render(
        request,
        "clients/clients.html",
        {
            "clients": clients,
            "query": query,
            "risque_scope": risque_scope,
            "clients_count": len(clients),
            "clients_critique": clients_critique,
            "total_risque": total_risque,
            "total_creance": total_creance,
            "clients_non_affectes": Client.objects.filter(commercial__isnull=True).count(),
        },
    )


def portefeuille_clients(request):
    query = request.GET.get("q", "").strip()
    commercial_id = request.GET.get("commercial", "").strip()
    scope = request.GET.get("scope", "").strip() or "all"

    commerciaux = User.objects.filter(
        groups__name__in=[ROLE_COMMERCIAL, ROLE_RESPONSABLE_COMMERCIAL]
    ).exclude(is_superuser=True).distinct().order_by("first_name", "last_name", "username")

    clients = Client.objects.select_related("prospect", "commercial").order_by("entreprise", "nom")
    if query:
        clients = clients.filter(
            Q(entreprise__icontains=query)
            | Q(nom__icontains=query)
            | Q(ville__icontains=query)
        )
    if commercial_id:
        clients = clients.filter(commercial_id=commercial_id)
    if scope == "unassigned":
        clients = clients.filter(commercial__isnull=True)

    if request.method == "POST":
        client_id = request.POST.get("client_id")
        target_commercial_id = request.POST.get("commercial_id")
        client = get_object_or_404(Client, id=client_id)
        commercial = get_object_or_404(commerciaux, id=target_commercial_id)
        client.commercial = commercial
        client.save(update_fields=["commercial"])
        journaliser_action(
            request.user,
            "Clients",
            "Affectation portefeuille",
            client.entreprise,
            (
                f"{request.user.username} a affecte le client {client.entreprise} "
                f"au portefeuille de {commercial.username}."
            ),
        )
        messages.success(
            request,
            f"Le client {client.entreprise} a ete affecte a {commercial.get_full_name() or commercial.username}.",
        )
        suffix = request.GET.urlencode()
        return redirect(f"/clients/portefeuilles/{'?' + suffix if suffix else ''}")

    return render(
        request,
        "clients/portefeuilles.html",
        {
            "clients": clients,
            "commerciaux": commerciaux,
            "query": query,
            "commercial_id": commercial_id,
            "scope": scope,
        },
    )


def ajouter_client(request):
    if request.method == "POST":
        form = ClientForm(request.POST, user=request.user)
        destination_formset = ClientDestinationFormSet(request.POST, prefix="destinations")
        if form.is_valid() and destination_formset.is_valid():
            client = form.save(commit=False)
            if get_user_role(request.user) == "commercial" and not is_admin_user(request.user):
                client.commercial = request.user
            client.save()
            destination_formset.instance = client
            destination_formset.save()
            journaliser_action(
                request.user,
                "Clients",
                "Ajout de client",
                client.entreprise,
                f"{request.user.username} a ajoute le client {client.entreprise}.",
            )
            return redirect("clients")
    else:
        form = ClientForm(user=request.user)
        destination_formset = ClientDestinationFormSet(prefix="destinations")

    return render(
        request,
        "clients/ajouter_client.html",
        {"form": form, "destination_formset": destination_formset},
    )


def modifier_client(request, id):
    client_queryset = Client.objects.all()
    if get_user_role(request.user) == "commercial" and not is_admin_user(request.user):
        client_queryset = client_queryset.filter(commercial=request.user)
    client = get_object_or_404(client_queryset, id=id)
    if request.method == "POST":
        form = ClientForm(request.POST, instance=client, user=request.user)
        destination_formset = ClientDestinationFormSet(request.POST, instance=client, prefix="destinations")
        if form.is_valid() and destination_formset.is_valid():
            client = form.save(commit=False)
            if get_user_role(request.user) == "commercial" and not is_admin_user(request.user):
                client.commercial = request.user
            client.save()
            destination_formset.save()
            journaliser_action(
                request.user,
                "Clients",
                "Modification de client",
                client.entreprise,
                f"{request.user.username} a modifie le client {client.entreprise}.",
            )
            return redirect("clients")
    else:
        form = ClientForm(instance=client, user=request.user)
        destination_formset = ClientDestinationFormSet(instance=client, prefix="destinations")

    return render(
        request,
        "clients/modifier_client.html",
        {"form": form, "client": client, "destination_formset": destination_formset},
    )


def villes_perequation(request):
    if not is_admin_user(request.user):
        messages.error(request, "Seul l'administrateur peut gerer les villes de perequation.")
        return redirect("/comptes/parametres/")

    editing_id = (request.GET.get("edit") or "").strip()
    editing_instance = None
    if editing_id.isdigit():
        editing_instance = VillePerequation.objects.filter(id=editing_id).first()

    if request.method == "POST":
        form_id = (request.POST.get("perequation_id") or "").strip()
        instance = VillePerequation.objects.filter(id=form_id).first() if form_id.isdigit() else None
        form = VillePerequationForm(request.POST, instance=instance)
        if form.is_valid():
            ville = form.save()
            action = "Modification" if instance else "Ajout"
            journaliser_action(
                request.user,
                "Parametres",
                f"{action} ville de perequation",
                ville.nom,
                f"{request.user.username} a {'modifie' if instance else 'ajoute'} la ville de perequation {ville.nom}.",
            )
            messages.success(request, "La ville de perequation a ete enregistree.")
            return redirect("villes_perequation")
        editing_instance = instance
    else:
        form = VillePerequationForm(instance=editing_instance)

    villes = VillePerequation.objects.order_by("nom")
    return render(
        request,
        "clients/villes_perequation.html",
        {
            "form": form,
            "villes_perequation": villes,
            "editing_instance": editing_instance,
        },
    )


def supprimer_client(request, id):
    client_queryset = Client.objects.all()
    if get_user_role(request.user) == "commercial" and not is_admin_user(request.user):
        client_queryset = client_queryset.filter(commercial=request.user)
    client = get_object_or_404(client_queryset, id=id)
    client_label = client.entreprise
    client.delete()
    journaliser_action(
        request.user,
        "Clients",
        "Suppression de client",
        client_label,
        f"{request.user.username} a supprime le client {client_label}.",
    )
    return redirect("clients")


def prospect_infos(request):
    prospect_id = request.GET.get("prospect_id")
    if not prospect_id:
        return JsonResponse(
            {"success": False, "errors": {"prospect": ["Prospect manquant."]}},
            status=400,
        )

    from prospects.models import Prospect

    prospect = Prospect.objects.filter(id=prospect_id).first()
    if not prospect:
        return JsonResponse(
            {"success": False, "errors": {"prospect": ["Prospect introuvable."]}},
            status=404,
        )

    return JsonResponse(
        {
            "success": True,
            "prospect": {
                "nom": prospect.nom,
                "fonction": prospect.fonction,
                "telephone": prospect.telephone,
                "entreprise": prospect.entreprise,
                "ville": prospect.ville,
                "commercial_id": prospect.commercial_id,
            },
        }
    )


def ajouter_client_modal(request):
    if request.method != "POST":
        return JsonResponse(
            {"success": False, "errors": {"__all__": ["Requete invalide."]}},
            status=405,
        )

    form = ClientForm(request.POST, user=request.user)
    if form.is_valid():
        client = form.save(commit=False)
        if get_user_role(request.user) == "commercial" and not is_admin_user(request.user):
            client.commercial = request.user
        client.save()
        journaliser_action(
            request.user,
            "Clients",
            "Ajout de client",
            client.entreprise,
            f"{request.user.username} a ajoute le client {client.entreprise} depuis une fenetre modale.",
        )
        return JsonResponse(
            {
                "success": True,
                "client": {
                    "id": client.id,
                    "label": client.entreprise,
                },
            }
        )

    errors = {
        field: [item["message"] for item in messages]
        for field, messages in form.errors.get_json_data().items()
    }
    return JsonResponse({"success": False, "errors": errors}, status=400)


def _build_encaissements_history_context(request):
    query = request.GET.get("q", "").strip()
    client_query = request.GET.get("client", "").strip()
    banque = request.GET.get("banque", "").strip()
    mode_paiement = request.GET.get("mode_paiement", "").strip()
    reference = request.GET.get("reference", "").strip()
    date_debut = request.GET.get("date_debut", "").strip()
    date_fin = request.GET.get("date_fin", "").strip()
    commande_query = request.GET.get("commande", "").strip()

    encaissements = EncaissementClient.objects.select_related("client", "commande").prefetch_related("allocations__commande").order_by("-date_encaissement", "-id")
    if get_user_role(request.user) == ROLE_COMMERCIAL and not is_admin_user(request.user):
        encaissements = encaissements.filter(client__commercial=request.user)
    if query:
        encaissements = encaissements.filter(
            Q(client__entreprise__icontains=query)
            | Q(client__nom__icontains=query)
            | Q(banque__icontains=query)
            | Q(reference__icontains=query)
            | Q(commande__reference__icontains=query)
            | Q(commande__operations__numero_bl__icontains=query)
            | Q(allocations__commande__reference__icontains=query)
            | Q(allocations__commande__operations__numero_bl__icontains=query)
        )
    if mode_paiement:
        encaissements = encaissements.filter(mode_paiement=mode_paiement)
    if date_debut:
        encaissements = encaissements.filter(date_encaissement__gte=date_debut)
    if date_fin:
        encaissements = encaissements.filter(date_encaissement__lte=date_fin)
    encaissements = encaissements.distinct()

    return {
        "encaissements": encaissements,
        "query": query,
        "client_query": client_query,
        "commande_query": commande_query,
        "banque": banque,
        "mode_paiement": mode_paiement,
        "reference": reference,
        "date_debut": date_debut,
        "date_fin": date_fin,
    }


def _commercial_display_name(user):
    if not user:
        return "Commercial non affecte"
    full_name = user.get_full_name().strip()
    return full_name or user.username


def _commande_product_family(commande):
    produit = getattr(commande, "produit", None)
    label = (getattr(produit, "nom", "") or "").strip().upper()
    if "ESS" in label:
        return "essence"
    if "GAS" in label:
        return "gasoil"
    return ""


def _build_rapport_encaissements_commerciaux_context(request):
    date_debut = (request.GET.get("date_debut") or "").strip()
    date_fin = (request.GET.get("date_fin") or "").strip()
    commercial_id = (request.GET.get("commercial") or "").strip()

    encaissements = (
        EncaissementClient.objects.select_related(
            "client__commercial",
            "commande__client__commercial",
            "commande__produit",
        )
        .prefetch_related(
            "allocations__commande__client__commercial",
            "allocations__commande__produit",
        )
        .order_by("-date_encaissement", "-id")
    )

    parsed_date_debut = parse_date(date_debut) if date_debut else None
    parsed_date_fin = parse_date(date_fin) if date_fin else None
    if parsed_date_debut:
        encaissements = encaissements.filter(date_encaissement__gte=parsed_date_debut)
    if parsed_date_fin:
        encaissements = encaissements.filter(date_encaissement__lte=parsed_date_fin)
    if commercial_id.isdigit():
        encaissements = encaissements.filter(
            Q(client__commercial_id=commercial_id)
            | Q(commande__client__commercial_id=commercial_id)
            | Q(allocations__commande__client__commercial_id=commercial_id)
        )

    summary_map = {}
    detail_rows = []
    totals = {
        "essence": Decimal("0.00"),
        "gasoil": Decimal("0.00"),
        "global": Decimal("0.00"),
    }

    def ensure_row(commercial):
        key = commercial.pk if commercial else 0
        if key not in summary_map:
            summary_map[key] = {
                "commercial": commercial,
                "commercial_name": _commercial_display_name(commercial),
                "essence": Decimal("0.00"),
                "gasoil": Decimal("0.00"),
                "total": Decimal("0.00"),
                "encaissement_ids": set(),
            }
        return summary_map[key]

    def add_amount(encaissement, commande, amount):
        family = _commande_product_family(commande)
        if family not in {"essence", "gasoil"}:
            return
        commercial = getattr(getattr(commande, "client", None), "commercial", None)
        if not commercial:
            commercial = getattr(getattr(encaissement, "client", None), "commercial", None)
        row = ensure_row(commercial)
        amount = amount or Decimal("0.00")
        essence_amount = amount if family == "essence" else Decimal("0.00")
        gasoil_amount = amount if family == "gasoil" else Decimal("0.00")
        row[family] += amount
        row["total"] += amount
        row["encaissement_ids"].add(encaissement.pk)
        totals[family] += amount
        totals["global"] += amount
        detail_rows.append(
            {
                "date": encaissement.date_encaissement,
                "commercial": commercial,
                "commercial_name": _commercial_display_name(commercial),
                "client": encaissement.client,
                "client_name": getattr(encaissement.client, "entreprise", ""),
                "commande": commande,
                "commande_reference": getattr(commande, "reference_affichee", ""),
                "produit": getattr(getattr(commande, "produit", None), "nom", ""),
                "mode_paiement": encaissement.get_mode_paiement_display(),
                "banque": encaissement.banque or "-",
                "reference": encaissement.reference or "-",
                "deposant": encaissement.nom_deposant or "-",
                "essence": essence_amount,
                "gasoil": gasoil_amount,
                "total": amount,
            }
        )

    for encaissement in encaissements.distinct():
        allocations = [
            allocation
            for allocation in encaissement.allocations.all()
            if allocation.cible_type == "commande" and allocation.commande_id
        ]
        if allocations:
            for allocation in allocations:
                add_amount(encaissement, allocation.commande, allocation.montant_affecte)
        elif encaissement.commande_id:
            add_amount(encaissement, encaissement.commande, encaissement.montant)

    summary_rows = sorted(summary_map.values(), key=lambda item: item["commercial_name"].lower())
    for row in summary_rows:
        row["encaissements_count"] = len(row["encaissement_ids"])
    detail_rows = sorted(
        detail_rows,
        key=lambda item: (item["commercial_name"].lower(), item["date"] or timezone.localdate(), item["commande_reference"]),
    )

    commercial_ids = (
        Client.objects.filter(commercial__isnull=False)
        .values_list("commercial_id", flat=True)
        .distinct()
    )
    commerciaux = User.objects.filter(id__in=commercial_ids).order_by("first_name", "last_name", "username")

    return {
        "summary_rows": summary_rows,
        "payment_rows": detail_rows,
        "totals": totals,
        "commerciaux": commerciaux,
        "selected_commercial": commercial_id,
        "date_debut": date_debut,
        "date_fin": date_fin,
        "active_entity": get_active_supervision_entity(request),
        "current_filters": request.GET.urlencode(),
    }


def rapport_encaissements_commerciaux(request):
    return render(
        request,
        "clients/rapport_encaissements_commerciaux.html",
        _build_rapport_encaissements_commerciaux_context(request),
    )


def export_rapport_encaissements_commerciaux_xls(request):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return HttpResponse("Le module openpyxl n'est pas installe sur cet environnement Python.", status=500)

    context = _build_rapport_encaissements_commerciaux_context(request)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Encaissements commerciaux"
    headers = ["Date", "Commercial", "Client", "Commande", "Produit", "Mode", "Banque", "Reference", "Deposant", "ESS", "GAS", "Total"]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="123047")
    for row in context["payment_rows"]:
        worksheet.append(
            [
                row["date"],
                row["commercial_name"],
                row["client_name"],
                row["commande_reference"],
                row["produit"],
                row["mode_paiement"],
                row["banque"],
                row["reference"],
                row["deposant"],
                float(row["essence"]),
                float(row["gasoil"]),
                float(row["total"]),
            ]
        )
    worksheet.append([])
    worksheet.append(["Totaux", "", "", "", "", "", "", "", "", float(context["totals"]["essence"]), float(context["totals"]["gasoil"]), float(context["totals"]["global"])])
    for column in worksheet.columns:
        column_letter = column[0].column_letter
        worksheet.column_dimensions[column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 32)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="rapport_encaissements_commerciaux.xlsx"'
    workbook.save(response)
    return response


def export_rapport_encaissements_commerciaux_pdf(request):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        return HttpResponse("Le module reportlab n'est pas installe sur cet environnement Python.", status=500)

    context = _build_rapport_encaissements_commerciaux_context(request)
    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="rapport_encaissements_commerciaux.pdf"'
    document = SimpleDocTemplate(response, pagesize=landscape(A4), rightMargin=20, leftMargin=20, topMargin=22, bottomMargin=18)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph("Rapport encaissements par commercial", styles["Title"]),
        Paragraph(f"ESS: {context['totals']['essence']} GNF | GAS: {context['totals']['gasoil']} GNF | Total: {context['totals']['global']} GNF", styles["Normal"]),
        Spacer(1, 12),
    ]
    data = [["Date", "Commercial", "Client", "Commande", "Produit", "Mode", "Banque", "Ref.", "Deposant", "ESS", "GAS", "Total"]]
    for row in context["payment_rows"]:
        data.append(
            [
                row["date"].strftime("%d/%m/%Y") if row["date"] else "",
                row["commercial_name"],
                row["client_name"],
                row["commande_reference"],
                row["produit"],
                row["mode_paiement"],
                row["banque"],
                row["reference"],
                row["deposant"],
                f"{row['essence']:.0f}",
                f"{row['gasoil']:.0f}",
                f"{row['total']:.0f}",
            ]
        )
    if len(data) == 1:
        data.append(["Aucun encaissement", "", "", "", "", "", "", "", "", "", "", ""])
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#123047")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d7e8f2")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f9fc")]),
            ]
        )
    )
    elements.append(table)
    document.build(elements)
    return response


def _operation_facture_amount(operation):
    if operation.montant_facture is not None:
        return Decimal(operation.montant_facture or 0)
    commande = operation.commande
    if commande and commande.prix_negocie is not None:
        return Decimal(operation.quantite or 0) * Decimal(commande.prix_negocie or 0)
    return Decimal("0.00")


def _build_rapport_factures_clients_context(request):
    commercial_id = (request.GET.get("commercial") or "").strip()
    client_id = (request.GET.get("client") or "").strip()
    client_search = (request.GET.get("client_search") or "").strip()
    statut = (request.GET.get("statut") or "").strip()
    if statut not in {"", "non_soldees", "soldees"}:
        statut = ""
    date_debut = (request.GET.get("date_debut") or "").strip()
    date_fin = (request.GET.get("date_fin") or "").strip()

    operations = (
        Operation.objects.filter(remplace_par__isnull=True)
        .filter(Q(numero_facture__gt="") | Q(date_facture__isnull=False) | Q(montant_facture__isnull=False))
        .select_related("commande__client__commercial", "client__commercial", "produit")
        .order_by("commande_id", "date_facture", "date_creation", "id")
    )
    if commercial_id.isdigit():
        operations = operations.filter(Q(commande__client__commercial_id=commercial_id) | Q(client__commercial_id=commercial_id))
    if client_id.isdigit():
        operations = operations.filter(Q(commande__client_id=client_id) | Q(client_id=client_id))
    elif client_search:
        operations = operations.filter(
            Q(commande__client__entreprise__icontains=client_search)
            | Q(commande__client__nom__icontains=client_search)
            | Q(client__entreprise__icontains=client_search)
            | Q(client__nom__icontains=client_search)
        )
    if date_debut:
        operations = operations.filter(date_facture__gte=date_debut)
    if date_fin:
        operations = operations.filter(date_facture__lte=date_fin)

    rows = []
    paid_remaining_by_commande = {}
    totals = {"facture": Decimal("0.00"), "encaisse": Decimal("0.00"), "reste": Decimal("0.00")}
    summary_map = {}

    for operation in operations:
        commande = operation.commande
        client = commande.client if getattr(commande, "client_id", None) else operation.client
        commercial = getattr(client, "commercial", None)
        montant_facture = _operation_facture_amount(operation)
        if commande and commande.id not in paid_remaining_by_commande:
            paid_remaining_by_commande[commande.id] = total_encaisse_sur_commande(commande)
        encaissé = Decimal("0.00")
        if commande:
            encaissé = min(paid_remaining_by_commande.get(commande.id, Decimal("0.00")), montant_facture)
            paid_remaining_by_commande[commande.id] = max(Decimal("0.00"), paid_remaining_by_commande.get(commande.id, Decimal("0.00")) - encaissé)
        reste = max(Decimal("0.00"), montant_facture - encaissé)
        row_status = "soldee" if montant_facture > 0 and reste <= 0 else "non_soldee"
        if statut == "soldees" and row_status != "soldee":
            continue
        if statut == "non_soldees" and row_status != "non_soldee":
            continue

        commercial_name = _commercial_display_name(commercial)
        client_name = getattr(client, "entreprise", "") or getattr(client, "nom", "") or "-"
        row = {
            "operation": operation,
            "commercial": commercial,
            "commercial_name": commercial_name,
            "client": client,
            "client_name": client_name,
            "commande_reference": commande.reference_affichee if commande else "-",
            "numero_facture": operation.numero_facture or f"Facture BL {operation.numero_bl}",
            "date_facture": operation.date_facture,
            "produit": operation.produit.nom if operation.produit_id else (commande.produit.nom if commande and commande.produit_id else "-"),
            "montant_facture": montant_facture,
            "montant_encaisse": encaissé,
            "reste": reste,
            "statut": row_status,
        }
        rows.append(row)
        totals["facture"] += montant_facture
        totals["encaisse"] += encaissé
        totals["reste"] += reste

        key = commercial.pk if commercial else 0
        if key not in summary_map:
            summary_map[key] = {
                "commercial_name": commercial_name,
                "facture": Decimal("0.00"),
                "encaisse": Decimal("0.00"),
                "reste": Decimal("0.00"),
                "count": 0,
            }
        summary_map[key]["facture"] += montant_facture
        summary_map[key]["encaisse"] += encaissé
        summary_map[key]["reste"] += reste
        summary_map[key]["count"] += 1

    rows = sorted(rows, key=lambda item: (item["commercial_name"].lower(), item["client_name"].lower(), item["date_facture"] or timezone.localdate()))
    summary_rows = sorted(summary_map.values(), key=lambda item: item["commercial_name"].lower())
    commercial_ids = Client.objects.filter(commercial__isnull=False).values_list("commercial_id", flat=True).distinct()
    return {
        "rows": rows,
        "summary_rows": summary_rows,
        "totals": totals,
        "commerciaux": User.objects.filter(id__in=commercial_ids).order_by("first_name", "last_name", "username"),
        "clients": Client.objects.order_by("entreprise", "nom"),
        "selected_commercial": commercial_id,
        "selected_client": client_id,
        "client_search": client_search,
        "statut": statut,
        "date_debut": date_debut,
        "date_fin": date_fin,
        "current_filters": request.GET.urlencode(),
    }


def rapport_factures_clients(request):
    return render(request, "clients/rapport_factures_clients.html", _build_rapport_factures_clients_context(request))


def export_rapport_factures_clients_xls(request):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return HttpResponse("Le module openpyxl n'est pas installe sur cet environnement Python.", status=500)
    context = _build_rapport_factures_clients_context(request)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Factures clients"
    headers = ["Commercial", "Client", "Facture", "Date facture", "Commande", "BL", "Produit", "Montant facture", "Encaisse", "Reste", "Statut"]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="123047")
    for row in context["rows"]:
        operation = row["operation"]
        worksheet.append([
            row["commercial_name"],
            row["client_name"],
            row["numero_facture"],
            row["date_facture"],
            row["commande_reference"],
            operation.numero_bl,
            row["produit"],
            float(row["montant_facture"]),
            float(row["montant_encaisse"]),
            float(row["reste"]),
            "Soldee" if row["statut"] == "soldee" else "Non soldee",
        ])
    worksheet.append([])
    worksheet.append(["Totaux", "", "", "", "", "", "", float(context["totals"]["facture"]), float(context["totals"]["encaisse"]), float(context["totals"]["reste"]), ""])
    for column in worksheet.columns:
        column_letter = column[0].column_letter
        worksheet.column_dimensions[column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 34)
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="rapport_factures_clients.xlsx"'
    workbook.save(response)
    return response


def export_rapport_factures_clients_pdf(request):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        return HttpResponse("Le module reportlab n'est pas installe sur cet environnement Python.", status=500)
    context = _build_rapport_factures_clients_context(request)
    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="rapport_factures_clients.pdf"'
    document = SimpleDocTemplate(response, pagesize=landscape(A4), rightMargin=20, leftMargin=20, topMargin=22, bottomMargin=18)
    styles = getSampleStyleSheet()
    data = [["Commercial", "Client", "Facture", "Date", "BL", "Montant", "Encaisse", "Reste", "Statut"]]
    for row in context["rows"]:
        data.append([
            row["commercial_name"],
            row["client_name"],
            row["numero_facture"],
            row["date_facture"].strftime("%d/%m/%Y") if row["date_facture"] else "",
            row["operation"].numero_bl,
            f"{row['montant_facture']:.0f}",
            f"{row['montant_encaisse']:.0f}",
            f"{row['reste']:.0f}",
            "Soldee" if row["statut"] == "soldee" else "Non soldee",
        ])
    if len(data) == 1:
        data.append(["Aucune facture", "", "", "", "", "", "", "", ""])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#123047")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d7e8f2")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f9fc")]),
    ]))
    elements = [
        Paragraph("Situation des factures clients", styles["Title"]),
        Paragraph(f"Facture: {context['totals']['facture']} GNF | Encaisse: {context['totals']['encaisse']} GNF | Reste: {context['totals']['reste']} GNF", styles["Normal"]),
        Spacer(1, 12),
        table,
    ]
    document.build(elements)
    return response


def encaissements_clients(request):
    history_context = _build_encaissements_history_context(request)
    clients = Client.objects.order_by("entreprise", "nom")
    if get_user_role(request.user) == ROLE_COMMERCIAL and not is_admin_user(request.user):
        clients = clients.filter(commercial=request.user)
    banques = Banque.objects.filter(actif=True).order_by("nom")
    allocation_rows = [{}]

    if request.method == "POST":
        form = EncaissementClientForm(request.POST, user=request.user)
        if form.is_valid():
            try:
                with transaction.atomic():
                    encaissement = form.save(commit=False)
                    allocations, allocation_rows = _build_encaissement_allocations(
                        request,
                        encaissement.client,
                        encaissement.type_encaissement,
                        encaissement.montant or Decimal("0.00"),
                        encaissement_instance=encaissement,
                    )
                    encaissement.save()
                    if allocations:
                        EncaissementClientAllocation.objects.bulk_create(
                            [
                                EncaissementClientAllocation(
                                    encaissement=encaissement,
                                    cible_type=item["cible_type"],
                                    commande=item["commande"],
                                    montant_affecte=item["montant_affecte"],
                                )
                                for item in allocations
                            ]
                        )
                journaliser_action(
                    request.user,
                    "Clients",
                    "Encaissement client",
                    encaissement.client.entreprise,
                    f"{request.user.username} a enregistre un encaissement de {encaissement.montant} pour {encaissement.client.entreprise}.",
                )
                messages.success(request, "Encaissement enregistre.")
                return redirect("encaissements_clients")
            except ValidationError as exc:
                form.add_error(None, exc.message if hasattr(exc, "message") else str(exc))
    else:
        form = EncaissementClientForm(user=request.user)

    return render(
        request,
        "clients/encaissements.html",
        {
            "form": form,
            "clients": clients,
            "banques": banques,
            "mode_paiement_choices": EncaissementClient.MODE_PAIEMENT_CHOICES,
            "can_manage_encaissements": is_admin_user(request.user),
            "can_impute_avances": _can_impute_client_avance(request.user),
            "allocation_rows": allocation_rows,
            "banque_form": BanqueForm(),
            **history_context,
        },
    )


def historique_encaissements_clients(request):
    history_context = _build_encaissements_history_context(request)
    return render(
        request,
        "clients/encaissements_historique.html",
        {
            "mode_paiement_choices": EncaissementClient.MODE_PAIEMENT_CHOICES,
            "can_manage_encaissements": is_admin_user(request.user),
            "can_impute_avances": _can_impute_client_avance(request.user),
            **history_context,
        },
    )


def modifier_encaissement_client(request, id):
    if not is_admin_user(request.user):
        messages.error(request, "Seul l'administrateur peut modifier un encaissement.")
        return redirect("encaissements_clients")

    encaissement = get_object_or_404(EncaissementClient.objects.prefetch_related("allocations__commande"), id=id)
    allocation_rows = [
        {
            "cible": f"commande:{item.commande_id}" if item.cible_type == "commande" and item.commande_id else item.cible_type,
            "montant": str(item.montant_affecte),
        }
        for item in encaissement.allocations.all()
    ] or [{}]
    if request.method == "POST":
        form = EncaissementClientForm(request.POST, instance=encaissement)
        if form.is_valid():
            try:
                with transaction.atomic():
                    encaissement = form.save(commit=False)
                    allocations, allocation_rows = _build_encaissement_allocations(
                        request,
                        encaissement.client,
                        encaissement.type_encaissement,
                        encaissement.montant or Decimal("0.00"),
                        encaissement_instance=encaissement,
                    )
                    encaissement.save()
                    encaissement.allocations.all().delete()
                    if allocations:
                        EncaissementClientAllocation.objects.bulk_create(
                            [
                                EncaissementClientAllocation(
                                    encaissement=encaissement,
                                    cible_type=item["cible_type"],
                                    commande=item["commande"],
                                    montant_affecte=item["montant_affecte"],
                                )
                                for item in allocations
                            ]
                        )
                journaliser_action(
                    request.user,
                    "Clients",
                    "Modification d'encaissement",
                    encaissement.client.entreprise,
                    f"{request.user.username} a modifie un encaissement de {encaissement.montant} pour {encaissement.client.entreprise}.",
                )
                messages.success(request, "Encaissement mis a jour.")
                return redirect("encaissements_clients")
            except ValidationError as exc:
                form.add_error(None, exc.message if hasattr(exc, "message") else str(exc))
    else:
        form = EncaissementClientForm(instance=encaissement)

    clients = Client.objects.order_by("entreprise", "nom")
    banques = Banque.objects.filter(actif=True).order_by("nom")
    encaissements = EncaissementClient.objects.select_related("client", "commande").prefetch_related("allocations__commande").order_by("-date_encaissement", "-id")
    return render(
        request,
        "clients/encaissements.html",
        {
            "form": form,
            "encaissements": encaissements,
            "clients": clients,
            "banques": banques,
            "query": "",
            "client_query": "",
            "commande_query": "",
            "banque": "",
            "mode_paiement": "",
            "reference": "",
            "date_debut": "",
            "date_fin": "",
            "mode_paiement_choices": EncaissementClient.MODE_PAIEMENT_CHOICES,
            "can_manage_encaissements": True,
            "can_impute_avances": _can_impute_client_avance(request.user),
            "editing_encaissement": encaissement,
            "allocation_rows": allocation_rows,
            "banque_form": BanqueForm(),
        },
    )


def supprimer_encaissement_client(request, id):
    if not is_admin_user(request.user):
        messages.error(request, "Seul l'administrateur peut supprimer un encaissement.")
        return redirect("encaissements_clients")

    encaissement = get_object_or_404(EncaissementClient, id=id)
    if request.method == "POST":
        client_label = encaissement.client.entreprise
        montant = encaissement.montant
        encaissement.delete()
        journaliser_action(
            request.user,
            "Clients",
            "Suppression d'encaissement",
            client_label,
            f"{request.user.username} a supprime un encaissement de {montant} pour {client_label}.",
        )
        messages.success(request, "Encaissement supprime.")
    return redirect("encaissements_clients")


def commandes_client_infos(request):
    client_id = request.GET.get("client_id")
    date_debut = parse_date((request.GET.get("date_debut") or "").strip() or "")
    date_fin = parse_date((request.GET.get("date_fin") or "").strip() or "")
    if not client_id:
        return JsonResponse({"success": False, "errors": {"client": ["Client manquant."]}}, status=400)

    client_queryset = Client.objects.all()
    if get_user_role(request.user) == ROLE_COMMERCIAL and not is_admin_user(request.user):
        client_queryset = client_queryset.filter(commercial=request.user)
    client = client_queryset.filter(id=client_id).first()
    if not client:
        return JsonResponse({"success": False, "errors": {"client": ["Client introuvable."]}}, status=404)
    snapshot = _build_client_snapshot(client, date_debut=date_debut, date_fin=date_fin)
    return JsonResponse({"success": True, **snapshot})


def ajouter_banque_modal(request):
    if request.method != "POST" or request.headers.get("X-Requested-With") != "XMLHttpRequest":
        return JsonResponse({"success": False, "errors": {"__all__": ["Requete invalide."]}}, status=400)

    form = BanqueForm(request.POST)
    if form.is_valid():
        banque = form.save()
        return JsonResponse(
            {
                "success": True,
                "banque": {
                    "id": banque.id,
                    "label": banque.nom,
                },
            }
        )

    errors = {
        field: [item["message"] for item in messages]
        for field, messages in form.errors.get_json_data().items()
    }
    return JsonResponse({"success": False, "errors": errors}, status=400)


def imputer_avance_client(request, id):
    if not _can_impute_client_avance(request.user):
        messages.error(request, "Seuls l'administrateur et la comptabilite peuvent imputer une avance client.")
        return redirect("encaissements_clients")

    encaissement = get_object_or_404(
        EncaissementClient.objects.select_related("client").prefetch_related("allocations__commande"),
        id=id,
        type_encaissement="avance_client",
    )
    if request.method != "POST":
        return redirect("encaissements_clients")

    commande_id = (request.POST.get("commande_id") or "").strip()
    montant_raw = (request.POST.get("montant_affecte") or "").strip()

    if not commande_id:
        messages.error(request, "Choisissez une commande pour imputer l'avance.")
        return redirect("encaissements_clients")
    if not montant_raw:
        messages.error(request, "Saisissez le montant a imputer sur la commande.")
        return redirect("encaissements_clients")

    try:
        montant_a_imputer = _parse_decimal_input(montant_raw)
    except Exception:
        messages.error(request, "Le montant a imputer est invalide.")
        return redirect("encaissements_clients")

    if montant_a_imputer <= Decimal("0.00"):
        messages.error(request, "Le montant a imputer doit etre strictement positif.")
        return redirect("encaissements_clients")

    commande = Commande.objects.filter(id=commande_id, client=encaissement.client).first()
    if not commande:
        messages.error(request, "La commande choisie est introuvable pour ce client.")
        return redirect("encaissements_clients")

    avance_disponible = encaissement.montant_non_affecte
    if montant_a_imputer > avance_disponible:
        messages.error(
            request,
            f"Le montant saisi depasse l'avance disponible ({avance_disponible}).",
        )
        return redirect("encaissements_clients")

    montant_commande = (commande.quantite or Decimal("0.00")) * (commande.prix_negocie or Decimal("0.00"))
    solde_commande = montant_commande - total_encaisse_sur_commande(commande)
    if solde_commande <= Decimal("0.00"):
        messages.error(request, "Cette commande est deja entierement couverte.")
        return redirect("encaissements_clients")
    if montant_a_imputer > solde_commande:
        messages.error(
            request,
            f"Le montant saisi depasse le solde restant de la commande ({solde_commande}).",
        )
        return redirect("encaissements_clients")

    with transaction.atomic():
        allocation, created = EncaissementClientAllocation.objects.get_or_create(
            encaissement=encaissement,
            commande=commande,
            defaults={"montant_affecte": montant_a_imputer},
        )
        if not created:
            allocation.montant_affecte = (allocation.montant_affecte or Decimal("0.00")) + montant_a_imputer
            allocation.full_clean()
            allocation.save(update_fields=["montant_affecte"])

    journaliser_action(
        request.user,
        "Clients",
        "Imputation d'avance client",
        encaissement.client.entreprise,
        (
            f"{request.user.username} a impute {montant_a_imputer} de l'avance client "
            f"sur la commande {commande.reference_affichee} pour {encaissement.client.entreprise}."
        ),
    )
    messages.success(
        request,
        f"Avance imputee sur {commande.reference_affichee} pour {montant_a_imputer}.",
    )
    return redirect("encaissements_clients")


liste_clients = role_required(
    "commercial",
    "responsable_commercial",
    "directeur",
    "comptable",
    "comptable_sogefi",
    "caissiere",
    "dga",
    "dga_sogefi",
    "logistique",
    "transitaire",
    "maintenancier",
    "responsable_achat",
    "controleur",
)(liste_clients)
ajouter_client = role_required("commercial", "responsable_commercial")(ajouter_client)
modifier_client = role_required("commercial", "responsable_commercial")(modifier_client)
supprimer_client = role_required("directeur")(supprimer_client)
prospect_infos = role_required("commercial", "responsable_commercial")(prospect_infos)
ajouter_client_modal = role_required("commercial", "responsable_commercial")(ajouter_client_modal)
portefeuille_clients = role_required("responsable_commercial")(portefeuille_clients)
encaissements_clients = role_required("commercial", "responsable_commercial", "directeur", "comptable", "comptable_sogefi")(encaissements_clients)
historique_encaissements_clients = role_required("commercial", "responsable_commercial", "directeur", "comptable", "comptable_sogefi")(historique_encaissements_clients)
modifier_encaissement_client = role_required("commercial", "responsable_commercial", "directeur", "comptable", "comptable_sogefi")(modifier_encaissement_client)
supprimer_encaissement_client = role_required("commercial", "responsable_commercial", "directeur", "comptable", "comptable_sogefi")(supprimer_encaissement_client)
commandes_client_infos = role_required("commercial", "responsable_commercial", "directeur", "comptable", "comptable_sogefi")(commandes_client_infos)
ajouter_banque_modal = role_required("commercial", "responsable_commercial", "directeur", "comptable", "comptable_sogefi")(ajouter_banque_modal)
imputer_avance_client = role_required("commercial", "responsable_commercial", "directeur", "comptable", "comptable_sogefi")(imputer_avance_client)


def detail_client(request, id):
    client_queryset = Client.objects.select_related("prospect", "commercial").prefetch_related("destinations__ville_perequation", "encaissements")
    if get_user_role(request.user) == ROLE_COMMERCIAL and not is_admin_user(request.user):
        client_queryset = client_queryset.filter(commercial=request.user)
    client = get_object_or_404(client_queryset, id=id)

    from commandes.models import Commande
    from operations.models import Operation

    statut_commande = request.GET.get("statut_commande", "").strip()
    etat_bl = request.GET.get("etat_bl", "").strip()
    date_debut = request.GET.get("date_debut", "").strip()
    date_fin = request.GET.get("date_fin", "").strip()

    commandes = Commande.objects.filter(client=client).select_related("produit", "camion", "chauffeur").prefetch_related("operations").order_by("-date_creation")
    operations = Operation.objects.filter(client=client, remplace_par__isnull=True).select_related("commande", "produit", "camion", "chauffeur").order_by("-date_creation")
    encaissements = client.encaissements.select_related("commande").prefetch_related("allocations__commande").all()

    if date_debut:
        commandes = commandes.filter(date_commande__gte=date_debut)
        operations = operations.filter(date_creation__date__gte=date_debut)
        encaissements = encaissements.filter(date_encaissement__gte=date_debut)
    if date_fin:
        commandes = commandes.filter(date_commande__lte=date_fin)
        operations = operations.filter(date_creation__date__lte=date_fin)
        encaissements = encaissements.filter(date_encaissement__lte=date_fin)

    if statut_commande:
        commandes = commandes.filter(statut=statut_commande)
    if etat_bl:
        operations = operations.filter(etat_bon=etat_bl, remplace_par__isnull=True)
        commandes_filtrees = []
        for commande in commandes:
            latest_operation = commande.operations.filter(remplace_par__isnull=True).order_by("-date_creation").first()
            if latest_operation and latest_operation.etat_bon == etat_bl:
                commandes_filtrees.append(commande.id)
        commandes = commandes.filter(id__in=commandes_filtrees)

    valorisation_commandes = Decimal("0.00")
    valorisation_livree = Decimal("0.00")
    total_essence = Decimal("0.00")
    total_gasoil = Decimal("0.00")
    for commande in commandes:
        commande.montant_commande_affiche = montant_total_commande(commande)
        valorisation_commandes += commande.montant_commande_affiche
        produit_nom = (commande.produit.nom or "").strip().upper() if commande.produit_id else ""
        if produit_nom == "ESSENCE":
            total_essence += commande.montant_commande_affiche
        elif produit_nom == "GASOIL":
            total_gasoil += commande.montant_commande_affiche

        montant_livre = Decimal("0.00")
        if commande_est_facturee(commande):
            montant_livre = commande.montant_commande_affiche
        commande.montant_livre_affiche = montant_livre
        valorisation_livree += montant_livre
        commande.total_encaisse_affiche = total_encaisse_sur_commande(commande)
        commande.solde_commande_affiche = max(Decimal("0.00"), commande.montant_commande_affiche - commande.total_encaisse_affiche)
        commande.dernier_reglement_affiche = dernier_encaissement_sur_commande(commande)

    # Calcul manuel pour garder une valorisation fiable sur quantite x prix_negocie.
    synthese_statuts = []
    for statut, _label in Commande.STATUT_CHOICES:
        total_statut = Decimal("0.00")
        for commande in Commande.objects.filter(client=client, statut=statut):
            total_statut += (commande.quantite or Decimal("0.00")) * (commande.prix_negocie or Decimal("0.00"))
        if total_statut > 0 or statut == statut_commande:
            synthese_statuts.append({"statut": statut, "total": total_statut})

    total_encaissements = encaissements.aggregate(
        total=Coalesce(Sum("montant"), Decimal("0.00"))
    ).get("total") or Decimal("0.00")
    total_encaissements_solde_initial = encaissements.filter(type_encaissement="solde_initial").aggregate(
        total=Coalesce(Sum("montant"), Decimal("0.00"))
    ).get("total") or Decimal("0.00")
    disponible_decouvert = (client.decouvert_maximum_autorise or Decimal("0.00")) - (client.risque_client or Decimal("0.00"))

    return render(
        request,
        "clients/detail_client.html",
        {
            "client": client,
            "commandes": commandes,
            "operations": operations,
            "encaissements": encaissements,
            "statut_commande": statut_commande,
            "etat_bl": etat_bl,
            "date_debut": date_debut,
            "date_fin": date_fin,
            "statut_choices": Commande.STATUT_CHOICES,
            "etat_bl_choices": Operation.ETAT_BON_CHOICES,
            "valorisation_commandes": valorisation_commandes,
            "valorisation_livree": valorisation_livree,
            "total_encaissements": total_encaissements,
            "total_encaissements_solde_initial": total_encaissements_solde_initial,
            "solde_initial_restant": client.solde_initial_restant,
            "disponible_decouvert": disponible_decouvert,
            "paiements_anticipes": client.paiements_anticipes,
            "engagement_net": client.engagement_net,
            "risque_client": client.risque_client,
            "synthese_statuts": synthese_statuts,
            "total_essence": total_essence,
            "total_gasoil": total_gasoil,
        },
    )


def _build_line_chart(labels, series_specs):
    chart_width = 640
    chart_height = 220
    left_pad = 54
    right_pad = 18
    top_pad = 16
    bottom_pad = 26
    plot_width = chart_width - left_pad - right_pad
    plot_height = chart_height - top_pad - bottom_pad

    max_value = Decimal("0.00")
    for _name, _color, values in series_specs:
        for value in values:
            max_value = max(max_value, Decimal(value or Decimal("0.00")))
    if max_value <= Decimal("0.00"):
        max_value = Decimal("1.00")

    count = max(len(labels), 1)
    x_step = Decimal("0.00") if count <= 1 else Decimal(plot_width) / Decimal(count - 1)

    x_points = []
    for index, label in enumerate(labels):
        x = left_pad + float(x_step * index)
        x_points.append({"label": label, "x": round(x, 2)})

    series = []
    for name, color, values in series_specs:
        dots = []
        for index, raw_value in enumerate(values):
            value = Decimal(raw_value or Decimal("0.00"))
            ratio = float(value / max_value) if max_value > 0 else 0
            x = left_pad + float(x_step * index) if count > 1 else left_pad + (plot_width / 2)
            y = top_pad + (plot_height - (ratio * plot_height))
            dots.append(
                {
                    "x": round(x, 2),
                    "y": round(y, 2),
                    "value": value,
                    "label": labels[index],
                }
            )
        series.append(
            {
                "name": name,
                "color": color,
                "points": " ".join(f"{dot['x']},{dot['y']}" for dot in dots),
                "dots": dots,
            }
        )

    y_ticks = []
    for tick_ratio in [1, Decimal("0.5"), Decimal("0.0")]:
        tick_value = max_value * Decimal(str(tick_ratio))
        y = top_pad + (plot_height - (float(tick_ratio) * plot_height))
        y_ticks.append({"label": tick_value, "y": round(y, 2)})

    return {
        "width": chart_width,
        "height": chart_height,
        "left_pad": left_pad,
        "right_pad": right_pad,
        "top_pad": top_pad,
        "bottom_pad": bottom_pad,
        "x_points": x_points,
        "y_ticks": y_ticks,
        "series": series,
        "plot_bottom": top_pad + plot_height,
        "plot_right": left_pad + plot_width,
    }


def rapport_financier_client(request, id):
    client_queryset = Client.objects.select_related("prospect", "commercial").prefetch_related("destinations__ville_perequation", "encaissements")
    if get_user_role(request.user) == ROLE_COMMERCIAL and not is_admin_user(request.user):
        client_queryset = client_queryset.filter(commercial=request.user)
    client = get_object_or_404(client_queryset, id=id)

    from commandes.models import Commande
    from operations.models import Operation

    date_debut = parse_date((request.GET.get("date_debut") or "").strip() or "")
    date_fin = parse_date((request.GET.get("date_fin") or "").strip() or "")

    commandes = _apply_date_range(
        Commande.objects.filter(client=client).select_related("produit").prefetch_related("operations").order_by("-date_creation"),
        "date_commande",
        date_debut,
        date_fin,
    )
    operations = _apply_date_range(
        Operation.objects.filter(client=client, remplace_par__isnull=True).select_related("commande", "produit").order_by("-date_creation"),
        "date_creation__date",
        date_debut,
        date_fin,
    )
    encaissements = _apply_date_range(
        client.encaissements.select_related("commande").prefetch_related("allocations__commande").all(),
        "date_encaissement",
        date_debut,
        date_fin,
    ).order_by("-date_encaissement", "-id")

    total_commandes = Decimal("0.00")
    total_livre = Decimal("0.00")
    total_encaisse_commandes = Decimal("0.00")
    total_essence = Decimal("0.00")
    total_gasoil = Decimal("0.00")
    total_solde = Decimal("0.00")
    top_impayes = []
    statut_totaux = []
    statut_map = {}

    for commande in commandes:
        montant = montant_total_commande(commande)
        total_commandes += montant
        total_paye = total_encaisse_sur_commande(commande)
        total_encaisse_commandes += min(montant, total_paye)
        solde = max(Decimal("0.00"), montant - total_paye)
        total_solde += solde
        produit_nom = (commande.produit.nom or "").strip().upper() if commande.produit_id else ""
        if produit_nom == "ESSENCE":
            total_essence += montant
        elif produit_nom == "GASOIL":
            total_gasoil += montant

        latest_operation = latest_operation_for_commande(commande)
        est_livree = commande_est_facturee(commande)
        if est_livree:
            total_livre += montant

        statut_label = latest_operation.get_etat_bon_display() if latest_operation else commande.get_statut_display()
        statut_map.setdefault(statut_label, Decimal("0.00"))
        statut_map[statut_label] += montant

        top_impayes.append(
            {
                "reference": commande.reference_affichee,
                "client": client.entreprise,
                "produit": commande.produit.nom if commande.produit_id else "-",
                "montant": montant,
                "paye": total_paye,
                "solde": solde,
                "statut": statut_label,
            }
        )

    paiements_anticipes = client.paiements_anticipes or Decimal("0.00")
    risque_net = client.risque_client or Decimal("0.00")
    dma = client.decouvert_maximum_autorise or Decimal("0.00")
    disponible = dma - risque_net
    ratio_dma = Decimal("0.00")
    if dma > 0:
        ratio_dma = (risque_net / dma) * Decimal("100.00")
    reste_a_encaisser = max(Decimal("0.00"), total_solde - paiements_anticipes)

    total_encaissements = (
        encaissements.aggregate(
            total=Coalesce(
                Sum("montant"),
                Decimal("0.00"),
                output_field=DecimalField(max_digits=18, decimal_places=2),
            )
        ).get("total")
        or Decimal("0.00")
    )

    top_impayes = sorted(top_impayes, key=lambda item: item["solde"], reverse=True)[:6]
    statut_totaux = sorted(statut_map.items(), key=lambda item: item[1], reverse=True)
    statut_max = max((amount for _, amount in statut_totaux), default=Decimal("0.00"))
    statut_graph = []
    for label, amount in statut_totaux:
        percent = 0
        if statut_max > 0:
            percent = int((amount / statut_max) * 100)
        statut_graph.append({"label": label, "amount": amount, "percent": max(percent, 4 if amount > 0 else 0)})

    produit_total = total_essence + total_gasoil
    essence_percent = int((total_essence / produit_total) * 100) if produit_total > 0 else 0
    gasoil_percent = int((total_gasoil / produit_total) * 100) if produit_total > 0 else 0

    monthly_map = {}
    monthly_livres_map = {}
    monthly_essence_map = {}
    monthly_gasoil_map = {}
    for encaissement in encaissements:
        if not encaissement.date_encaissement:
            continue
        key = encaissement.date_encaissement.strftime("%Y-%m")
        monthly_map.setdefault(key, Decimal("0.00"))
        monthly_map[key] += encaissement.montant or Decimal("0.00")

    for commande in commandes:
        montant = (commande.quantite or Decimal("0.00")) * (commande.prix_negocie or Decimal("0.00"))
        produit_nom = (commande.produit.nom or "").strip().upper() if commande.produit_id else ""
        commande_key = commande.date_commande.strftime("%Y-%m") if commande.date_commande else ""
        if commande_key:
            if produit_nom == "ESSENCE":
                monthly_essence_map.setdefault(commande_key, Decimal("0.00"))
                monthly_essence_map[commande_key] += montant
            elif produit_nom == "GASOIL":
                monthly_gasoil_map.setdefault(commande_key, Decimal("0.00"))
                monthly_gasoil_map[commande_key] += montant
        latest_operation = commande.operations.filter(remplace_par__isnull=True).order_by("-date_creation").first()
        if latest_operation and latest_operation.etat_bon == "livre" and latest_operation.date_bons_livres:
            key = latest_operation.date_bons_livres.strftime("%Y-%m")
            monthly_livres_map.setdefault(key, Decimal("0.00"))
            monthly_livres_map[key] += montant

    today = timezone.localdate()
    monthly_points = []
    year = today.year
    month = today.month
    month_keys = []
    for _ in range(8):
        key = f"{year:04d}-{month:02d}"
        month_keys.append(key)
        monthly_points.append((key, monthly_map.get(key, Decimal("0.00"))))
        month -= 1
        if month == 0:
            month = 12
            year -= 1
    month_keys.reverse()
    monthly_points.reverse()
    monthly_max = max((amount for _, amount in monthly_points), default=Decimal("0.00"))
    monthly_graph = []
    for key, amount in monthly_points:
        percent = 0
        if monthly_max > 0:
            percent = int((amount / monthly_max) * 100)
        monthly_graph.append(
            {
                "label": key,
                "amount": amount,
                "percent": max(percent, 8 if amount > 0 else 0),
            }
        )

    month_labels = [key[5:7] + "/" + key[0:4] for key in month_keys]
    encaissements_series = [monthly_map.get(key, Decimal("0.00")) for key in month_keys]
    livres_series = [monthly_livres_map.get(key, Decimal("0.00")) for key in month_keys]
    essence_series = [monthly_essence_map.get(key, Decimal("0.00")) for key in month_keys]
    gasoil_series = [monthly_gasoil_map.get(key, Decimal("0.00")) for key in month_keys]
    cashflow_chart = _build_line_chart(
        month_labels,
        [
            ("Encaissements", "#132f88", encaissements_series),
            ("Commandes livrees", "#7fc1ff", livres_series),
        ],
    )
    product_line_chart = _build_line_chart(
        month_labels,
        [
            ("Essence", "#0f766e", essence_series),
            ("Gasoil", "#d4a017", gasoil_series),
        ],
    )

    return render(
        request,
        "clients/rapport_financier_client.html",
        {
            "client": client,
            "date_debut": request.GET.get("date_debut", "").strip(),
            "date_fin": request.GET.get("date_fin", "").strip(),
            "total_commandes": total_commandes,
            "total_livre": total_livre,
            "total_encaisse": total_encaisse_commandes,
            "total_encaissements": total_encaissements,
            "total_essence": total_essence,
            "total_gasoil": total_gasoil,
            "paiements_anticipes": paiements_anticipes,
            "risque_net": risque_net,
            "reste_a_encaisser": reste_a_encaisser,
            "disponible": disponible,
            "dma": dma,
            "ratio_dma": ratio_dma,
            "top_impayes": top_impayes,
            "statut_graph": statut_graph,
            "monthly_graph": monthly_graph,
            "cashflow_chart": cashflow_chart,
            "product_line_chart": product_line_chart,
            "encaissements": encaissements[:8],
            "operations_count": operations.count(),
            "essence_percent": essence_percent,
            "gasoil_percent": gasoil_percent,
        },
    )


detail_client = role_required(
    "commercial",
    "responsable_commercial",
    "directeur",
    "comptable",
    "comptable_sogefi",
    "caissiere",
    "dga",
    "dga_sogefi",
    "logistique",
    "transitaire",
    "maintenancier",
    "responsable_achat",
    "controleur",
)(detail_client)
rapport_financier_client = role_required(
    "directeur",
    "dga",
)(rapport_financier_client)
rapport_encaissements_commerciaux = role_required(
    "dga",
    "directeur",
)(rapport_encaissements_commerciaux)
export_rapport_encaissements_commerciaux_xls = role_required(
    "dga",
    "directeur",
)(export_rapport_encaissements_commerciaux_xls)
export_rapport_encaissements_commerciaux_pdf = role_required(
    "dga",
    "directeur",
)(export_rapport_encaissements_commerciaux_pdf)
rapport_factures_clients = role_required(
    "dga",
    "directeur",
    "responsable_commercial",
)(rapport_factures_clients)
export_rapport_factures_clients_xls = role_required(
    "dga",
    "directeur",
    "responsable_commercial",
)(export_rapport_factures_clients_xls)
export_rapport_factures_clients_pdf = role_required(
    "dga",
    "directeur",
    "responsable_commercial",
)(export_rapport_factures_clients_pdf)
