from io import BytesIO
from decimal import Decimal

from django.db import transaction
from django.core.exceptions import ValidationError
from django.db.models import Case, Count, DecimalField, F, Q, Sum, Value, When
from django.db.models.functions import Coalesce
from django.contrib import messages
from django.contrib.auth.models import User
from django.http import FileResponse, Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST
from django.utils import timezone
from django.utils.formats import date_format
from django.utils.safestring import mark_safe
from django.utils.text import get_valid_filename
from urllib.parse import urlencode
import json

from clients.forms import BanqueForm
from clients.models import Banque
from chauffeurs.models import Chauffeur
from camions.models import Camion
from depenses.forms import TypePieceIdentiteForm
from depenses.models import Depense, TypePieceIdentite
from utilisateurs.models import journaliser_action
from utilisateurs.permissions import get_user_role, is_admin_user, role_required

from .forms import (
    ArticleStockForm,
    ArticleStockConversionFormSet,
    ApprovisionnementCaisseForm,
    FournisseurForm,
    MaintenanceAchatForm,
    MaintenanceFactureFormSet,
    MaintenanceGarageForm,
    MaintenanceGarageLigneFormSet,
    MaintenanceDecisionPaiementForm,
    MaintenancePaiementForm,
    MouvementStockForm,
    PanneCatalogueForm,
    PanneFournisseurPrixForm,
    PrestataireForm,
    SoldeInitialCaisseForm,
    TypeMaintenanceForm,
)
from .models import (
    ArticleStock,
    ApprovisionnementCaisse,
    Fournisseur,
    Maintenance,
    MaintenanceFacture,
    PanneCatalogue,
    PanneFournisseurPrix,
    MaintenanceSousLigne,
    MouvementStock,
    Prestataire,
    SoldeInitialCaisse,
    TypeMaintenance,
)


CAISSIERE_ROLES = {"caissiere", "caissiere_soni", "caissiere_avena"}


def _caisse_internal_entity_for_role(role):
    if role == "comptable":
        return Depense.ENTITE_SONI
    if role == "caissiere_soni":
        return Depense.ENTITE_SONI
    if role == "caissiere_avena":
        return Depense.ENTITE_AVENA
    if role == "comptable_avena":
        return Depense.ENTITE_AVENA
    return ""


def _maintenance_queryset():
    _normalize_stock_only_workflow()
    return Maintenance.objects.select_related("camion", "fournisseur").prefetch_related(
        "factures_achat__fournisseur",
        "lignes__type_maintenance",
        "lignes__sous_lignes__article_stock",
    )


def _safe_file_response(field_file):
    if not field_file:
        raise Http404("Fichier introuvable.")
    storage = field_file.storage
    name = field_file.name
    if not name or not storage.exists(name):
        raise Http404("Fichier introuvable.")
    filename = get_valid_filename(name.split("/")[-1]) or "facture"
    return FileResponse(storage.open(name, "rb"), as_attachment=False, filename=filename)


def _fournisseur_portefeuille_for_role(user, fallback=Fournisseur.PORTEFEUILLE_LOGISTIQUE):
    role = get_user_role(user)
    if role in {"responsable_achat", "dga_sogefi", "dga_avena", "comptable_avena"}:
        return Fournisseur.PORTEFEUILLE_INTERNE
    return fallback


def _fournisseur_entite_for_user(user, portefeuille):
    role = get_user_role(user)
    if role in {"logistique", "dga", "comptable", "caissiere_soni"}:
        return Fournisseur.ENTITE_SONI
    if role in {"dga_avena", "comptable_avena", "caissiere_avena"}:
        return Fournisseur.ENTITE_AVENA
    if portefeuille != Fournisseur.PORTEFEUILLE_INTERNE:
        return ""
    return Fournisseur.ENTITE_SOGEFI


def _fournisseurs_queryset_for_user(user):
    portefeuille = _fournisseur_portefeuille_for_role(user)
    queryset = Fournisseur.objects.filter(portefeuille=portefeuille)
    entite_reference = _fournisseur_entite_for_user(user, portefeuille)
    if entite_reference:
        if portefeuille == Fournisseur.PORTEFEUILLE_LOGISTIQUE:
            queryset = queryset.filter(Q(entite_reference=entite_reference) | Q(entite_reference=""))
        else:
            queryset = queryset.filter(entite_reference=entite_reference)
    return queryset.order_by("nom_fournisseur", "entreprise")


def _default_fournisseurs_return_url(portefeuille):
    if portefeuille == Fournisseur.PORTEFEUILLE_INTERNE:
        return "/depenses/"
    return "/maintenance/achat/"


def _fournisseurs_return_context(request, portefeuille):
    next_url = (request.GET.get("next") or request.POST.get("next") or "").strip()
    return_url = next_url or _default_fournisseurs_return_url(portefeuille)
    return {
        "next_url": next_url,
        "return_url": return_url,
    }


def _normalize_stock_only_workflow():
    candidates = (
        Maintenance.objects.exclude(statut__in=["payee", "rejetee_dga", "rejetee_dg", "validee_stock"])
        .prefetch_related("lignes__sous_lignes")
    )
    updates = []
    for maintenance in candidates:
        if not maintenance.is_stock_only():
            continue
        if maintenance.validation_dg_at:
            target_status = "validee_stock"
        elif maintenance.validation_dga_at:
            target_status = "attente_dg"
        else:
            target_status = "attente_dga"
        if maintenance.statut != target_status:
            updates.append((maintenance.pk, target_status))

    for maintenance_id, target_status in updates:
        Maintenance.objects.filter(pk=maintenance_id).update(statut=target_status)


def _maintenance_payment_allowed(user, maintenance):
    role = get_user_role(user)
    if is_admin_user(user):
        return True
    if maintenance.mode_paiement == Maintenance.MODE_CHEQUE:
        return role == "comptable_sogefi"
    if maintenance.mode_paiement == Maintenance.MODE_ESPECE:
        return role == "caissiere"
    return role in {"directeur", "comptable"}


def _apply_maintenance_filters(request, queryset):
    q = (request.GET.get("q") or "").strip()
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()
    statut = (request.GET.get("statut") or "").strip()

    if q:
        queryset = queryset.filter(
            Q(reference__icontains=q)
            | Q(numero_facture__icontains=q)
            | Q(camion__code_camion__icontains=q)
            | Q(camion__numero_tracteur__icontains=q)
            | Q(camion__numero_citerne__icontains=q)
            | Q(camion__chauffeur__nom__icontains=q)
        ).distinct()

    if date_from:
        queryset = queryset.filter(date_debut__date__gte=date_from)

    if date_to:
        queryset = queryset.filter(date_debut__date__lte=date_to)

    if statut:
        queryset = queryset.filter(statut=statut)

    return queryset, {
        "q": q,
        "date_from": date_from,
        "date_to": date_to,
        "statut": statut,
    }


def _maintenance_export_rows(queryset):
    rows = []
    for maintenance in queryset:
        rows.append(
            [
                maintenance.reference,
                maintenance.camion.code_camion,
                maintenance.camion.numero_tracteur,
                maintenance.camion.numero_citerne or "",
                (
                    Chauffeur.objects.filter(camion=maintenance.camion)
                    .values_list("nom", flat=True)
                    .first()
                    or ""
                ),
                maintenance.get_statut_display(),
                maintenance.date_debut.strftime("%Y-%m-%d %H:%M"),
                maintenance.date_fin.strftime("%Y-%m-%d %H:%M") if maintenance.date_fin else "",
                str(maintenance.total_facture),
                maintenance.numero_facture or "",
                maintenance.date_paiement.strftime("%Y-%m-%d") if maintenance.date_paiement else "",
                maintenance.mode_paiement or "",
            ]
        )
    return rows


def _export_maintenance_xls(queryset, filename):
    try:
        from openpyxl import Workbook
    except ImportError:
        return HttpResponse(
            "Le module openpyxl n'est pas installe sur cet environnement Python.",
            status=503,
            content_type="text/plain; charset=utf-8",
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Maintenance"
    sheet.append(
        [
            "Reference",
            "Code camion",
            "Tracteur",
            "Citerne",
            "Chauffeur",
            "Statut",
            "Date entree",
            "Date sortie",
            "Montant",
            "Numero facture",
            "Date paiement",
            "Mode paiement",
        ]
    )
    for row in _maintenance_export_rows(queryset):
        sheet.append(row)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    workbook.save(response)
    return response


def _export_maintenance_pdf(queryset, filename):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    except ImportError:
        return HttpResponse(
            "Le module reportlab n'est pas installe sur cet environnement Python.",
            status=503,
            content_type="text/plain; charset=utf-8",
        )

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
    data = [[
        "Reference",
        "Code camion",
        "Tracteur",
        "Citerne",
        "Chauffeur",
        "Statut",
        "Entree",
        "Sortie",
        "Montant",
        "Facture",
        "Paiement",
        "Mode paiement",
    ]]
    data.extend(_maintenance_export_rows(queryset))

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#123047")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d9e2e8")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f8fb")]),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("PADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    doc.build([table])

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
    response.write(buffer.getvalue())
    return response


_UNITS = {
    0: "zero",
    1: "un",
    2: "deux",
    3: "trois",
    4: "quatre",
    5: "cinq",
    6: "six",
    7: "sept",
    8: "huit",
    9: "neuf",
    10: "dix",
    11: "onze",
    12: "douze",
    13: "treize",
    14: "quatorze",
    15: "quinze",
    16: "seize",
}
_TENS = {
    20: "vingt",
    30: "trente",
    40: "quarante",
    50: "cinquante",
    60: "soixante",
}


def _number_to_french(n):
    n = int(n)
    if n < 0:
        return "moins " + _number_to_french(-n)
    if n in _UNITS:
        return _UNITS[n]
    if n < 20:
        return "dix-" + _UNITS[n - 10]
    if n < 70:
        tens = (n // 10) * 10
        unit = n % 10
        base = _TENS[tens]
        if unit == 0:
            return base
        if unit == 1:
            return f"{base} et un"
        return f"{base}-{_number_to_french(unit)}"
    if n < 80:
        if n == 71:
            return "soixante et onze"
        return f"soixante-{_number_to_french(n - 60)}"
    if n < 100:
        if n == 80:
            return "quatre-vingts"
        return f"quatre-vingt-{_number_to_french(n - 80)}"
    if n < 1000:
        hundreds = n // 100
        rest = n % 100
        if hundreds == 1:
            prefix = "cent"
        else:
            prefix = f"{_number_to_french(hundreds)} cent"
        if rest == 0:
            return prefix
        return f"{prefix} {_number_to_french(rest)}"
    if n < 1_000_000:
        thousands = n // 1000
        rest = n % 1000
        if thousands == 1:
            prefix = "mille"
        else:
            prefix = f"{_number_to_french(thousands)} mille"
        if rest == 0:
            return prefix
        return f"{prefix} {_number_to_french(rest)}"
    millions = n // 1_000_000
    rest = n % 1_000_000
    prefix = "un million" if millions == 1 else f"{_number_to_french(millions)} millions"
    if rest == 0:
        return prefix
    return f"{prefix} {_number_to_french(rest)}"


def _amount_to_words(amount):
    amount = Decimal(amount or 0)
    entier = int(amount)
    decimals = int((amount - Decimal(entier)) * 100)
    words = _number_to_french(entier) + " francs guineens"
    if decimals:
        words += f" et {_number_to_french(decimals)} centimes"
    return words


def _format_amount(amount):
    amount = Decimal(amount or 0)
    quantized = amount.quantize(Decimal("0.01"))
    text = f"{quantized:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if text.endswith(",00"):
        return text[:-3]
    return text


def _format_step_date(value):
    if not value:
        return "-"
    if hasattr(value, "hour"):
        return date_format(value, "d/m/Y H:i")
    return date_format(value, "d/m/Y")


def _get_duplicate_facture_matches(numero_facture, maintenance_id=None):
    numero_facture = (numero_facture or "").strip()
    if not numero_facture:
        return []
    invoice_maintenance_ids = MaintenanceFacture.objects.filter(
        numero_facture__iexact=numero_facture
    ).values_list("maintenance_id", flat=True)
    queryset = (
        Maintenance.objects.exclude(numero_facture="")
        .filter(Q(numero_facture__iexact=numero_facture) | Q(id__in=invoice_maintenance_ids))
        .select_related("camion")
        .order_by("-date_creation")
    )
    if maintenance_id:
        queryset = queryset.exclude(pk=maintenance_id)
    return list(queryset[:5])


def _attach_subline_values(formset, request=None):
    for ligne_form in formset.forms:
        if request and request.method == "POST":
            ids = request.POST.getlist(f"subline-{ligne_form.prefix}-ids")
            labels = request.POST.getlist(f"subline-{ligne_form.prefix}-labels")
            quantities = request.POST.getlist(f"subline-{ligne_form.prefix}-quantites")
            article_ids = request.POST.getlist(f"subline-{ligne_form.prefix}-articles")
            values = [
                {
                    "id": ids[index] if index < len(ids) else "",
                    "libelle": labels[index] if index < len(labels) else "",
                    "quantite": quantities[index] if index < len(quantities) else "1",
                    "article_stock_id": article_ids[index] if index < len(article_ids) else "",
                    "panne_catalogue_id": "",
                    "prix_unitaire": "0",
                    "montant": "0",
                }
                for index in range(max(len(labels), len(quantities), len(ids), len(article_ids)))
            ]
        else:
            values = (
                list(
                    ligne_form.instance.sous_lignes.values(
                        "id",
                        "article_stock_id",
                        "libelle",
                        "panne_catalogue_id",
                        "quantite",
                        "prix_unitaire",
                        "montant",
                    )
                )
                if ligne_form.instance.pk
                else []
            )
        ligne_form.subline_values = values or [
            {
                "id": "",
                "article_stock_id": "",
                "panne_catalogue_id": "",
                "libelle": "",
                "quantite": "1",
                "prix_unitaire": "0",
                "montant": "0",
            }
        ]

    formset.empty_form.subline_values = [
        {
            "id": "",
            "article_stock_id": "",
            "panne_catalogue_id": "",
            "libelle": "",
            "quantite": "1",
            "prix_unitaire": "0",
            "montant": "0",
        }
    ]


def _resolve_panne_catalogue(ligne, label):
    if not ligne.type_maintenance_id or not label:
        return None
    return PanneCatalogue.objects.filter(
        type_maintenance=ligne.type_maintenance,
        libelle__iexact=label.strip(),
    ).first()


def _save_subline_items(request, formset, allow_create=True, allow_delete=True):
    for ligne_form in formset.forms:
        if not hasattr(ligne_form, "cleaned_data"):
            continue
        if not ligne_form.cleaned_data or ligne_form.cleaned_data.get("DELETE"):
            continue

        ligne = ligne_form.instance
        if not ligne.pk:
            continue

        posted_ids = request.POST.getlist(f"subline-{ligne_form.prefix}-ids")
        labels = request.POST.getlist(f"subline-{ligne_form.prefix}-labels")
        quantities = request.POST.getlist(f"subline-{ligne_form.prefix}-quantites")
        article_ids = request.POST.getlist(f"subline-{ligne_form.prefix}-articles")

        kept_ids = []
        existing_by_id = {str(item.id): item for item in ligne.sous_lignes.all()}

        for index in range(max(len(labels), len(quantities), len(posted_ids), len(article_ids))):
            subline_id = posted_ids[index].strip() if index < len(posted_ids) and posted_ids[index] else ""
            label = labels[index].strip() if index < len(labels) and labels[index] else ""
            quantity_raw = quantities[index].strip() if index < len(quantities) and quantities[index] else "1"
            article_id = article_ids[index].strip() if index < len(article_ids) and article_ids[index] else ""
            article = ArticleStock.objects.filter(pk=article_id).first() if article_id else None
            if article and not label:
                label = article.libelle
            panne_catalogue = _resolve_panne_catalogue(ligne, label)

            if not label:
                continue

            if subline_id and subline_id in existing_by_id:
                subline = existing_by_id[subline_id]
                if subline.mouvement_stock_id and (
                    str(subline.article_stock_id or "") != str(article.id if article else "")
                    or str(subline.quantite) != str(quantity_raw or "1")
                ):
                    raise ValidationError(
                        f"La piece '{subline.libelle}' a deja ete sortie du stock. "
                        "Seul l'administrateur peut corriger ce mouvement manuellement."
                )
                subline.article_stock = article
                subline.panne_catalogue = panne_catalogue
                subline.libelle = label
                subline.quantite = quantity_raw or "1"
                if article and not subline.mouvement_stock_id:
                    subline.prix_unitaire = article.prix_unitaire or Decimal("0")
                subline.save()
                kept_ids.append(subline.id)
            elif allow_create:
                subline = MaintenanceSousLigne.objects.create(
                    maintenance_ligne=ligne,
                    article_stock=article,
                    panne_catalogue=panne_catalogue,
                    libelle=label,
                    quantite=quantity_raw or "1",
                    prix_unitaire=(article.prix_unitaire or Decimal("0")) if article else Decimal("0"),
                )
                kept_ids.append(subline.id)

        if allow_delete:
            ligne.sous_lignes.exclude(id__in=kept_ids).delete()


def _attach_achat_piece_rows(maintenance):
    piece_rows = []
    for ligne in maintenance.lignes.select_related("type_maintenance").prefetch_related("sous_lignes"):
        pieces = list(ligne.sous_lignes.all())
        if pieces:
            for piece in pieces:
                piece.is_stock_information = bool(piece.article_stock_id)
            piece_rows.append({"ligne": ligne, "pieces": pieces})
        else:
            piece_rows.append({"ligne": ligne, "pieces": []})
    return piece_rows


def _maintenance_invoice_rows(maintenance):
    rows = list(maintenance.factures_achat.select_related("fournisseur").all())
    if rows:
        return rows
    if maintenance.fournisseur_id or maintenance.numero_facture or maintenance.facture_fichier:
        return [
            {
                "fournisseur": maintenance.fournisseur,
                "numero_facture": maintenance.numero_facture,
                "facture_fichier": maintenance.facture_fichier,
                "is_legacy": True,
            }
        ]
    return []


def _sync_maintenance_invoice_snapshot(maintenance):
    first_invoice = maintenance.factures_achat.select_related("fournisseur").order_by("id").first()
    if first_invoice:
        maintenance.fournisseur = first_invoice.fournisseur
        maintenance.numero_facture = first_invoice.numero_facture
        maintenance.facture_fichier = first_invoice.facture_fichier.name if first_invoice.facture_fichier else ""
    else:
        maintenance.fournisseur = None
        maintenance.numero_facture = ""
        maintenance.facture_fichier = ""
    Maintenance.objects.filter(pk=maintenance.pk).update(
        fournisseur=maintenance.fournisseur,
        numero_facture=maintenance.numero_facture,
        facture_fichier=maintenance.facture_fichier,
    )


def _build_pannes_catalog():
    catalog = {}
    for panne in PanneCatalogue.objects.select_related("type_maintenance").order_by(
        "type_maintenance__libelle", "libelle"
    ):
        catalog.setdefault(str(panne.type_maintenance_id), []).append(
            {"id": panne.id, "label": panne.libelle}
        )
    return catalog


def _issue_stock_for_maintenance(maintenance, user):
    pieces = list(
        MaintenanceSousLigne.objects.select_related("article_stock", "mouvement_stock")
        .filter(maintenance_ligne__maintenance=maintenance, article_stock__isnull=False)
    )
    if not pieces:
        return

    for piece in pieces:
        if piece.mouvement_stock_id:
            continue
        stock_disponible = piece.article_stock.quantite_stock or Decimal("0")
        if piece.quantite > stock_disponible:
            raise ValidationError(
                f"Stock insuffisant pour {piece.article_stock.libelle}. "
                f"Disponible: {_format_amount(stock_disponible)} {piece.article_stock.unite}. "
                f"Demande: {_format_amount(piece.quantite)}."
            )

    for piece in pieces:
        if piece.mouvement_stock_id:
            continue
        mouvement = MouvementStock.objects.create(
            article=piece.article_stock,
            type_mouvement="sortie",
            quantite=piece.quantite,
            quantite_saisie=piece.quantite,
            unite_saisie=piece.article_stock.unite,
            reference=maintenance.reference,
            observation=f"Consommation maintenance {maintenance.reference} - {piece.libelle}",
            date_mouvement=timezone.now(),
            created_by=user,
        )
        MaintenanceSousLigne.objects.filter(pk=piece.pk).update(
            mouvement_stock=mouvement,
            prix_unitaire=mouvement.prix_unitaire,
            montant=mouvement.montant_net,
        )


def _save_achat_piece_prices(request, maintenance):
    def _parse_price(raw_value):
        normalized = (
            (raw_value or "0")
            .strip()
            .replace(" ", "")
            .replace("\u00a0", "")
            .replace(",", ".")
        )
        try:
            return Decimal(normalized or "0")
        except Exception:
            return Decimal("0")

    for ligne in maintenance.lignes.prefetch_related("sous_lignes"):
        if ligne.sous_lignes.exists():
            for piece in ligne.sous_lignes.all():
                if piece.article_stock_id:
                    continue
                piece.prix_unitaire = _parse_price(request.POST.get(f"piece-price-{piece.id}") or "0")
                piece.save()
        else:
            ligne.prix_unitaire = _parse_price(request.POST.get(f"ligne-price-{ligne.id}") or "0")
            ligne.save()


def _memoize_panne_prices_from_maintenance(maintenance):
    fournisseur = maintenance.factures_achat.select_related("fournisseur").order_by("id").first()
    if not fournisseur:
        return
    fournisseur = fournisseur.fournisseur
    pieces = MaintenanceSousLigne.objects.select_related("panne_catalogue").filter(
        maintenance_ligne__maintenance=maintenance,
        article_stock__isnull=True,
        panne_catalogue__isnull=False,
        prix_unitaire__gt=0,
    )
    for piece in pieces:
        PanneFournisseurPrix.objects.update_or_create(
            panne=piece.panne_catalogue,
            fournisseur=fournisseur,
            defaults={
                "montant": piece.prix_unitaire,
                "date_reference": timezone.localdate(),
                "observation": f"Memo auto {maintenance.reference}",
            },
        )


def _maintenance_tabs_context(active_tab):
    return {"active_tab": active_tab}


def _caissiere_users_queryset():
    return User.objects.filter(groups__name__in=["caissiere", "caissiere_soni", "caissiere_avena"], is_active=True).order_by("first_name", "last_name", "username").distinct()


def _caissiere_users_queryset_for_user(user):
    role = get_user_role(user)
    if is_admin_user(user) or role == "directeur":
        return _caissiere_users_queryset()
    if role == "comptable_sogefi":
        return User.objects.filter(groups__name="caissiere", is_active=True).order_by("first_name", "last_name", "username").distinct()
    if role == "comptable":
        return User.objects.filter(groups__name="caissiere_soni", is_active=True).order_by("first_name", "last_name", "username").distinct()
    if role == "comptable_avena":
        return User.objects.filter(groups__name="caissiere_avena", is_active=True).order_by("first_name", "last_name", "username").distinct()
    if role in {"caissiere", "caissiere_soni", "caissiere_avena"}:
        return User.objects.filter(id=user.id, is_active=True)
    return _caissiere_users_queryset()


def _caissiere_scope_filter_for_user(user):
    return {"caissiere__in": _caissiere_users_queryset_for_user(user)}


def _resolve_caissiere_for_request(request):
    role = get_user_role(request.user)
    if role in {"caissiere", "caissiere_soni", "caissiere_avena"}:
        return request.user

    caissiere_id = (request.GET.get("caissiere") or request.POST.get("caissiere") or "").strip()
    queryset = _caissiere_users_queryset_for_user(request.user)
    if caissiere_id:
        selected = queryset.filter(id=caissiere_id).first()
        if selected:
            return selected
    return queryset.first()


def _get_caisse_metrics(caissiere, date_from="", date_to=""):
    appro_qs = ApprovisionnementCaisse.objects.none()
    maintenance_qs = Maintenance.objects.none()
    depenses_qs = Depense.objects.none()
    solde_initial_obj = None
    solde_initial = Decimal("0")

    if caissiere:
        solde_initial_obj = SoldeInitialCaisse.objects.filter(caissiere=caissiere).first()
        solde_initial = solde_initial_obj.montant_initial if solde_initial_obj else Decimal("0")
        appro_qs = ApprovisionnementCaisse.objects.filter(caissiere=caissiere)
        maintenance_qs = _maintenance_queryset().filter(
            statut="payee",
            paiement_saisi_par=caissiere,
        )
        depenses_qs = Depense.objects.select_related(
            "type_depense",
            "operation",
            "operation__camion",
        ).prefetch_related("lignes").filter(
            statut=Depense.STATUT_PAYEE,
            mode_reglement=Depense.MODE_ESPECE,
            paiement_saisi_par=caissiere,
        )

    if date_from:
        appro_qs = appro_qs.filter(date_approvisionnement__gte=date_from)
        maintenance_qs = maintenance_qs.filter(date_paiement__gte=date_from)
        depenses_qs = depenses_qs.filter(date_paiement__gte=date_from)
    if date_to:
        appro_qs = appro_qs.filter(date_approvisionnement__lte=date_to)
        maintenance_qs = maintenance_qs.filter(date_paiement__lte=date_to)
        depenses_qs = depenses_qs.filter(date_paiement__lte=date_to)

    caisse_mouvements_qs = appro_qs.filter(
        nature_approvisionnement__in=[
            ApprovisionnementCaisse.NATURE_URGENCE_DG,
            ApprovisionnementCaisse.NATURE_REMBOURSEMENT_DG_ESPECE,
            ApprovisionnementCaisse.NATURE_RETOUR_CAISSE,
            ApprovisionnementCaisse.NATURE_CHEQUE_DIRECT,
        ]
    )
    total_appro = Decimal("0")
    total_retours_caisse = Decimal("0")
    total_remboursements_dg_caisse = Decimal("0")
    for mouvement in caisse_mouvements_qs:
        impact = mouvement.impact_caisse
        if impact > 0:
            total_appro += impact
        elif impact < 0:
            total_remboursements_dg_caisse += abs(impact)
        if mouvement.nature_approvisionnement == ApprovisionnementCaisse.NATURE_RETOUR_CAISSE:
            total_retours_caisse += mouvement.montant or Decimal("0")
    total_sorties_maintenance = maintenance_qs.aggregate(total=Sum("total_facture")).get("total") or Decimal("0")
    total_sorties_depenses = sum((depense.montant_comptable or Decimal("0")) for depense in depenses_qs)
    solde = solde_initial + total_appro - total_remboursements_dg_caisse - total_sorties_maintenance - total_sorties_depenses

    mouvements = []
    if solde_initial_obj:
        mouvements.append(
            {
                "date": solde_initial_obj.date_reference,
                "sort_key": (
                    solde_initial_obj.date_reference or timezone.localdate(),
                    solde_initial_obj.created_at or timezone.now(),
                    solde_initial_obj.id,
                ),
                "type": "Solde initial",
                "reference": "INITIAL",
                "designation": f"Ouverture de caisse - {caissiere.get_full_name() or caissiere.username}",
                "detail": solde_initial_obj.observation or "-",
                "entree": solde_initial,
                "sortie": Decimal("0"),
            }
        )
    for appro in caisse_mouvements_qs:
        impact_caisse = appro.impact_caisse
        mouvements.append(
            {
                "date": appro.date_approvisionnement,
                "sort_key": (
                    appro.date_approvisionnement or timezone.localdate(),
                    appro.created_at or timezone.now(),
                    appro.id,
                ),
                "type": "Remboursement DG" if impact_caisse < 0 else "Retour caisse" if appro.nature_approvisionnement == ApprovisionnementCaisse.NATURE_RETOUR_CAISSE else "Approvisionnement",
                "reference": appro.reference,
                "designation": appro.get_nature_approvisionnement_display(),
                "detail": appro.observation or "-",
                "entree": impact_caisse if impact_caisse > 0 else Decimal("0"),
                "sortie": abs(impact_caisse) if impact_caisse < 0 else Decimal("0"),
            }
        )

    for maintenance in maintenance_qs:
        invoice_rows = _maintenance_invoice_rows(maintenance)
        detail_items = []
        for invoice in invoice_rows:
            if isinstance(invoice, dict):
                fournisseur_value = str(invoice.get("fournisseur") or "-")
                numero_facture_value = invoice.get("numero_facture") or "-"
            else:
                fournisseur_value = str(getattr(invoice, "fournisseur", None) or "-")
                numero_facture_value = getattr(invoice, "numero_facture", "") or "-"
            detail_items.append(f"{fournisseur_value} / {numero_facture_value}")
        mouvements.append(
            {
                "date": maintenance.date_paiement,
                "sort_key": (
                    maintenance.date_paiement or timezone.localdate(),
                    maintenance.paiement_saisi_le or maintenance.date_creation or timezone.now(),
                    maintenance.id,
                ),
                "type": "Paiement maintenance",
                "reference": maintenance.reference,
                "designation": f"Camion {maintenance.camion.numero_tracteur}",
                "detail": " | ".join(detail_items) if detail_items else (maintenance.observation or "-"),
                "entree": Decimal("0"),
                "sortie": maintenance.total_facture or Decimal("0"),
            }
        )

    for depense in depenses_qs:
        detail_parts = []
        if depense.source_depense == Depense.SOURCE_CHARGEMENT and depense.operation_id:
            detail_parts.append(f"BL {depense.operation.numero_bl}")
        if depense.type_depense_id:
            detail_parts.append(depense.type_depense.libelle)
        if depense.numero_facture:
            detail_parts.append(f"Facture {depense.numero_facture}")
        mouvements.append(
            {
                "date": depense.date_paiement,
                "sort_key": (
                    depense.date_paiement or timezone.localdate(),
                    depense.paiement_saisi_le or depense.date_creation or timezone.now(),
                    depense.id,
                ),
                "type": "Depense BL" if depense.source_depense == Depense.SOURCE_CHARGEMENT else "Autre depense",
                "reference": depense.reference,
                "designation": depense.libelle_depense or depense.titre,
                "detail": " | ".join(detail_parts) if detail_parts else (depense.description or "-"),
                "entree": Decimal("0"),
                "sortie": depense.montant_comptable or Decimal("0"),
            }
        )

    mouvements.sort(key=lambda item: item["sort_key"])
    running_balance = Decimal("0")
    for mouvement in mouvements:
        running_balance += mouvement["entree"] - mouvement["sortie"]
        mouvement["solde"] = running_balance
        mouvement["entree_display"] = _format_amount(mouvement["entree"])
        mouvement["sortie_display"] = _format_amount(mouvement["sortie"])
        mouvement["solde_display"] = _format_amount(mouvement["solde"])

    return {
        "appro_qs": appro_qs,
        "appro_caisse_qs": caisse_mouvements_qs,
        "maintenance_qs": maintenance_qs,
        "depenses_qs": depenses_qs,
        "solde_initial_obj": solde_initial_obj,
        "solde_initial": solde_initial,
        "mouvements": mouvements,
        "total_appro": total_appro,
        "total_retours_caisse": total_retours_caisse,
        "total_remboursements_dg_caisse": total_remboursements_dg_caisse,
        "total_sorties_maintenance": total_sorties_maintenance,
        "total_sorties_depenses": total_sorties_depenses,
        "solde": solde,
    }


@role_required("logistique", "maintenancier", "caissiere", "comptable_sogefi", "responsable_achat", "dga_sogefi", "directeur", "admin")
def voir_facture_maintenance(request, id):
    facture = get_object_or_404(MaintenanceFacture, pk=id)
    return _safe_file_response(facture.facture_fichier)


def _get_dg_metrics(date_from="", date_to="", caissiere=None, user=None):
    dg_qs = ApprovisionnementCaisse.objects.filter(
        nature_approvisionnement__in=[
            ApprovisionnementCaisse.NATURE_URGENCE_DG,
            ApprovisionnementCaisse.NATURE_REMBOURSEMENT_DG_ESPECE,
            ApprovisionnementCaisse.NATURE_RETRAIT_REMBOURSEMENT,
        ]
    ).select_related("caissiere", "saisi_par")
    if caissiere:
        dg_qs = dg_qs.filter(caissiere=caissiere)
    elif user:
        dg_qs = dg_qs.filter(**_caissiere_scope_filter_for_user(user))

    if date_from:
        dg_qs = dg_qs.filter(date_approvisionnement__gte=date_from)
    if date_to:
        dg_qs = dg_qs.filter(date_approvisionnement__lte=date_to)

    total_avances = dg_qs.filter(
        nature_approvisionnement=ApprovisionnementCaisse.NATURE_URGENCE_DG
    ).aggregate(total=Sum("montant")).get("total") or Decimal("0")
    total_remboursements = dg_qs.filter(
        nature_approvisionnement__in=[
            ApprovisionnementCaisse.NATURE_RETRAIT_REMBOURSEMENT,
            ApprovisionnementCaisse.NATURE_REMBOURSEMENT_DG_ESPECE,
        ]
    ).aggregate(total=Sum("montant")).get("total") or Decimal("0")
    solde_dg = total_avances - total_remboursements

    mouvements = []
    running_balance = Decimal("0")
    for mouvement in dg_qs.order_by("date_approvisionnement", "id"):
        avance = mouvement.impact_dg_avance
        remboursement = mouvement.impact_dg_remboursement
        running_balance += avance - remboursement
        mouvements.append(
            {
                "date": mouvement.date_approvisionnement,
                "reference": mouvement.reference,
                "nature": mouvement.get_nature_approvisionnement_display(),
                "caissiere": mouvement.caissiere,
                "mode": mouvement.get_mode_approvisionnement_display(),
                "observation": mouvement.observation or "-",
                "reference_cheque": mouvement.reference_cheque,
                "banque_cheque": mouvement.banque_cheque,
                "date_cheque": mouvement.date_cheque,
                "avance": avance,
                "remboursement": remboursement,
                "solde": running_balance,
            }
        )

    return {
        "queryset": dg_qs.order_by("-date_approvisionnement", "-id"),
        "mouvements": list(reversed(mouvements)) if False else mouvements,
        "total_avances": total_avances,
        "total_remboursements": total_remboursements,
        "solde_dg": solde_dg,
    }


def _can_manage_stock(user):
    return is_admin_user(user) or get_user_role(user) in {"logistique", "directeur"}


def _build_stock_report_context(request):
    report_article = (request.GET.get("report_article") or "").strip()
    report_camion = (request.GET.get("report_camion") or "").strip()
    report_categorie = (request.GET.get("report_categorie") or "").strip()
    report_panne = (request.GET.get("report_panne") or "").strip()
    report_date_from = (request.GET.get("report_date_from") or "").strip()
    report_date_to = (request.GET.get("report_date_to") or "").strip()

    consommation_base_qs = MaintenanceSousLigne.objects.filter(
        article_stock__isnull=False,
        mouvement_stock__isnull=False,
    ).select_related(
        "article_stock",
        "maintenance_ligne__maintenance__camion",
        "mouvement_stock",
    )
    article_filter_choices = list(
        consommation_base_qs.values("article_stock_id", "article_stock__libelle")
        .distinct()
        .order_by("article_stock__libelle")
    )
    camion_filter_choices = list(
        consommation_base_qs.values(
            "maintenance_ligne__maintenance__camion_id",
            "maintenance_ligne__maintenance__camion__numero_tracteur",
            "maintenance_ligne__maintenance__camion__numero_citerne",
            "maintenance_ligne__maintenance__camion__code_camion",
        )
        .distinct()
        .order_by("maintenance_ligne__maintenance__camion__numero_tracteur")
    )
    categorie_filter_choices = list(
        MaintenanceSousLigne.objects.exclude(maintenance_ligne__type_maintenance__libelle="")
        .values("maintenance_ligne__type_maintenance__libelle")
        .distinct()
        .order_by("maintenance_ligne__type_maintenance__libelle")
    )
    panne_filter_choices = list(
        MaintenanceSousLigne.objects.exclude(libelle="")
        .values("libelle")
        .distinct()
        .order_by("libelle")
    )
    consommation_qs = consommation_base_qs
    if report_article:
        consommation_qs = consommation_qs.filter(article_stock_id=report_article)
    if report_camion:
        consommation_qs = consommation_qs.filter(maintenance_ligne__maintenance__camion_id=report_camion)
    if report_date_from:
        consommation_qs = consommation_qs.filter(mouvement_stock__date_mouvement__date__gte=report_date_from)
    if report_date_to:
        consommation_qs = consommation_qs.filter(mouvement_stock__date_mouvement__date__lte=report_date_to)

    report_total_quantite_raw = consommation_qs.aggregate(total=Sum("quantite")).get("total") or Decimal("0")
    report_total_montant_raw = consommation_qs.aggregate(total=Sum("mouvement_stock__montant_net")).get("total") or Decimal("0")
    report_total_passages = consommation_qs.count()
    report_total_camions = consommation_qs.values("maintenance_ligne__maintenance__camion_id").distinct().count()
    consommation_par_camion = list(
        consommation_qs
        .values(
            "maintenance_ligne__maintenance__camion__code_camion",
            "maintenance_ligne__maintenance__camion__numero_tracteur",
            "maintenance_ligne__maintenance__camion__numero_citerne",
            "article_stock__libelle",
            "article_stock__unite",
        )
        .annotate(
            total_quantite=Sum("quantite"),
            total_maintenances=Count("maintenance_ligne__maintenance", distinct=True),
            total_passages=Count("id"),
            total_montant=Sum("mouvement_stock__montant_net"),
        )
        .order_by("-total_montant", "-total_quantite", "article_stock__libelle")
    )
    top_camions_consommation = list(
        consommation_qs
        .values(
            "maintenance_ligne__maintenance__camion__code_camion",
            "maintenance_ligne__maintenance__camion__numero_tracteur",
            "maintenance_ligne__maintenance__camion__numero_citerne",
        )
        .annotate(
            total_quantite=Sum("quantite"),
            total_passages=Count("id"),
            total_montant=Sum("mouvement_stock__montant_net"),
        )
        .order_by("-total_montant", "-total_quantite")[:5]
    )
    top_articles_consommation = list(
        consommation_qs
        .values(
            "article_stock__libelle",
            "article_stock__unite",
        )
        .annotate(
            total_quantite=Sum("quantite"),
            total_passages=Count("id"),
            total_montant=Sum("mouvement_stock__montant_net"),
        )
        .order_by("-total_montant", "-total_quantite")[:5]
    )

    panne_base_qs = MaintenanceSousLigne.objects.filter(
        maintenance_ligne__maintenance__isnull=False,
    ).select_related(
        "article_stock",
        "panne_catalogue",
        "maintenance_ligne__type_maintenance",
        "maintenance_ligne__maintenance__camion",
        "mouvement_stock",
    )
    if report_camion:
        panne_base_qs = panne_base_qs.filter(maintenance_ligne__maintenance__camion_id=report_camion)
    if report_date_from:
        panne_base_qs = panne_base_qs.filter(maintenance_ligne__maintenance__date_debut__date__gte=report_date_from)
    if report_date_to:
        panne_base_qs = panne_base_qs.filter(maintenance_ligne__maintenance__date_debut__date__lte=report_date_to)
    if report_categorie:
        panne_base_qs = panne_base_qs.filter(maintenance_ligne__type_maintenance__libelle=report_categorie)
    if report_panne:
        panne_base_qs = panne_base_qs.filter(
            Q(panne_catalogue__libelle__icontains=report_panne) | Q(libelle__icontains=report_panne)
        )

    cout_panne_expr = Case(
        When(
            article_stock__isnull=False,
            then=Coalesce(F("mouvement_stock__montant_net"), Value(0)),
        ),
        default=Coalesce(F("montant"), Value(0)),
        output_field=DecimalField(max_digits=14, decimal_places=2),
    )
    panne_rapport_qs = panne_base_qs.annotate(
        categorie_panne=Coalesce(
            F("maintenance_ligne__type_maintenance__libelle"),
            Value("Sans categorie"),
        ),
        libelle_panne=Coalesce(
            F("panne_catalogue__libelle"),
            F("libelle"),
            Value("Sans libelle"),
        ),
        cout_consomme=cout_panne_expr,
    )
    report_total_pannes = panne_rapport_qs.count()
    report_total_montant_pannes_raw = (
        panne_rapport_qs.aggregate(total=Sum("cout_consomme")).get("total") or Decimal("0")
    )
    report_total_categories = (
        panne_rapport_qs.values("categorie_panne").distinct().count()
    )
    top_pannes_consommation = list(
        panne_rapport_qs
        .values("libelle_panne", "categorie_panne")
        .annotate(
            total_quantite=Sum("quantite"),
            total_passages=Count("id"),
            total_montant=Sum("cout_consomme"),
        )
        .order_by("-total_montant", "-total_quantite", "libelle_panne")[:5]
    )
    consommation_par_panne = list(
        panne_rapport_qs
        .values("categorie_panne", "libelle_panne")
        .annotate(
            total_quantite=Sum("quantite"),
            total_maintenances=Count("maintenance_ligne__maintenance", distinct=True),
            total_passages=Count("id"),
            total_montant=Sum("cout_consomme"),
        )
        .order_by("categorie_panne", "-total_montant", "libelle_panne")
    )
    return {
        "article_filter_choices": article_filter_choices,
        "camion_filter_choices": camion_filter_choices,
        "categorie_filter_choices": categorie_filter_choices,
        "panne_filter_choices": panne_filter_choices,
        "report_filter_values": {
            "article": report_article,
            "camion": report_camion,
            "categorie": report_categorie,
            "panne": report_panne,
            "date_from": report_date_from,
            "date_to": report_date_to,
        },
        "report_total_passages": report_total_passages,
        "report_total_camions": report_total_camions,
        "report_total_quantite": _format_amount(report_total_quantite_raw),
        "report_total_montant": _format_amount(report_total_montant_raw),
        "consommation_par_camion": consommation_par_camion,
        "top_camions_consommation": top_camions_consommation,
        "top_articles_consommation": top_articles_consommation,
        "report_total_pannes": report_total_pannes,
        "report_total_categories": report_total_categories,
        "report_total_montant_pannes": _format_amount(report_total_montant_pannes_raw),
        "top_pannes_consommation": top_pannes_consommation,
        "consommation_par_panne": consommation_par_panne,
    }


def _build_conversion_map(article, formset, principal_unit):
    conversions = {}
    if formset is not None:
        for form in formset.forms:
            if not hasattr(form, "cleaned_data") or not form.cleaned_data or form.cleaned_data.get("DELETE"):
                continue
            unite_source = (form.cleaned_data.get("unite_source") or "").strip().lower()
            quantite_equivalente = form.cleaned_data.get("quantite_equivalente")
            if unite_source and quantite_equivalente:
                conversions[unite_source] = Decimal(quantite_equivalente)
    if article is not None:
        for conversion in article.conversions.all():
            conversions.setdefault(conversion.unite_source.lower(), conversion.quantite_equivalente)
    conversions[(principal_unit or "").strip().lower()] = Decimal("1")
    return conversions


def _map_stock_article_error_field(field_name):
    field_map = {
        "prix_conditionnement": "prix_achat_saisi",
        "unite_prix_conditionnement": "unite_stock_saisie",
    }
    return field_map.get(field_name, field_name)


def _convert_to_principal_unit(quantity, source_unit, principal_unit, conversion_map):
    quantity = Decimal(quantity or 0)
    source_unit = (source_unit or principal_unit or "").strip().lower()
    principal_unit = (principal_unit or "").strip().lower()
    if source_unit == principal_unit:
        return quantity
    if source_unit in conversion_map:
        return quantity * Decimal(conversion_map[source_unit])
    raise ValidationError(
        {"unite_stock_saisie": f"Aucune conversion definie entre {source_unit} et {principal_unit}."}
    )


def _compute_display_purchase_price(article, source_unit):
    if article is None:
        return Decimal("0")
    source_unit = (source_unit or article.unite or "").strip().lower()
    coefficient = article.get_conversion_factor(source_unit)
    return (article.prix_unitaire or Decimal("0")) * coefficient


def _get_preferred_purchase_unit(article):
    if article is None:
        return "piece"
    stored = (article.unite_prix_conditionnement or "").strip().lower()
    principal = (article.unite or "").strip().lower() or "piece"
    if stored and stored != principal:
        return stored
    first_conversion = article.conversions.order_by("id").first()
    if first_conversion and first_conversion.unite_source:
        return first_conversion.unite_source.strip().lower()
    return principal


def _build_stock_unit_choices(principal_unit, formset=None, article=None, post_data=None):
    principal = (principal_unit or "").strip().lower() or "piece"
    choices = [(principal, principal)]
    seen = {principal}

    if post_data is not None:
        for key, value in post_data.items():
            if key.endswith("-unite_source"):
                unit = (value or "").strip().lower()
                if unit and unit not in seen:
                    seen.add(unit)
                    choices.append((unit, unit))

    if formset is not None:
        for form in formset.forms:
            unit = ""
            if hasattr(form, "cleaned_data") and form.cleaned_data and not form.cleaned_data.get("DELETE"):
                unit = (form.cleaned_data.get("unite_source") or "").strip().lower()
            elif form.is_bound:
                unit = (form.data.get(f"{form.prefix}-unite_source") or "").strip().lower()
            else:
                unit = (form.initial.get("unite_source") or "").strip().lower()
            if unit and unit not in seen:
                seen.add(unit)
                choices.append((unit, unit))

    if article is not None:
        for conversion in article.conversions.all():
            unit = (conversion.unite_source or "").strip().lower()
            if unit and unit not in seen:
                seen.add(unit)
                choices.append((unit, unit))

    return choices


def _garage_camions_catalog():
    chauffeurs_by_camion = {
        chauffeur.camion_id: chauffeur.nom
        for chauffeur in Chauffeur.objects.select_related("camion").filter(camion__isnull=False)
    }
    return [
        {
            "id": camion.id,
            "code_camion": camion.code_camion,
            "numero_tracteur": camion.numero_tracteur,
            "numero_citerne": camion.numero_citerne,
            "capacite": camion.capacite,
            "chauffeur": chauffeurs_by_camion.get(camion.id, ""),
        }
        for camion in Camion.objects.filter(est_affrete=False).order_by("numero_tracteur")
    ]


def garage_maintenances(request):
    historique = request.GET.get("scope") == "historique"
    maintenances_qs = _maintenance_queryset()
    if historique:
        maintenances_qs = maintenances_qs.filter(
            statut__in=["payee", "validee_stock", "rejetee_dga", "rejetee_dg"]
        )
    else:
        maintenances_qs = maintenances_qs.exclude(
            statut__in=["payee", "validee_stock", "rejetee_dga", "rejetee_dg"]
        )
    maintenances_qs, filter_values = _apply_maintenance_filters(request, maintenances_qs)
    maintenances = list(maintenances_qs)
    user_role = get_user_role(request.user)
    is_admin = is_admin_user(request.user)
    for maintenance in maintenances:
        maintenance.invoice_rows = _maintenance_invoice_rows(maintenance)
        maintenance.pricing_complete = maintenance.is_pricing_complete()
        maintenance.status_display_label = maintenance.get_statut_display()
        maintenance.can_terminate = False
        maintenance.can_reject_dga = (
            user_role == "dga"
            and maintenance.statut == "attente_dga"
        )
        maintenance.can_validate_dga = (
            user_role == "dga"
            and maintenance.statut == "attente_dga"
        )
        maintenance.can_reject_dg = (
            user_role == "directeur"
            and maintenance.statut == "attente_dg"
        )
        maintenance.can_validate_dg = (
            user_role == "directeur"
            and maintenance.statut == "attente_dg"
        )
        maintenance.can_enter_prices = (
            user_role == "logistique"
            and maintenance.statut == "attente_prix"
        )
        if maintenance.statut == "rejetee_dga":
            maintenance.validation_status_label = "Rejete par le DGA"
            maintenance.validation_status_variant = "danger"
        elif maintenance.statut == "rejetee_dg":
            maintenance.validation_status_label = "Rejete par le DG"
            maintenance.validation_status_variant = "danger"
        elif maintenance.statut == "payee":
            maintenance.validation_status_label = "Payee"
            maintenance.validation_status_variant = "ok"
        elif maintenance.statut == "validee_stock":
            maintenance.validation_status_label = "Validee stock"
            maintenance.validation_status_variant = "ok"
        elif maintenance.statut == "en_cours":
            maintenance.validation_status_label = "Diagnostic en cours"
            maintenance.validation_status_variant = "warning"
        elif maintenance.statut == "attente_prix":
            maintenance.validation_status_label = "En attente de saisie de prix"
            maintenance.validation_status_variant = "warning"
        elif maintenance.statut == "attente_dga":
            maintenance.validation_status_label = "En attente validation DGA"
            maintenance.validation_status_variant = "warning"
        elif maintenance.statut == "attente_dg":
            maintenance.validation_status_label = "En attente validation DG"
            maintenance.validation_status_variant = "warning"
        elif maintenance.statut == "attente_paiement":
            if maintenance.mode_paiement == Maintenance.MODE_CHEQUE:
                maintenance.status_display_label = "En attente de paiement chez le comptable"
                maintenance.validation_status_label = "En attente de paiement comptable"
            else:
                maintenance.status_display_label = "En attente de paiement chez la caissiere"
                maintenance.validation_status_label = "En attente de paiement caisse"
            maintenance.validation_status_variant = "warning"
        else:
            maintenance.validation_status_label = maintenance.get_statut_display()
            maintenance.validation_status_variant = "ok"
    depenses_camions = (
        Maintenance.objects.values("camion__numero_tracteur", "camion__numero_citerne")
        .annotate(total_depense=Sum("total_facture"), total_maintenances=Count("id"))
        .order_by("-total_depense", "-total_maintenances")[:5]
    )
    return render(
        request,
        "maintenance/garage.html",
        {
            "maintenances": maintenances,
            "historique": historique,
            "depenses_camions": depenses_camions,
            "filter_values": filter_values,
            "statut_choices": Maintenance.STATUT_CHOICES,
            "is_admin_maintenance": is_admin,
            **_maintenance_tabs_context("garage"),
        },
    )


def achat_maintenances(request):
    historique = request.GET.get("scope") == "historique"
    user_role = get_user_role(request.user)
    can_edit_achat = is_admin_user(request.user) or user_role == "logistique"
    maintenances = _maintenance_queryset()
    if historique:
        maintenances = maintenances.exclude(
            Q(validation_dga_at__isnull=True) & Q(statut__in=["attente_prix", "attente_dga"])
        )
    else:
        maintenances = maintenances.filter(
            validation_dga_at__isnull=True,
            statut__in=["attente_prix", "attente_dga"],
        )
    maintenances, filter_values = _apply_maintenance_filters(request, maintenances)
    maintenances = list(maintenances)
    for maintenance in maintenances:
        maintenance.invoice_rows = _maintenance_invoice_rows(maintenance)
        maintenance.can_edit_prices = bool(
            can_edit_achat
            and maintenance.validation_dga_at is None
            and maintenance.statut in {"attente_prix", "attente_dga"}
        )
        if maintenance.is_stock_only():
            maintenance.pricing_action_label = "Valider le stock"
        elif maintenance.statut == "attente_dga" or (maintenance.total_facture or Decimal("0")) > 0:
            maintenance.pricing_action_label = "Modifier"
        else:
            maintenance.pricing_action_label = "Saisir les prix"
    return render(
        request,
        "maintenance/achat.html",
        {
            "maintenances": maintenances,
            "historique": historique,
            "filter_values": filter_values,
            "statut_choices": Maintenance.STATUT_CHOICES,
            "is_admin_maintenance": is_admin_user(request.user),
            "can_edit_achat": can_edit_achat,
            **_maintenance_tabs_context("achat"),
        },
    )


def paiements_maintenances(request):
    historique = request.GET.get("scope") == "historique"
    role = get_user_role(request.user)
    q = (request.GET.get("q") or "").strip()
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()
    payment_type = (request.GET.get("payment_type") or "").strip()

    maintenances = _maintenance_queryset()
    if historique:
        maintenances = maintenances.filter(statut="payee")
    else:
        maintenances = maintenances.filter(statut="attente_paiement")

    if not is_admin_user(request.user):
        if role == "caissiere":
            maintenances = maintenances.filter(mode_paiement=Maintenance.MODE_ESPECE)
        elif role == "caissiere_soni":
            maintenances = maintenances.none()
        elif role == "caissiere_avena":
            maintenances = maintenances.none()
        elif role == "comptable":
            maintenances = maintenances.none()
        elif role == "comptable_avena":
            maintenances = maintenances.none()
        elif role == "comptable_sogefi":
            maintenances = maintenances.filter(mode_paiement=Maintenance.MODE_CHEQUE)

    if q:
        maintenances = maintenances.filter(
            Q(reference__icontains=q)
            | Q(numero_facture__icontains=q)
            | Q(camion__code_camion__icontains=q)
            | Q(camion__numero_tracteur__icontains=q)
            | Q(camion__numero_citerne__icontains=q)
            | Q(camion__chauffeur__nom__icontains=q)
            | Q(fournisseur__nom_fournisseur__icontains=q)
            | Q(fournisseur__entreprise__icontains=q)
        ).distinct()
    if date_from:
        if historique:
            maintenances = maintenances.filter(date_paiement__gte=date_from)
        else:
            maintenances = maintenances.filter(date_debut__date__gte=date_from)
    if date_to:
        if historique:
            maintenances = maintenances.filter(date_paiement__lte=date_to)
        else:
            maintenances = maintenances.filter(date_debut__date__lte=date_to)

    depenses = Depense.objects.select_related(
        "demandeur",
        "type_depense",
        "fournisseur",
        "operation",
        "commande",
        "paiement_saisi_par",
    ).prefetch_related("lignes")

    if historique:
        depenses = depenses.filter(statut=Depense.STATUT_PAYEE)
    else:
        depenses = depenses.filter(
            statut__in=[
                Depense.STATUT_ATTENTE_PAIEMENT_CAISSIERE,
                Depense.STATUT_ATTENTE_PAIEMENT_COMPTABLE,
            ]
        )

    if not is_admin_user(request.user):
        if role == "caissiere":
            if historique:
                depenses = depenses.filter(
                    statut=Depense.STATUT_PAYEE,
                    mode_reglement=Depense.MODE_ESPECE,
                    paiement_saisi_par=request.user,
                )
            else:
                depenses = depenses.filter(
                    Q(statut=Depense.STATUT_ATTENTE_PAIEMENT_CAISSIERE, source_depense=Depense.SOURCE_CHARGEMENT)
                    | Q(
                        statut=Depense.STATUT_ATTENTE_PAIEMENT_CAISSIERE,
                        source_depense=Depense.SOURCE_GENERALE,
                        entite_depense=Depense.ENTITE_SOGEFI,
                    )
                )
        elif role == "caissiere_soni":
            if historique:
                depenses = depenses.filter(
                    statut=Depense.STATUT_PAYEE,
                    source_depense=Depense.SOURCE_GENERALE,
                    entite_depense=Depense.ENTITE_SONI,
                    mode_reglement=Depense.MODE_ESPECE,
                    paiement_saisi_par=request.user,
                )
            else:
                depenses = depenses.filter(
                    statut=Depense.STATUT_ATTENTE_PAIEMENT_CAISSIERE,
                    source_depense=Depense.SOURCE_GENERALE,
                    entite_depense=Depense.ENTITE_SONI,
                )
        elif role == "caissiere_avena":
            if historique:
                depenses = depenses.filter(
                    statut=Depense.STATUT_PAYEE,
                    source_depense=Depense.SOURCE_GENERALE,
                    entite_depense=Depense.ENTITE_AVENA,
                    mode_reglement=Depense.MODE_ESPECE,
                    paiement_saisi_par=request.user,
                )
            else:
                depenses = depenses.filter(
                    statut=Depense.STATUT_ATTENTE_PAIEMENT_CAISSIERE,
                    source_depense=Depense.SOURCE_GENERALE,
                    entite_depense=Depense.ENTITE_AVENA,
                )
        elif role == "comptable":
            depenses = depenses.filter(
                source_depense=Depense.SOURCE_GENERALE,
                entite_depense=Depense.ENTITE_SONI,
                mode_reglement=Depense.MODE_CHEQUE,
            )
        elif role == "comptable_avena":
            depenses = depenses.filter(
                source_depense=Depense.SOURCE_GENERALE,
                entite_depense=Depense.ENTITE_AVENA,
                mode_reglement=Depense.MODE_CHEQUE,
            )
        elif role == "comptable_sogefi":
            depenses = depenses.filter(
                mode_reglement=Depense.MODE_CHEQUE,
            ).exclude(
                source_depense=Depense.SOURCE_GENERALE,
                entite_depense=Depense.ENTITE_SONI,
            )

    if q:
        depenses = depenses.filter(
            Q(reference__icontains=q)
            | Q(titre__icontains=q)
            | Q(description__icontains=q)
            | Q(demandeur__username__icontains=q)
            | Q(fournisseur__nom_fournisseur__icontains=q)
            | Q(fournisseur__entreprise__icontains=q)
            | Q(operation__numero_bl__icontains=q)
            | Q(operation__camion__numero_tracteur__icontains=q)
            | Q(operation__camion__numero_citerne__icontains=q)
            | Q(operation__chauffeur__nom__icontains=q)
            | Q(numero_facture__icontains=q)
            | Q(numero_cheque__icontains=q)
        ).distinct()
    if date_from:
        depenses = depenses.filter((Q(date_paiement__gte=date_from) if historique else Q(date_creation__date__gte=date_from)))
    if date_to:
        depenses = depenses.filter((Q(date_paiement__lte=date_to) if historique else Q(date_creation__date__lte=date_to)))

    payment_items = []
    current_cash_balance = None
    if role in {"caissiere", "caissiere_soni", "caissiere_avena"}:
        current_cash_balance = _get_caisse_metrics(request.user)["solde"]

    maintenance_destination_label = "Caissiere" if role == "caissiere" else "Comptable SOGEFI"
    depense_destination_cheque_label = "Comptable SOGEFI"
    depense_destination_espece_label = "Caissiere"
    if role in {"comptable_avena", "caissiere_avena"}:
        depense_destination_cheque_label = "Comptable Avena"
        depense_destination_espece_label = "Caissiere Avena"
    elif role in {"comptable", "caissiere_soni"}:
        depense_destination_cheque_label = "Comptable SONI"
        depense_destination_espece_label = "Caissiere SONI"

    if payment_type in {"maintenance", ""}:
        for maintenance in maintenances:
            invoice_rows = _maintenance_invoice_rows(maintenance)
            fournisseur_label = " / ".join(
                str(invoice.fournisseur)
                for invoice in invoice_rows
                if getattr(invoice, "fournisseur", None)
            )
            numero_factures = ", ".join(
                str(invoice.numero_facture)
                for invoice in invoice_rows
                if getattr(invoice, "numero_facture", None)
            )
            camion_label = maintenance.camion.numero_tracteur if maintenance.camion_id else "-"
            if maintenance.camion_id and maintenance.camion.numero_citerne:
                camion_label = f"{camion_label} / {maintenance.camion.numero_citerne}"
            montant = maintenance.total_facture or Decimal("0")
            is_cash_item = maintenance.mode_paiement == Maintenance.MODE_ESPECE
            payment_items.append(
                {
                    "family": "maintenance",
                    "id": maintenance.id,
                    "reference": maintenance.reference,
                    "item_type": "Maintenance",
                    "designation": camion_label,
                    "tiers": fournisseur_label or str(maintenance.fournisseur or maintenance.prestataire or "-"),
                    "support": numero_factures or "-",
                    "mode_paiement": maintenance.get_mode_paiement_display() if maintenance.mode_paiement else "-",
                    "beneficiaire": maintenance.beneficiaire_cheque or "-",
                    "montant_raw": montant,
                    "montant_display": f"{_format_amount(montant)} GNF",
                    "date_raw": maintenance.date_paiement if historique and maintenance.date_paiement else (maintenance.date_debut.date() if maintenance.date_debut else None),
                    "date_display": date_format(maintenance.date_paiement if historique and maintenance.date_paiement else maintenance.date_debut, "d/m/Y"),
                    "action_url": f"/maintenance/paiements/modifier/{maintenance.id}/",
                    "action_label": "Consulter" if historique else "Payer",
                    "support_url": f"/maintenance/facture/{invoice_rows[0].id}/" if invoice_rows and getattr(invoice_rows[0], "facture_fichier", None) else "",
                    "is_cash_item": is_cash_item,
                    "balance_warning": bool(not historique and is_cash_item and current_cash_balance is not None and montant > current_cash_balance),
                    "destination_label": maintenance_destination_label if is_cash_item else "Comptable SOGEFI",
                }
            )

    if payment_type in {"depense_bl", "depense_generale", ""}:
        depense_items = []
        seen_chargement_operations = set()
        for depense in depenses.order_by("-date_creation", "-id"):
            if depense.source_depense == Depense.SOURCE_CHARGEMENT and depense.operation_id:
                if depense.operation_id in seen_chargement_operations:
                    continue
                seen_chargement_operations.add(depense.operation_id)
                depense = (
                    Depense.objects.filter(
                        source_depense=Depense.SOURCE_CHARGEMENT,
                        operation_id=depense.operation_id,
                        portee_chargement=Depense.PORTEE_BL,
                    )
                    .select_related("demandeur", "type_depense", "fournisseur", "operation", "commande", "paiement_saisi_par")
                    .prefetch_related("lignes")
                    .order_by("date_creation", "id")
                    .first()
                    or depense
                )
            current_type = "depense_bl" if depense.source_depense == Depense.SOURCE_CHARGEMENT else "depense_generale"
            if payment_type and payment_type != current_type:
                continue
            montant = depense.montant_comptable or Decimal("0")
            if current_type == "depense_bl" and montant <= 0:
                continue
            operation = depense.operation
            is_cash_item = depense.mode_reglement == Depense.MODE_ESPECE
            depense_items.append(
                {
                    "family": "depense",
                    "id": depense.id,
                    "reference": depense.reference,
                    "item_type": "Depense BL" if current_type == "depense_bl" else "Autre depense",
                    "designation": depense.titre or depense.libelle_depense or "-",
                    "tiers": str(depense.fournisseur or depense.demandeur or "-"),
                    "support": depense.numero_facture or (operation.numero_bl if operation else "-"),
                    "mode_paiement": depense.get_mode_reglement_display() if depense.mode_reglement else "-",
                    "beneficiaire": depense.beneficiaire_cheque or depense.receveur_nom or "-",
                    "montant_raw": montant,
                    "montant_display": f"{_format_amount(montant)} GNF",
                    "date_raw": depense.date_paiement if historique and depense.date_paiement else (depense.date_creation.date() if depense.date_creation else None),
                    "date_display": date_format(depense.date_paiement if historique and depense.date_paiement else depense.date_creation, "d/m/Y"),
                    "action_url": f"/depenses/paiement/{depense.id}/",
                    "action_label": "Consulter" if historique else "Payer",
                    "support_url": f"/depenses/piece/{depense.id}/" if depense.piece_justificative else "",
                    "is_cash_item": is_cash_item,
                    "balance_warning": bool(not historique and is_cash_item and current_cash_balance is not None and montant > current_cash_balance),
                    "destination_label": depense_destination_espece_label if is_cash_item else depense_destination_cheque_label,
                }
            )
        payment_items.extend(depense_items)

    payment_items.sort(key=lambda item: (item["date_raw"] or timezone.localdate(), item["reference"]), reverse=True)
    total_montant = sum((item["montant_raw"] for item in payment_items), Decimal("0"))
    payment_type_choices = [
        ("", "Tous les paiements"),
        ("maintenance", "Maintenance"),
        ("depense_bl", "Depenses BL"),
        ("depense_generale", "Autres depenses"),
    ]
    page_subtitle = "Maintenance, depenses BL et autres depenses sont regroupes dans une seule file de traitement."
    if role in {"comptable_avena", "caissiere_avena"}:
        payment_type_choices = [
            ("", "Toutes les depenses Avena"),
            ("depense_generale", "Depenses internes Avena"),
        ]
        page_subtitle = "Les depenses internes Avena sont regroupees ici, sans maintenance, BL ni paiements SOGEFI."
    return render(
        request,
        "maintenance/paiements.html",
        {
            "payment_items": payment_items,
            "historique": historique,
            "filter_values": {
                "q": q,
                "date_from": date_from,
                "date_to": date_to,
                "payment_type": payment_type,
            },
            "payment_type_choices": payment_type_choices,
            "page_subtitle": page_subtitle,
            "can_edit_paiement": True,
            "is_admin_maintenance": is_admin_user(request.user),
            "payment_count": len(payment_items),
            "total_montant_display": _format_amount(total_montant),
            "solde_caisse_display": _format_amount(current_cash_balance) if current_cash_balance is not None else None,
            **_maintenance_tabs_context("paiements"),
        },
    )


def appro_caisse(request):
    caissieres = _caissiere_users_queryset_for_user(request.user)
    if request.method == "POST" and request.POST.get("form_type") == "solde_initial":
        solde_instance = None
        selected_solde_caissiere_id = (request.POST.get("caissiere") or "").strip()
        if selected_solde_caissiere_id:
            solde_instance = SoldeInitialCaisse.objects.filter(caissiere_id=selected_solde_caissiere_id).first()
        solde_form = SoldeInitialCaisseForm(request.POST, instance=solde_instance, caissiere_queryset=caissieres)
        if solde_form.is_valid():
            solde_initial = solde_form.save(commit=False)
            solde_initial.saisi_par = request.user
            solde_initial.save()
            journaliser_action(
                request.user,
                "Caisse",
                "Solde initial caisse",
                solde_initial.caissiere.username,
                f"{request.user.username} a defini le solde initial de caisse de {solde_initial.caissiere.username} a {solde_initial.montant_initial} GNF.",
            )
            messages.success(request, "Le solde initial de caisse a ete enregistre.")
            return redirect(f"/maintenance/caisse/appro/?caissiere={solde_initial.caissiere_id}")
        form = ApprovisionnementCaisseForm(caissiere_queryset=caissieres, user=request.user)
        messages.error(request, "Impossible d'enregistrer le solde initial. Verifiez les champs.")
    elif request.method == "POST":
        form = ApprovisionnementCaisseForm(request.POST, caissiere_queryset=caissieres, user=request.user)
        if form.is_valid():
            appro = form.save(commit=False)
            appro.saisi_par = request.user
            appro.save()
            journaliser_action(
                request.user,
                "Caisse",
                "Mouvement caisse",
                appro.reference,
                f"{request.user.username} a saisi {appro.get_nature_approvisionnement_display().lower()} pour {appro.caissiere.username} ({appro.montant} GNF).",
            )
            messages.success(request, f"Le mouvement {appro.reference} a ete enregistre.")
            return redirect("appro_caisse")
        solde_form = SoldeInitialCaisseForm(caissiere_queryset=caissieres)
        messages.error(request, "Impossible d'enregistrer ce mouvement. Verifiez les champs.")
    else:
        initial = {}
        if caissieres.count() == 1:
            initial["caissiere"] = caissieres.first()
        form = ApprovisionnementCaisseForm(caissiere_queryset=caissieres, initial=initial, user=request.user)
        solde_form = SoldeInitialCaisseForm(caissiere_queryset=caissieres, initial=initial)

    selected_caissiere = _resolve_caissiere_for_request(request)
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()
    metrics = _get_caisse_metrics(selected_caissiere, date_from=date_from, date_to=date_to)
    if selected_caissiere:
        solde_form = SoldeInitialCaisseForm(
            caissiere_queryset=caissieres,
            instance=metrics["solde_initial_obj"],
            initial={"caissiere": selected_caissiere},
        )
    approvisionnements = ApprovisionnementCaisse.objects.select_related("caissiere", "saisi_par").filter(
        **_caissiere_scope_filter_for_user(request.user)
    )
    if selected_caissiere:
        approvisionnements = approvisionnements.filter(caissiere=selected_caissiere)
    if date_from:
        approvisionnements = approvisionnements.filter(date_approvisionnement__gte=date_from)
    if date_to:
        approvisionnements = approvisionnements.filter(date_approvisionnement__lte=date_to)

    return render(
        request,
        "maintenance/appro_caisse.html",
        {
            "form": form,
            "caissieres": caissieres,
            "selected_caissiere": selected_caissiere,
            "filter_values": {
                "caissiere": str(selected_caissiere.id) if selected_caissiere else "",
                "date_from": date_from,
                "date_to": date_to,
            },
            "banques": Banque.objects.filter(actif=True).order_by("nom"),
            "banque_form": BanqueForm(),
            "solde_form": solde_form,
            "approvisionnements": approvisionnements[:20],
            "solde_initial": _format_amount(metrics["solde_initial"]),
            "total_appro": _format_amount(metrics["total_appro"]),
            "total_retours_caisse": _format_amount(metrics["total_retours_caisse"]),
            "total_remboursements_dg_caisse": _format_amount(metrics["total_remboursements_dg_caisse"]),
            "total_sorties": _format_amount(
                metrics["total_sorties_maintenance"]
                + metrics["total_sorties_depenses"]
                + metrics["total_remboursements_dg_caisse"]
            ),
            "solde_caisse": _format_amount(metrics["solde"]),
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("appro_caisse"),
        },
    )


def situation_caisse(request):
    caissieres = _caissiere_users_queryset_for_user(request.user)
    selected_caissiere = _resolve_caissiere_for_request(request)
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()
    metrics = _get_caisse_metrics(selected_caissiere, date_from=date_from, date_to=date_to)
    mouvements = list(metrics["mouvements"])

    return render(
        request,
        "maintenance/situation_caisse.html",
        {
            "caissieres": caissieres,
            "selected_caissiere": selected_caissiere,
            "filter_values": {
                "caissiere": str(selected_caissiere.id) if selected_caissiere else "",
                "date_from": date_from,
                "date_to": date_to,
            },
            "mouvements": mouvements,
            "total_appro": _format_amount(metrics["total_appro"]),
            "total_retours_caisse": _format_amount(metrics["total_retours_caisse"]),
            "total_remboursements_dg_caisse": _format_amount(metrics["total_remboursements_dg_caisse"]),
            "total_sorties_maintenance": _format_amount(metrics["total_sorties_maintenance"]),
            "total_sorties_depenses": _format_amount(metrics["total_sorties_depenses"]),
            "solde_caisse": _format_amount(metrics["solde"]),
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("situation_caisse"),
        },
    )


def situation_dg(request):
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()
    user_role = get_user_role(request.user)
    caissiere = request.user if user_role in {"caissiere", "caissiere_soni", "caissiere_avena"} else None
    metrics = _get_dg_metrics(date_from=date_from, date_to=date_to, caissiere=caissiere, user=request.user)
    return render(
        request,
        "maintenance/situation_dg.html",
        {
            "filter_values": {
                "date_from": date_from,
                "date_to": date_to,
            },
            "mouvements": metrics["mouvements"],
            "total_avances": _format_amount(metrics["total_avances"]),
            "total_remboursements": _format_amount(metrics["total_remboursements"]),
            "solde_dg": _format_amount(metrics["solde_dg"]),
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("situation_dg"),
        },
    )


def rapport_maintenances(request):
    user_role = get_user_role(request.user)
    if user_role in CAISSIERE_ROLES or user_role in {"comptable", "comptable_avena"}:
        internal_entity = _caisse_internal_entity_for_role(user_role)
        q = (request.GET.get("q") or "").strip()
        date_from = (request.GET.get("date_from") or "").strip()
        date_to = (request.GET.get("date_to") or "").strip()
        type_depense = (request.GET.get("type_depense") or "").strip()

        if user_role in CAISSIERE_ROLES:
            maintenance_qs = _maintenance_queryset().filter(
                statut="payee",
            ).filter(
                Q(paiement_saisi_par=request.user)
                | Q(paiement_saisi_par__isnull=True)
            )
            payment_filter = {"paiement_saisi_par": request.user}
        else:
            maintenance_qs = _maintenance_queryset().none()
            payment_filter = {}
        depenses_qs = Depense.objects.select_related(
            "type_depense",
            "fournisseur",
            "operation",
            "operation__camion",
            "operation__chauffeur",
            "commande",
            "paiement_saisi_par",
        ).filter(
            statut=Depense.STATUT_PAYEE,
            mode_reglement=Depense.MODE_ESPECE,
            **payment_filter,
        )
        if internal_entity:
            maintenance_qs = maintenance_qs.none()
            depenses_qs = depenses_qs.filter(
                source_depense=Depense.SOURCE_GENERALE,
                entite_depense=internal_entity,
            )

        if q:
            maintenance_qs = maintenance_qs.filter(
                Q(reference__icontains=q)
                | Q(camion__numero_tracteur__icontains=q)
                | Q(camion__numero_citerne__icontains=q)
                | Q(prestataire__icontains=q)
                | Q(fournisseur__nom_fournisseur__icontains=q)
                | Q(numero_facture__icontains=q)
                | Q(receveur_nom__icontains=q)
            ).distinct()
            depenses_qs = depenses_qs.filter(
                Q(reference__icontains=q)
                | Q(titre__icontains=q)
                | Q(libelle_depense__icontains=q)
                | Q(description__icontains=q)
                | Q(type_depense__libelle__icontains=q)
                | Q(fournisseur__nom_fournisseur__icontains=q)
                | Q(numero_facture__icontains=q)
                | Q(operation__numero_bl__icontains=q)
                | Q(receveur_nom__icontains=q)
            ).distinct()
        if date_from:
            maintenance_qs = maintenance_qs.filter(date_paiement__gte=date_from)
            depenses_qs = depenses_qs.filter(date_paiement__gte=date_from)
        if date_to:
            maintenance_qs = maintenance_qs.filter(date_paiement__lte=date_to)
            depenses_qs = depenses_qs.filter(date_paiement__lte=date_to)
        if internal_entity and type_depense == "maintenance":
            type_depense = ""
        if type_depense == "maintenance":
            depenses_qs = depenses_qs.none()
        elif type_depense == "bl":
            maintenance_qs = maintenance_qs.none()
            depenses_qs = depenses_qs.filter(source_depense=Depense.SOURCE_CHARGEMENT)
        elif type_depense == "generale":
            maintenance_qs = maintenance_qs.none()
            depenses_qs = depenses_qs.filter(source_depense=Depense.SOURCE_GENERALE)

        payment_rows = []
        total_montant = Decimal("0")

        for maintenance in maintenance_qs:
            maintenance.invoice_rows = _maintenance_invoice_rows(maintenance)
            fournisseur_label = "-"
            facture_label = maintenance.numero_facture or "-"
            invoice_details = []
            if maintenance.invoice_rows:
                first_invoice = maintenance.invoice_rows[0]
                if isinstance(first_invoice, dict):
                    fournisseur_label = str(first_invoice.get("fournisseur") or "-")
                    facture_label = first_invoice.get("numero_facture") or facture_label
                else:
                    fournisseur_label = str(getattr(first_invoice, "fournisseur", None) or "-")
                    facture_label = getattr(first_invoice, "numero_facture", "") or facture_label
                for invoice in maintenance.invoice_rows:
                    if isinstance(invoice, dict):
                        fournisseur_value = str(invoice.get("fournisseur") or "-")
                        numero_facture_value = invoice.get("numero_facture") or ""
                    else:
                        fournisseur_value = str(getattr(invoice, "fournisseur", None) or "-")
                        numero_facture_value = getattr(invoice, "numero_facture", "") or ""
                    invoice_details.append(
                        " / ".join(
                            part
                            for part in [
                                fournisseur_value,
                                numero_facture_value,
                            ]
                            if part
                        )
                    )
            elif maintenance.prestataire:
                fournisseur_label = maintenance.prestataire
            montant = maintenance.total_facture or Decimal("0")
            total_montant += montant
            detail_sections = []
            designation_parts = [
                f"Camion {maintenance.camion.numero_tracteur}",
            ]
            if maintenance.camion.numero_citerne:
                designation_parts.append(f"Citerne {maintenance.camion.numero_citerne}")
            if invoice_details:
                detail_sections.append(
                    {
                        "label": "Factures",
                        "items": invoice_details,
                    }
                )
            if maintenance.observation:
                detail_sections.append(
                    {
                        "label": "Observation",
                        "items": [maintenance.observation],
                    }
                )
            article_items = []
            for ligne in maintenance.lignes.all():
                sous_lignes = list(ligne.sous_lignes.all())
                if sous_lignes:
                    for sous_ligne in sous_lignes:
                        article_items.append(
                            f"{sous_ligne.libelle} - Qté {_format_amount(sous_ligne.quantite)} x PU {_format_amount(sous_ligne.prix_unitaire)} = {_format_amount(sous_ligne.montant)} GNF"
                        )
                else:
                    article_items.append(
                        f"{ligne.libelle} - Qté {_format_amount(ligne.quantite)} x PU {_format_amount(ligne.prix_unitaire)} = {_format_amount(ligne.montant)} GNF"
                    )
            if article_items:
                detail_sections.append(
                    {
                        "label": "Pieces / travaux",
                        "items": article_items,
                    }
                )
            payment_rows.append(
                {
                    "famille": "Maintenance",
                    "famille_key": "maintenance",
                    "reference": maintenance.reference,
                    "designation": " | ".join(designation_parts),
                    "designation_lines": designation_parts,
                    "detail_sections": detail_sections,
                    "piece": facture_label,
                    "montant": montant,
                    "montant_display": _format_amount(montant),
                    "date_paiement": maintenance.date_paiement,
                    "mode_paiement": maintenance.mode_paiement or "Espece",
                    "beneficiaire": maintenance.receveur_nom or "-",
                }
            )

        for depense in depenses_qs:
            montant = depense.montant_comptable or Decimal("0")
            total_montant += montant
            detail_sections = []
            if depense.source_depense == Depense.SOURCE_CHARGEMENT:
                famille = "Depense BL"
                famille_key = "bl"
                designation = depense.libelle_depense or depense.titre
                designation_lines = [designation]
                detail_parts = []
                if depense.operation_id:
                    detail_parts.append(f"BL {depense.operation.numero_bl}")
                    if getattr(depense.operation, "camion", None):
                        designation_lines.append(f"Camion {depense.operation.camion.numero_tracteur}")
                    if getattr(depense.operation, "chauffeur", None):
                        designation_lines.append(str(depense.operation.chauffeur))
                if depense.commande_id:
                    detail_parts.append(f"Commande {depense.commande.reference}")
            else:
                famille = "Depense generale"
                famille_key = "generale"
                designation = depense.libelle_depense or depense.titre
                designation_lines = [designation]
                detail_parts = []
            if depense.type_depense_id:
                detail_parts.append(f"Type {depense.type_depense.libelle}")
            if depense.fournisseur_id:
                detail_parts.append(f"Fournisseur {str(depense.fournisseur)}")
            if depense.numero_facture:
                detail_parts.append(f"Facture {depense.numero_facture}")
            if depense.description:
                detail_parts.append(f"Motif {depense.description}")
            if detail_parts:
                detail_sections.append(
                    {
                        "label": "Details",
                        "items": detail_parts,
                    }
                )
            ligne_items = [
                f"{ligne.designation} - Qté {_format_amount(ligne.quantite)} x PU {_format_amount(ligne.prix_unitaire)} = {_format_amount(ligne.montant)} GNF"
                for ligne in depense.lignes.all()
            ]
            if ligne_items:
                detail_sections.append(
                    {
                        "label": "Lignes de depense",
                        "items": ligne_items,
                    }
                )
            payment_rows.append(
                {
                    "famille": famille,
                    "famille_key": famille_key,
                    "reference": depense.reference,
                    "designation": designation,
                    "designation_lines": designation_lines,
                    "detail_sections": detail_sections,
                    "piece": depense.numero_cheque or depense.numero_facture or "-",
                    "montant": montant,
                    "montant_display": _format_amount(montant),
                    "date_paiement": depense.date_paiement,
                    "mode_paiement": depense.get_mode_reglement_display(),
                    "beneficiaire": depense.receveur_nom or "-",
                }
            )

        payment_rows.sort(
            key=lambda item: (
                item["date_paiement"] or timezone.datetime.min.date(),
                item["reference"],
            ),
            reverse=True,
        )

        filter_values = {
            "q": q,
            "date_from": date_from,
            "date_to": date_to,
            "type_depense": type_depense,
        }
        return render(
            request,
            "maintenance/rapport_caisse.html",
            {
                "payment_rows": payment_rows,
                "filter_values": filter_values,
                "total_paiements": len(payment_rows),
                "total_montant": _format_amount(total_montant),
                "total_maintenance": sum(1 for row in payment_rows if row["famille_key"] == "maintenance"),
                "total_depenses_bl": sum(1 for row in payment_rows if row["famille_key"] == "bl"),
                "total_depenses_generales": sum(1 for row in payment_rows if row["famille_key"] == "generale"),
                "caisse_report_internal_only": bool(internal_entity),
                "caisse_report_entity_label": "Avena" if internal_entity == Depense.ENTITE_AVENA else "SONI" if internal_entity == Depense.ENTITE_SONI else "",
                "is_admin_maintenance": is_admin_user(request.user),
                **_maintenance_tabs_context("rapports"),
            },
        )

    maintenances_qs = _maintenance_queryset()
    q = (request.GET.get("q") or "").strip()
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()
    statut = (request.GET.get("statut") or "").strip()
    panne = (request.GET.get("panne") or "").strip()

    if q:
        maintenances_qs = maintenances_qs.filter(
            Q(reference__icontains=q)
            | Q(camion__code_camion__icontains=q)
            | Q(camion__numero_tracteur__icontains=q)
            | Q(camion__numero_citerne__icontains=q)
            | Q(camion__chauffeur__nom__icontains=q)
            | Q(fournisseur__nom_fournisseur__icontains=q)
            | Q(prestataire__icontains=q)
            | Q(numero_facture__icontains=q)
        ).distinct()
    if date_from:
        maintenances_qs = maintenances_qs.filter(date_paiement__gte=date_from)
    if date_to:
        maintenances_qs = maintenances_qs.filter(date_paiement__lte=date_to)
    if statut:
        maintenances_qs = maintenances_qs.filter(statut=statut)
    if panne:
        maintenances_qs = maintenances_qs.filter(
            Q(lignes__sous_lignes__panne_catalogue__libelle__icontains=panne)
            | Q(lignes__sous_lignes__libelle__icontains=panne)
        ).distinct()

    filter_values = {
        "q": q,
        "date_from": date_from,
        "date_to": date_to,
        "statut": statut,
        "panne": panne,
    }
    maintenances = list(maintenances_qs)
    panne_choices = list(
        PanneCatalogue.objects.order_by("libelle").values_list("libelle", flat=True).distinct()
    )

    total_situations = len(maintenances)
    total_montant = sum((maintenance.total_facture or Decimal("0")) for maintenance in maintenances)
    total_cout_global = sum((maintenance.total_global_information or Decimal("0")) for maintenance in maintenances)
    total_payees = sum(1 for maintenance in maintenances if maintenance.statut == "payee")
    total_attente = sum(
        1
        for maintenance in maintenances
        if maintenance.statut in {"attente_prix", "attente_dga", "attente_dg", "attente_paiement"}
    )

    for maintenance in maintenances:
        maintenance.invoice_rows = _maintenance_invoice_rows(maintenance)
        maintenance.chauffeur_nom = (
            Chauffeur.objects.filter(camion=maintenance.camion)
            .values_list("nom", flat=True)
            .first()
            or "-"
        )
        diagnostics = []
        for ligne in maintenance.lignes.all():
            piece_labels = [piece.libelle for piece in ligne.sous_lignes.all()]
            diagnostic = f"{ligne.type_maintenance.libelle}: {ligne.libelle}"
            if piece_labels:
                diagnostic += f" ({', '.join(piece_labels)})"
            diagnostics.append(diagnostic)
        maintenance.diagnostics_resume = diagnostics
        maintenance.cout_global_affiche = _format_amount(maintenance.total_global_information)
        maintenance.stock_information_affiche = _format_amount(maintenance.total_information_stock)

    return render(
        request,
        "maintenance/rapport.html",
        {
            "maintenances": maintenances,
            "filter_values": filter_values,
            "statut_choices": Maintenance.STATUT_CHOICES,
            "panne_choices": panne_choices,
            "total_situations": total_situations,
            "total_montant": _format_amount(total_montant),
            "total_cout_global": _format_amount(total_cout_global),
            "total_payees": total_payees,
            "total_attente": total_attente,
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("rapports"),
        },
    )


def stock_maintenances(request):
    q = (request.GET.get("q") or "").strip()
    articles = ArticleStock.objects.select_related("fournisseur")
    if q:
        articles = articles.filter(
            Q(code_article__icontains=q)
            | Q(libelle__icontains=q)
            | Q(categorie__icontains=q)
            | Q(fournisseur__nom_fournisseur__icontains=q)
        )
    articles = list(articles.order_by("libelle"))
    recent_movements = list(
        MouvementStock.objects.select_related("article", "created_by")[:12]
    )
    maintenance_refs = {
        mouvement.reference
        for mouvement in recent_movements
        if (mouvement.reference or "").startswith("MAIN")
    }
    maintenance_by_ref = {
        maintenance.reference: maintenance
        for maintenance in Maintenance.objects.select_related("camion").filter(reference__in=maintenance_refs)
    }
    for mouvement in recent_movements:
        maintenance = maintenance_by_ref.get(mouvement.reference)
        if maintenance and maintenance.camion_id:
            mouvement.camion_display = maintenance.camion.numero_tracteur
            if maintenance.camion.numero_citerne:
                mouvement.camion_display += f" / {maintenance.camion.numero_citerne}"
        else:
            mouvement.camion_display = "-"
    total_articles = len(articles)
    articles_en_alerte = sum(1 for article in articles if article.en_alerte)
    valeur_stock = sum((article.valeur_stock for article in articles), Decimal("0"))
    for article in articles:
        article.valeur_stock_affichee = f"{_format_amount(article.valeur_stock)} GNF"
        article.prix_conditionnement_affiche = article.prix_conditionnement_display or "-"
        article.prix_unitaire_affiche = article.prix_unitaire_display
        article.prix_equivalent_affiche = article.prix_conditionnement_equivalent_display
    can_manage_stock = _can_manage_stock(request.user)
    return render(
        request,
        "maintenance/stock.html",
        {
            "articles": articles,
            "recent_movements": recent_movements,
            "filter_values": {"q": q},
            "total_articles": total_articles,
            "articles_en_alerte": articles_en_alerte,
            "valeur_stock": _format_amount(valeur_stock),
            "can_manage_stock": can_manage_stock,
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("stock"),
        },
    )


def rapport_pannes_main_oeuvre(request):
    q = (request.GET.get("q") or "").strip()
    camion_id = (request.GET.get("camion") or "").strip()
    categorie = (request.GET.get("categorie") or "").strip()
    libelle = (request.GET.get("libelle") or "").strip()
    source = (request.GET.get("source") or "").strip()
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()

    maintenances = list(
        _maintenance_queryset().prefetch_related(
            "lignes__type_maintenance",
            "lignes__sous_lignes__article_stock",
            "lignes__sous_lignes__panne_catalogue",
        )
    )

    rows = []
    for maintenance in maintenances:
        camion = maintenance.camion
        if camion_id and str(camion.id) != camion_id:
            continue

        maintenance_date = maintenance.date_fin or maintenance.date_debut
        if date_from and maintenance_date and maintenance_date.date().isoformat() < date_from:
            continue
        if date_to and maintenance_date and maintenance_date.date().isoformat() > date_to:
            continue

        camion_label = camion.numero_tracteur
        if camion.numero_citerne:
            camion_label = f"{camion.numero_tracteur} / {camion.numero_citerne}"

        fournisseur_label = str(maintenance.fournisseur) if maintenance.fournisseur_id else (maintenance.prestataire or "-")

        for ligne in maintenance.lignes.all():
            categorie_label = ligne.type_maintenance.libelle if ligne.type_maintenance_id else "-"

            if (ligne.prix_unitaire or Decimal("0")) > 0 or (ligne.montant or Decimal("0")) > 0:
                row = {
                    "date": maintenance_date,
                    "date_display": maintenance_date.strftime("%d/%m/%Y") if maintenance_date else "-",
                    "reference": maintenance.reference,
                    "camion_id": camion.id,
                    "camion_label": camion_label,
                    "categorie": categorie_label,
                    "libelle": ligne.libelle or categorie_label,
                    "source": "Main d'oeuvre",
                    "quantite": ligne.quantite or Decimal("0"),
                    "quantite_display": _format_amount(ligne.quantite or Decimal("0")),
                    "prix_unitaire": ligne.prix_unitaire or Decimal("0"),
                    "prix_unitaire_display": _format_amount(ligne.prix_unitaire or Decimal("0")),
                    "montant": ligne.montant or Decimal("0"),
                    "montant_display": _format_amount(ligne.montant or Decimal("0")),
                    "fournisseur": fournisseur_label,
                    "statut": maintenance.get_statut_display(),
                }
                rows.append(row)

            for piece in ligne.sous_lignes.all():
                row = {
                    "date": maintenance_date,
                    "date_display": maintenance_date.strftime("%d/%m/%Y") if maintenance_date else "-",
                    "reference": maintenance.reference,
                    "camion_id": camion.id,
                    "camion_label": camion_label,
                    "categorie": categorie_label,
                    "libelle": piece.libelle or (piece.panne_catalogue.libelle if piece.panne_catalogue_id else "-"),
                    "source": "Stock" if piece.article_stock_id else "Hors stock",
                    "quantite": piece.quantite or Decimal("0"),
                    "quantite_display": _format_amount(piece.quantite or Decimal("0")),
                    "prix_unitaire": piece.prix_unitaire or Decimal("0"),
                    "prix_unitaire_display": _format_amount(piece.prix_unitaire or Decimal("0")),
                    "montant": piece.montant or Decimal("0"),
                    "montant_display": _format_amount(piece.montant or Decimal("0")),
                    "fournisseur": fournisseur_label,
                    "statut": maintenance.get_statut_display(),
                }
                rows.append(row)

    categories = sorted({row["categorie"] for row in rows if row["categorie"] and row["categorie"] != "-"})
    libelles = sorted({row["libelle"] for row in rows if row["libelle"] and row["libelle"] != "-"})
    camions = sorted(
        {(
            maintenance.camion.id,
            f"{maintenance.camion.numero_tracteur}{' / ' + maintenance.camion.numero_citerne if maintenance.camion.numero_citerne else ''}",
        ) for maintenance in maintenances},
        key=lambda item: item[1],
    )

    filtered_rows = []
    q_lower = q.lower()
    libelle_lower = libelle.lower()
    for row in rows:
        if categorie and row["categorie"] != categorie:
            continue
        if source and row["source"] != source:
            continue
        if libelle and libelle_lower not in row["libelle"].lower():
            continue
        if q and not (
            q_lower in row["reference"].lower()
            or q_lower in row["camion_label"].lower()
            or q_lower in row["categorie"].lower()
            or q_lower in row["libelle"].lower()
            or q_lower in row["fournisseur"].lower()
        ):
            continue
        filtered_rows.append(row)

    filtered_rows.sort(key=lambda item: (item["date"] or timezone.now(), item["reference"], item["libelle"]), reverse=True)

    total_lignes = len(filtered_rows)
    total_camions = len({row["camion_id"] for row in filtered_rows})
    total_maintenances = len({row["reference"] for row in filtered_rows})
    total_montant = sum((row["montant"] for row in filtered_rows), Decimal("0"))

    grouped_camions = {}
    grouped_libelles = {}
    for row in filtered_rows:
        camion_group = grouped_camions.setdefault(
            row["camion_label"],
            {"camion_label": row["camion_label"], "nb_fois": 0, "montant": Decimal("0")},
        )
        camion_group["nb_fois"] += 1
        camion_group["montant"] += row["montant"]

        libelle_group = grouped_libelles.setdefault(
            (row["categorie"], row["libelle"]),
            {"categorie": row["categorie"], "libelle": row["libelle"], "nb_fois": 0, "montant": Decimal("0")},
        )
        libelle_group["nb_fois"] += 1
        libelle_group["montant"] += row["montant"]

    top_camions = sorted(grouped_camions.values(), key=lambda item: (-item["montant"], -item["nb_fois"], item["camion_label"]))[:5]
    top_libelles = sorted(grouped_libelles.values(), key=lambda item: (-item["montant"], -item["nb_fois"], item["libelle"]))[:5]

    for item in top_camions:
        item["montant_display"] = _format_amount(item["montant"])
    for item in top_libelles:
        item["montant_display"] = _format_amount(item["montant"])

    return render(
        request,
        "maintenance/rapport_pannes.html",
        {
            "rows": filtered_rows,
            "camion_choices": camions,
            "categorie_choices": categories,
            "libelle_choices": libelles,
            "source_choices": ["Main d'oeuvre", "Stock", "Hors stock"],
            "filter_values": {
                "q": q,
                "camion": camion_id,
                "categorie": categorie,
                "libelle": libelle,
                "source": source,
                "date_from": date_from,
                "date_to": date_to,
            },
            "total_lignes": total_lignes,
            "total_camions": total_camions,
            "total_maintenances": total_maintenances,
            "total_montant": _format_amount(total_montant),
            "top_camions": top_camions,
            "top_libelles": top_libelles,
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("rapport_pannes"),
        },
    )


def rapport_stock_maintenances(request):
    report_context = _build_stock_report_context(request)
    return render(
        request,
        "maintenance/rapport_stock.html",
        {
            **report_context,
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("rapport_stock"),
        },
    )


def ajouter_article_stock(request):
    if request.method == "POST":
        formset = ArticleStockConversionFormSet(request.POST)
        principal_unit = ((request.POST.get("unite") or "piece").strip().lower() or "piece")
        unite_choices = _build_stock_unit_choices(principal_unit, formset=formset, post_data=request.POST)
        form = ArticleStockForm(request.POST, unite_choices=unite_choices)
        if form.is_valid() and formset.is_valid():
            try:
                with transaction.atomic():
                    article = form.save(commit=False)
                    principal_unit = (form.cleaned_data.get("unite") or "").strip().lower()
                    source_unit = (form.cleaned_data.get("unite_stock_saisie") or principal_unit).strip().lower()
                    conversion_map = _build_conversion_map(None, formset, principal_unit)
                    article.quantite_stock = _convert_to_principal_unit(
                        form.cleaned_data.get("quantite_stock"),
                        source_unit,
                        principal_unit,
                        conversion_map,
                    )
                    prix_achat_saisi = Decimal(form.cleaned_data.get("prix_achat_saisi") or 0)
                    remise_globale = Decimal(form.cleaned_data.get("remise_globale") or 0)
                    coefficient_prix = conversion_map.get(source_unit, Decimal("1"))
                    montant_brut = (Decimal(form.cleaned_data.get("quantite_stock") or 0) * prix_achat_saisi)
                    if remise_globale > montant_brut:
                        raise ValidationError({"remise_globale": "La remise ne peut pas depasser le montant brut du stock initial."})
                    montant_net = montant_brut - remise_globale
                    article.prix_conditionnement = prix_achat_saisi
                    article.unite_prix_conditionnement = source_unit
                    article.prix_unitaire = (
                        montant_net / article.quantite_stock if article.quantite_stock and montant_net > 0 else Decimal("0")
                    )
                    article.save()
                    formset.instance = article
                    formset.save()
            except ValidationError as error:
                if hasattr(error, "message_dict"):
                    for field_name, field_errors in error.message_dict.items():
                        form.add_error(
                            _map_stock_article_error_field(field_name),
                            field_errors[0] if field_errors else str(error),
                        )
                else:
                    form.add_error("unite_stock_saisie", str(error))
            else:
                journaliser_action(
                    request.user,
                    "Maintenance",
                    "Creation article de stock",
                    article.code_article,
                    f"{request.user.username} a ajoute l'article de stock {article.code_article} - {article.libelle}.",
                )
                messages.success(request, f"L'article {article.libelle} a ete ajoute au stock.")
                return redirect("stock_maintenances")
    else:
        formset = ArticleStockConversionFormSet()
        unite_choices = _build_stock_unit_choices("piece", formset=formset)
        form = ArticleStockForm(
            initial={"unite_stock_saisie": "piece", "prix_achat_saisi": Decimal("0"), "remise_globale": Decimal("0")},
            unite_choices=unite_choices,
        )
    return render(
        request,
        "maintenance/ajouter_article_stock.html",
        {
            "form": form,
            "formset": formset,
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("stock"),
        },
    )


def modifier_article_stock(request, id):
    article = get_object_or_404(ArticleStock, id=id)
    if request.method == "POST":
        formset = ArticleStockConversionFormSet(request.POST, instance=article)
        principal_unit = ((request.POST.get("unite") or article.unite).strip().lower() or article.unite)
        unite_choices = _build_stock_unit_choices(principal_unit, formset=formset, article=article, post_data=request.POST)
        form = ArticleStockForm(request.POST, instance=article, unite_choices=unite_choices)
        if form.is_valid() and formset.is_valid():
            try:
                with transaction.atomic():
                    article = form.save(commit=False)
                    principal_unit = (form.cleaned_data.get("unite") or article.unite).strip().lower()
                    source_unit = (form.cleaned_data.get("unite_stock_saisie") or principal_unit).strip().lower()
                    conversion_map = _build_conversion_map(article, formset, principal_unit)
                    article.quantite_stock = _convert_to_principal_unit(
                        form.cleaned_data.get("quantite_stock"),
                        source_unit,
                        principal_unit,
                        conversion_map,
                    )
                    prix_achat_saisi = Decimal(form.cleaned_data.get("prix_achat_saisi") or 0)
                    remise_globale = Decimal(form.cleaned_data.get("remise_globale") or 0)
                    coefficient_prix = conversion_map.get(source_unit, Decimal("1"))
                    montant_brut = (Decimal(form.cleaned_data.get("quantite_stock") or 0) * prix_achat_saisi)
                    if remise_globale > montant_brut:
                        raise ValidationError({"remise_globale": "La remise ne peut pas depasser le montant brut du stock initial."})
                    montant_net = montant_brut - remise_globale
                    article.prix_conditionnement = prix_achat_saisi
                    article.unite_prix_conditionnement = source_unit
                    article.prix_unitaire = (
                        montant_net / article.quantite_stock if article.quantite_stock and montant_net > 0 else Decimal("0")
                    )
                    article.save()
                    formset.save()
            except ValidationError as error:
                if hasattr(error, "message_dict"):
                    for field_name, field_errors in error.message_dict.items():
                        form.add_error(
                            _map_stock_article_error_field(field_name),
                            field_errors[0] if field_errors else str(error),
                        )
                else:
                    form.add_error("unite_stock_saisie", str(error))
            else:
                journaliser_action(
                    request.user,
                    "Maintenance",
                    "Modification article de stock",
                    article.code_article,
                    f"{request.user.username} a modifie l'article de stock {article.code_article}.",
                )
                messages.success(request, f"L'article {article.libelle} a ete mis a jour.")
                return redirect("stock_maintenances")
    else:
        unite_choices = _build_stock_unit_choices(article.unite, article=article)
        preferred_unit = _get_preferred_purchase_unit(article)
        form = ArticleStockForm(
            instance=article,
            initial={
                "unite_stock_saisie": preferred_unit,
                "prix_achat_saisi": article.prix_conditionnement or _compute_display_purchase_price(article, preferred_unit),
                "remise_globale": Decimal("0"),
            },
            unite_choices=unite_choices,
        )
        formset = ArticleStockConversionFormSet(instance=article)
    return render(
        request,
        "maintenance/modifier_article_stock.html",
        {
            "form": form,
            "formset": formset,
            "article": article,
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("stock"),
        },
    )


@require_POST
def supprimer_article_stock(request, id):
    if not is_admin_user(request.user):
        messages.error(request, "Seul l'administrateur peut supprimer un article de stock.")
        return redirect("stock_maintenances")

    article = get_object_or_404(ArticleStock, id=id)
    if article.mouvements.exists() or article.maintenances_consommees.exists():
        messages.error(
            request,
            "Impossible de supprimer cet article car il a deja des mouvements ou des maintenances liees.",
        )
        return redirect("stock_maintenances")

    label = f"{article.code_article} - {article.libelle}"
    article.delete()
    journaliser_action(
        request.user,
        "Maintenance",
        "Suppression article de stock",
        label,
        f"{request.user.username} a supprime l'article de stock {label}.",
    )
    messages.success(request, f"L'article {label} a ete supprime.")
    return redirect("stock_maintenances")


def ajouter_mouvement_stock(request, article_id):
    article = get_object_or_404(ArticleStock, id=article_id)
    conversions_map = {
        conversion.unite_source.lower(): float(conversion.quantite_equivalente)
        for conversion in article.conversions.all()
    }
    conversions_map[(article.unite or "").strip().lower()] = 1.0
    if request.method == "POST":
        form = MouvementStockForm(request.POST, article=article)
        form.instance.article = article
        if form.is_valid():
            mouvement = form.save(commit=False)
            mouvement.article = article
            mouvement.created_by = request.user
            mouvement.save()
            journaliser_action(
                request.user,
                "Maintenance",
                "Mouvement de stock",
                article.code_article,
                f"{request.user.username} a enregistre un mouvement {mouvement.get_type_mouvement_display().lower()} pour {article.code_article}.",
            )
            messages.success(request, f"Le mouvement de stock pour {article.libelle} a ete enregistre.")
            return redirect("stock_maintenances")
    else:
        form = MouvementStockForm(
            article=article,
            initial={
                "unite_saisie": article.unite,
                "prix_conditionnement": article.prix_conditionnement or Decimal("0"),
                "remise": Decimal("0"),
                "fournisseur": article.fournisseur_id,
            },
        )
    return render(
        request,
        "maintenance/ajouter_mouvement_stock.html",
        {
            "form": form,
            "article": article,
            "article_conversions_json": mark_safe(json.dumps(conversions_map)),
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("stock"),
        },
    )


def export_situation_caisse_xls(request):
    try:
        from openpyxl import Workbook
    except ImportError:
        return HttpResponse(
            "Le module openpyxl n'est pas installe sur cet environnement Python.",
            status=503,
            content_type="text/plain; charset=utf-8",
        )

    selected_caissiere = _resolve_caissiere_for_request(request)
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()
    metrics = _get_caisse_metrics(selected_caissiere, date_from=date_from, date_to=date_to)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Situation caisse"
    sheet.append(["Situation de caisse"])
    sheet.append(["Caissiere", selected_caissiere.get_full_name() or selected_caissiere.username if selected_caissiere else ""])
    sheet.append(["Periode", f"{date_from or 'Debut'} - {date_to or 'Aujourd hui'}"])
    sheet.append(["Entrees caisse", float(metrics["total_appro"])])
    sheet.append(["Remboursements DG caisse", float(metrics["total_remboursements_dg_caisse"])])
    sheet.append(["Sorties maintenance", float(metrics["total_sorties_maintenance"])])
    sheet.append(["Sorties depenses", float(metrics["total_sorties_depenses"])])
    sheet.append(["Solde caisse", float(metrics["solde"])])
    sheet.append([])
    sheet.append(["Date", "Type", "Reference", "Designation", "Detail", "Entree", "Sortie", "Solde"])
    for mouvement in metrics["mouvements"]:
        sheet.append(
            [
                mouvement["date"].strftime("%d/%m/%Y") if mouvement.get("date") else "",
                mouvement.get("type", ""),
                mouvement.get("reference", ""),
                mouvement.get("designation", ""),
                mouvement.get("detail", ""),
                float(mouvement.get("entree") or Decimal("0")),
                float(mouvement.get("sortie") or Decimal("0")),
                float(mouvement.get("solde") or Decimal("0")),
            ]
        )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="situation_caisse.xlsx"'
    workbook.save(response)
    return response


def _rapport_caisse_export_rows(request):
    user_role = get_user_role(request.user)
    internal_entity = _caisse_internal_entity_for_role(user_role)
    q = (request.GET.get("q") or "").strip()
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()
    type_depense = (request.GET.get("type_depense") or "").strip()

    if user_role in CAISSIERE_ROLES:
        maintenance_qs = _maintenance_queryset().filter(statut="payee").filter(
            Q(paiement_saisi_par=request.user) | Q(paiement_saisi_par__isnull=True)
        )
        payment_filter = {"paiement_saisi_par": request.user}
    elif internal_entity:
        maintenance_qs = _maintenance_queryset().none()
        payment_filter = {}
    else:
        maintenance_qs = _maintenance_queryset().filter(statut="payee").filter(
            Q(paiement_saisi_par=request.user) | Q(paiement_saisi_par__isnull=True)
        )
        payment_filter = {"paiement_saisi_par": request.user}
    depenses_qs = Depense.objects.select_related(
        "type_depense",
        "fournisseur",
        "operation",
        "operation__camion",
        "operation__chauffeur",
        "commande",
        "paiement_saisi_par",
    ).prefetch_related("lignes").filter(
        statut=Depense.STATUT_PAYEE,
        mode_reglement=Depense.MODE_ESPECE,
        **payment_filter,
    )
    if internal_entity:
        maintenance_qs = maintenance_qs.none()
        depenses_qs = depenses_qs.filter(
            source_depense=Depense.SOURCE_GENERALE,
            entite_depense=internal_entity,
        )

    if q:
        maintenance_qs = maintenance_qs.filter(
            Q(reference__icontains=q)
            | Q(camion__numero_tracteur__icontains=q)
            | Q(camion__numero_citerne__icontains=q)
            | Q(prestataire__icontains=q)
            | Q(fournisseur__nom_fournisseur__icontains=q)
            | Q(numero_facture__icontains=q)
            | Q(receveur_nom__icontains=q)
        ).distinct()
        depenses_qs = depenses_qs.filter(
            Q(reference__icontains=q)
            | Q(titre__icontains=q)
            | Q(libelle_depense__icontains=q)
            | Q(description__icontains=q)
            | Q(type_depense__libelle__icontains=q)
            | Q(fournisseur__nom_fournisseur__icontains=q)
            | Q(numero_facture__icontains=q)
            | Q(operation__numero_bl__icontains=q)
            | Q(receveur_nom__icontains=q)
        ).distinct()
    if date_from:
        maintenance_qs = maintenance_qs.filter(date_paiement__gte=date_from)
        depenses_qs = depenses_qs.filter(date_paiement__gte=date_from)
    if date_to:
        maintenance_qs = maintenance_qs.filter(date_paiement__lte=date_to)
        depenses_qs = depenses_qs.filter(date_paiement__lte=date_to)
    if internal_entity and type_depense == "maintenance":
        type_depense = ""
    if type_depense == "maintenance":
        depenses_qs = depenses_qs.none()
    elif type_depense == "bl":
        maintenance_qs = maintenance_qs.none()
        depenses_qs = depenses_qs.filter(source_depense=Depense.SOURCE_CHARGEMENT)
    elif type_depense == "generale":
        maintenance_qs = maintenance_qs.none()
        depenses_qs = depenses_qs.filter(source_depense=Depense.SOURCE_GENERALE)

    rows = []
    for maintenance in maintenance_qs:
        invoice_rows = _maintenance_invoice_rows(maintenance)
        invoice_details = []
        for invoice in invoice_rows:
            if isinstance(invoice, dict):
                fournisseur_value = str(invoice.get("fournisseur") or "-")
                numero_facture_value = invoice.get("numero_facture") or ""
            else:
                fournisseur_value = str(getattr(invoice, "fournisseur", None) or "-")
                numero_facture_value = getattr(invoice, "numero_facture", "") or ""
            invoice_details.append(" / ".join(part for part in [fournisseur_value, numero_facture_value] if part))
        designation_parts = [f"Camion {maintenance.camion.numero_tracteur}"]
        if maintenance.camion.numero_citerne:
            designation_parts.append(f"Citerne {maintenance.camion.numero_citerne}")
        rows.append(
            {
                "date": maintenance.date_paiement,
                "type": "Maintenance",
                "reference": maintenance.reference,
                "designation": " | ".join(designation_parts),
                "detail": " | ".join(invoice_details) if invoice_details else (maintenance.observation or ""),
                "beneficiaire": maintenance.receveur_nom or "",
                "mode": maintenance.get_mode_paiement_display() if maintenance.mode_paiement else "Espece",
                "montant": maintenance.total_facture or Decimal("0"),
            }
        )

    for depense in depenses_qs:
        detail_parts = []
        if depense.source_depense == Depense.SOURCE_CHARGEMENT and depense.operation_id:
            detail_parts.append(f"BL {depense.operation.numero_bl}")
        if depense.commande_id:
            detail_parts.append(f"Commande {depense.commande.reference}")
        if depense.type_depense_id:
            detail_parts.append(depense.type_depense.libelle)
        if depense.fournisseur_id:
            detail_parts.append(str(depense.fournisseur))
        if depense.numero_facture:
            detail_parts.append(f"Facture {depense.numero_facture}")
        rows.append(
            {
                "date": depense.date_paiement,
                "type": "Depense BL" if depense.source_depense == Depense.SOURCE_CHARGEMENT else "Autre depense",
                "reference": depense.reference,
                "designation": depense.libelle_depense or depense.titre,
                "detail": " | ".join(detail_parts) if detail_parts else (depense.description or ""),
                "beneficiaire": depense.receveur_nom or "",
                "mode": depense.get_mode_reglement_display(),
                "montant": depense.montant_comptable or Decimal("0"),
            }
        )

    rows.sort(key=lambda item: (item["date"] or timezone.datetime.min.date(), item["reference"]), reverse=True)
    return rows, {
        "q": q,
        "date_from": date_from,
        "date_to": date_to,
        "type_depense": type_depense,
    }


def export_rapport_caisse_xls(request):
    try:
        from openpyxl import Workbook
    except ImportError:
        return HttpResponse(
            "Le module openpyxl n'est pas installe sur cet environnement Python.",
            status=503,
            content_type="text/plain; charset=utf-8",
        )

    rows, filters = _rapport_caisse_export_rows(request)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rapport caisse"
    sheet.append(["Rapport caisse"])
    sheet.append(["Recherche", filters["q"]])
    sheet.append(["Periode", f"{filters['date_from'] or 'Debut'} - {filters['date_to'] or 'Aujourd hui'}"])
    sheet.append(["Type", filters["type_depense"] or "Tous"])
    sheet.append(["Total paiements", len(rows)])
    sheet.append(["Montant total", float(sum((row["montant"] for row in rows), Decimal("0")))])
    sheet.append([])
    sheet.append(["Date", "Type", "Reference", "Designation", "Detail", "Beneficiaire", "Mode", "Montant"])
    for row in rows:
        sheet.append(
            [
                row["date"].strftime("%d/%m/%Y") if row["date"] else "",
                row["type"],
                row["reference"],
                row["designation"],
                row["detail"],
                row["beneficiaire"],
                row["mode"],
                float(row["montant"] or Decimal("0")),
            ]
        )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="rapport_caisse.xlsx"'
    workbook.save(response)
    return response


def export_rapport_maintenances_xls(request):
    try:
        from openpyxl import Workbook
    except ImportError:
        return HttpResponse(
            "Le module openpyxl n'est pas installe sur cet environnement Python.",
            status=503,
            content_type="text/plain; charset=utf-8",
        )

    queryset = _maintenance_queryset()
    q = (request.GET.get("q") or "").strip()
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()
    statut = (request.GET.get("statut") or "").strip()

    if q:
        queryset = queryset.filter(
            Q(reference__icontains=q)
            | Q(camion__code_camion__icontains=q)
            | Q(camion__numero_tracteur__icontains=q)
            | Q(camion__numero_citerne__icontains=q)
            | Q(camion__chauffeur__nom__icontains=q)
            | Q(fournisseur__nom_fournisseur__icontains=q)
            | Q(prestataire__icontains=q)
            | Q(numero_facture__icontains=q)
        ).distinct()
    if date_from:
        queryset = queryset.filter(date_paiement__gte=date_from)
    if date_to:
        queryset = queryset.filter(date_paiement__lte=date_to)
    if statut:
        queryset = queryset.filter(statut=statut)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rapport maintenance"
    sheet.append(
        [
            "Reference",
            "Code camion",
            "Numero tracteur",
            "Numero citerne",
            "Chauffeur",
            "Date entree garage",
            "Date sortie",
            "Kilometrage entree",
            "Kilometrage sortie",
            "Diagnostics",
            "Fournisseur",
            "Prestataire",
            "Numero facture",
            "Total facture",
            "Statut",
            "Validation logistique",
            "Validation DGA",
            "Validation DG",
            "Date paiement",
            "Observation",
        ]
    )

    for maintenance in queryset:
        chauffeur_nom = (
            Chauffeur.objects.filter(camion=maintenance.camion)
            .values_list("nom", flat=True)
            .first()
            or ""
        )
        diagnostics = []
        for ligne in maintenance.lignes.all():
            piece_labels = [piece.libelle for piece in ligne.sous_lignes.all()]
            diagnostic = f"{ligne.type_maintenance.libelle}: {ligne.libelle}"
            if piece_labels:
                diagnostic += f" ({', '.join(piece_labels)})"
            diagnostics.append(diagnostic)
        sheet.append(
            [
                maintenance.reference,
                maintenance.camion.code_camion,
                maintenance.camion.numero_tracteur,
                maintenance.camion.numero_citerne or "",
                chauffeur_nom,
                maintenance.date_debut.strftime("%d/%m/%Y %H:%M") if maintenance.date_debut else "",
                maintenance.date_fin.strftime("%d/%m/%Y %H:%M") if maintenance.date_fin else "",
                maintenance.kilometrage_entree or "",
                maintenance.kilometrage_sortie or "",
                " | ".join(diagnostics),
                str(maintenance.fournisseur or ""),
                maintenance.prestataire or "",
                maintenance.numero_facture or "",
                str(maintenance.total_facture or ""),
                maintenance.get_statut_display(),
                maintenance.validation_logistique_at.strftime("%d/%m/%Y %H:%M") if maintenance.validation_logistique_at else "",
                maintenance.validation_dga_at.strftime("%d/%m/%Y %H:%M") if maintenance.validation_dga_at else "",
                maintenance.validation_dg_at.strftime("%d/%m/%Y %H:%M") if maintenance.validation_dg_at else "",
                maintenance.date_paiement.strftime("%d/%m/%Y") if maintenance.date_paiement else "",
                maintenance.observation or "",
            ]
        )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="rapport_maintenance_complet.xlsx"'
    workbook.save(response)
    return response


def fournisseurs_maintenance(request):
    query = (request.GET.get("q") or "").strip()
    portefeuille = _fournisseur_portefeuille_for_role(request.user)
    fournisseurs = _fournisseurs_queryset_for_user(request.user)
    if query:
        fournisseurs = fournisseurs.filter(
            Q(nom_fournisseur__icontains=query)
            | Q(entreprise__icontains=query)
            | Q(domaine_activite__icontains=query)
            | Q(numero_telephone__icontains=query)
        )
    form = FournisseurForm(portefeuille=portefeuille, entite_reference=_fournisseur_entite_for_user(request.user, portefeuille))
    return render(
        request,
        "maintenance/fournisseurs.html",
        {
            "fournisseurs": fournisseurs,
            "query": query,
            "form": form,
            "is_admin_maintenance": is_admin_user(request.user),
            "fournisseur_portefeuille": portefeuille,
            **_fournisseurs_return_context(request, portefeuille),
            **_maintenance_tabs_context("fournisseurs"),
        },
    )


def ajouter_fournisseur(request):
    if request.method != "POST":
        return redirect("fournisseurs_maintenance")

    portefeuille = _fournisseur_portefeuille_for_role(request.user)
    form = FournisseurForm(
        request.POST,
        portefeuille=portefeuille,
        entite_reference=_fournisseur_entite_for_user(request.user, portefeuille),
    )
    if form.is_valid():
        fournisseur = form.save()
        journaliser_action(
            request.user,
            "Maintenance",
            "Creation fournisseur",
            str(fournisseur),
            f"{request.user.username} a cree le fournisseur {fournisseur}.",
        )
        messages.success(request, f"Le fournisseur {fournisseur} a ete cree.")
        next_url = (request.POST.get("next") or "").strip()
        target_url = "/maintenance/fournisseurs/"
        if next_url:
            target_url = f"{target_url}?{urlencode({'next': next_url})}"
        return redirect(target_url)

    fournisseurs = _fournisseurs_queryset_for_user(request.user)
    messages.error(request, "Impossible de creer le fournisseur. Verifiez les champs puis reessayez.")
    return render(
        request,
        "maintenance/fournisseurs.html",
        {
            "fournisseurs": fournisseurs,
            "query": "",
            "form": form,
            "is_admin_maintenance": is_admin_user(request.user),
            "fournisseur_portefeuille": portefeuille,
            **_fournisseurs_return_context(request, portefeuille),
            **_maintenance_tabs_context("fournisseurs"),
        },
        status=400,
    )


def modifier_fournisseur(request, id):
    portefeuille = _fournisseur_portefeuille_for_role(request.user)
    fournisseur = get_object_or_404(_fournisseurs_queryset_for_user(request.user), pk=id)
    next_url = (request.GET.get("next") or request.POST.get("next") or "").strip()
    if request.method == "POST":
        form = FournisseurForm(
            request.POST,
            instance=fournisseur,
            portefeuille=portefeuille,
            entite_reference=_fournisseur_entite_for_user(request.user, portefeuille),
        )
        if form.is_valid():
            fournisseur = form.save()
            journaliser_action(
                request.user,
                "Maintenance",
                "Mise a jour fournisseur",
                str(fournisseur),
                f"{request.user.username} a modifie le fournisseur {fournisseur}.",
            )
            messages.success(request, f"Le fournisseur {fournisseur} a ete mis a jour.")
            target_url = "/maintenance/fournisseurs/"
            if next_url:
                target_url = f"{target_url}?{urlencode({'next': next_url})}"
            return redirect(target_url)
        messages.error(request, "Impossible de mettre a jour le fournisseur. Verifiez les champs.")
    else:
        form = FournisseurForm(
            instance=fournisseur,
            portefeuille=portefeuille,
            entite_reference=_fournisseur_entite_for_user(request.user, portefeuille),
        )

    return render(
        request,
        "maintenance/modifier_fournisseur.html",
        {
            "form": form,
            "fournisseur": fournisseur,
            "is_admin_maintenance": is_admin_user(request.user),
            "fournisseur_portefeuille": portefeuille,
            "next_url": next_url,
            "return_url": next_url or "/maintenance/fournisseurs/",
            **_maintenance_tabs_context("fournisseurs"),
        },
        status=400 if request.method == "POST" and form.errors else 200,
    )


@require_POST
def supprimer_fournisseur(request, id):
    if not is_admin_user(request.user):
        messages.error(request, "Seul l'administrateur peut supprimer un fournisseur.")
        return redirect("fournisseurs_maintenance")

    fournisseur = get_object_or_404(_fournisseurs_queryset_for_user(request.user), pk=id)
    label = str(fournisseur)
    fournisseur.delete()
    journaliser_action(
        request.user,
        "Maintenance",
        "Suppression fournisseur",
        label,
        f"{request.user.username} a supprime le fournisseur {label}.",
    )
    messages.success(request, f"Le fournisseur {label} a ete supprime.")
    return redirect("fournisseurs_maintenance")


def types_maintenance(request):
    return_url = request.GET.get("next") or "/maintenance/garage/"
    query = (request.GET.get("q") or "").strip()
    types_qs = TypeMaintenance.objects.prefetch_related("pannes_catalogue")
    if query:
        types_qs = types_qs.filter(libelle__icontains=query)
    return render(
        request,
        "maintenance/types_maintenance.html",
        {
            "types_maintenance": types_qs.order_by("libelle"),
            "query": query,
            "return_url": return_url,
            "form": TypeMaintenanceForm(),
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("garage"),
        },
    )


def _panne_price_columns(panne):
    prices = list(panne.prix_fournisseurs.all()[:3])
    columns = []
    for price in prices:
        columns.append(
            {
                "id": price.id,
                "fournisseur": str(price.fournisseur),
                "montant": _format_amount(price.montant),
                "date": price.date_reference,
                "observation": price.observation,
            }
        )
    while len(columns) < 3:
        columns.append(None)
    return columns


def _panne_usage_rows(panne):
    rows = []
    usages = (
        MaintenanceSousLigne.objects.select_related(
            "maintenance_ligne__maintenance",
            "maintenance_ligne__maintenance__camion",
        )
        .prefetch_related("maintenance_ligne__maintenance__factures_achat__fournisseur")
        .filter(panne_catalogue=panne)
        .order_by("-maintenance_ligne__maintenance__date_debut", "-id")
    )
    for index, usage in enumerate(usages, start=1):
        maintenance = usage.maintenance_ligne.maintenance
        camion = maintenance.camion
        factures = list(maintenance.factures_achat.all())
        facture = factures[0] if factures else None
        fournisseur = facture.fournisseur if facture else maintenance.fournisseur
        rows.append(
            {
                "index": index,
                "date": _format_step_date(maintenance.date_debut),
                "camion": camion.numero_tracteur if camion else "-",
                "citerne": camion.numero_citerne if camion and camion.numero_citerne else "-",
                "chauffeur": _chauffeur_for_camion(camion),
                "prix": _format_amount(usage.prix_unitaire),
                "fournisseur": str(fournisseur) if fournisseur else "-",
                "reference": maintenance.reference,
            }
        )
    return rows


def _chauffeur_for_camion(camion):
    if not camion:
        return "-"
    chauffeur = Chauffeur.objects.filter(camion=camion).order_by("nom").first()
    return chauffeur.nom if chauffeur else "-"


def gerer_pannes(request):
    query = (request.GET.get("q") or "").strip()
    type_filter = (request.GET.get("type") or "").strip()
    pannes_qs = (
        PanneCatalogue.objects.select_related("type_maintenance")
        .prefetch_related("prix_fournisseurs__fournisseur")
        .annotate(utilisations=Count("sous_lignes_maintenance", distinct=True))
    )
    if query:
        pannes_qs = pannes_qs.filter(
            Q(libelle__icontains=query)
            | Q(type_maintenance__libelle__icontains=query)
            | Q(prix_fournisseurs__fournisseur__nom_fournisseur__icontains=query)
            | Q(prix_fournisseurs__fournisseur__entreprise__icontains=query)
        ).distinct()
    if type_filter:
        pannes_qs = pannes_qs.filter(type_maintenance_id=type_filter)
    pannes = list(pannes_qs.order_by("type_maintenance__libelle", "libelle"))
    for panne in pannes:
        panne.price_columns = _panne_price_columns(panne)
        panne.usage_rows = _panne_usage_rows(panne)

    return render(
        request,
        "maintenance/gerer_pannes.html",
        {
            "pannes": pannes,
            "query": query,
            "type_filter": type_filter,
            "type_choices": TypeMaintenance.objects.order_by("libelle"),
            "panne_form": PanneCatalogueForm(),
            "prix_form": PanneFournisseurPrixForm(),
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("garage"),
        },
    )


@require_POST
def ajouter_panne_catalogue(request):
    form = PanneCatalogueForm(request.POST)
    if form.is_valid():
        panne = form.save()
        journaliser_action(
            request.user,
            "Maintenance",
            "Creation panne",
            str(panne),
            f"{request.user.username} a cree la panne {panne}.",
        )
        messages.success(request, "La panne a ete ajoutee au catalogue.")
    else:
        messages.error(request, "Impossible d'ajouter cette panne. Verifiez le type et le libelle.")
    return redirect("gerer_pannes")


@require_POST
def modifier_panne_catalogue(request, id):
    panne = get_object_or_404(PanneCatalogue, pk=id)
    form = PanneCatalogueForm(request.POST, instance=panne)
    if form.is_valid():
        panne = form.save()
        journaliser_action(
            request.user,
            "Maintenance",
            "Modification panne",
            str(panne),
            f"{request.user.username} a modifie la panne {panne}.",
        )
        messages.success(request, "La panne a ete mise a jour.")
    else:
        messages.error(request, "Impossible de modifier cette panne.")
    return redirect("gerer_pannes")


@require_POST
def supprimer_panne_catalogue(request, id):
    panne = get_object_or_404(PanneCatalogue, pk=id)
    label = str(panne)
    if panne.sous_lignes_maintenance.exists():
        messages.error(request, "Cette panne est deja utilisee dans une maintenance. Elle ne peut pas etre supprimee sans perdre l'historique.")
        return redirect("gerer_pannes")
    panne.delete()
    journaliser_action(
        request.user,
        "Maintenance",
        "Suppression panne",
        label,
        f"{request.user.username} a supprime la panne {label}.",
    )
    messages.success(request, "La panne a ete supprimee.")
    return redirect("gerer_pannes")


@require_POST
def ajouter_prix_panne(request, id):
    panne = get_object_or_404(PanneCatalogue, pk=id)
    form = PanneFournisseurPrixForm(request.POST)
    if form.is_valid():
        price = form.save(commit=False)
        PanneFournisseurPrix.objects.update_or_create(
            panne=panne,
            fournisseur=price.fournisseur,
            defaults={
                "montant": price.montant,
                "date_reference": price.date_reference,
                "observation": price.observation,
            },
        )
        journaliser_action(
            request.user,
            "Maintenance",
            "Prix fournisseur panne",
            str(panne),
            f"{request.user.username} a enregistre un prix fournisseur pour {panne}.",
        )
        messages.success(request, "Le prix fournisseur a ete memorise.")
    else:
        messages.error(request, "Impossible d'enregistrer ce prix fournisseur.")
    return redirect("gerer_pannes")


@require_POST
def supprimer_prix_panne(request, id):
    price = get_object_or_404(PanneFournisseurPrix.objects.select_related("panne", "fournisseur"), pk=id)
    label = str(price)
    price.delete()
    journaliser_action(
        request.user,
        "Maintenance",
        "Suppression prix fournisseur panne",
        label,
        f"{request.user.username} a supprime le prix {label}.",
    )
    messages.success(request, "Le prix fournisseur a ete supprime.")
    return redirect("gerer_pannes")


def ajouter_type_maintenance(request):
    if request.method != "POST":
        return redirect("types_maintenance")

    return_url = request.POST.get("next") or "/maintenance/garage/"
    form = TypeMaintenanceForm(request.POST)
    if form.is_valid():
        type_maintenance = form.save()
        journaliser_action(
            request.user,
            "Maintenance",
            "Creation type de maintenance",
            str(type_maintenance),
            f"{request.user.username} a cree le type de maintenance {type_maintenance}.",
        )
        messages.success(request, f"Le type {type_maintenance.libelle} a ete cree.")
        return redirect(f"/maintenance/types-maintenance/?next={return_url}")

    messages.error(request, "Impossible de creer ce type de maintenance. Verifiez le libelle.")
    return render(
        request,
        "maintenance/types_maintenance.html",
        {
            "types_maintenance": TypeMaintenance.objects.all().order_by("libelle"),
            "query": "",
            "return_url": return_url,
            "form": form,
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("garage"),
        },
        status=400,
    )


def modifier_type_maintenance(request, id):
    type_maintenance = get_object_or_404(TypeMaintenance, pk=id)
    return_url = request.GET.get("next") or request.POST.get("next") or "/maintenance/garage/"
    if request.method == "POST":
        form = TypeMaintenanceForm(request.POST, instance=type_maintenance)
        if form.is_valid():
            type_maintenance = form.save()
            journaliser_action(
                request.user,
                "Maintenance",
                "Mise a jour type de maintenance",
                str(type_maintenance),
                f"{request.user.username} a modifie le type de maintenance {type_maintenance}.",
            )
            messages.success(request, f"Le type {type_maintenance.libelle} a ete mis a jour.")
            return redirect(f"/maintenance/types-maintenance/?next={return_url}")
        messages.error(request, "Impossible de modifier ce type de maintenance.")
    else:
        form = TypeMaintenanceForm(instance=type_maintenance)

    return render(
        request,
        "maintenance/modifier_type_maintenance.html",
        {
            "form": form,
            "type_maintenance": type_maintenance,
            "return_url": return_url,
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("garage"),
        },
        status=400 if request.method == "POST" and form.errors else 200,
    )


@require_POST
def supprimer_type_maintenance(request, id):
    type_maintenance = get_object_or_404(TypeMaintenance, pk=id)
    return_url = request.POST.get("next") or "/maintenance/garage/"
    label = type_maintenance.libelle
    type_maintenance.actif = not type_maintenance.actif
    type_maintenance.save(update_fields=["actif"])
    journaliser_action(
        request.user,
        "Maintenance",
        "Archivage type de maintenance" if not type_maintenance.actif else "Reactivation type de maintenance",
        label,
        (
            f"{request.user.username} a archive le type de maintenance {label}."
            if not type_maintenance.actif
            else f"{request.user.username} a reactive le type de maintenance {label}."
        ),
    )
    messages.success(
        request,
        (
            f"Le type {label} a ete archive. Il ne sera plus propose dans les nouvelles fiches."
            if not type_maintenance.actif
            else f"Le type {label} a ete reactive."
        ),
    )
    return redirect(f"/maintenance/types-maintenance/?next={return_url}")


def _render_garage_form(request, template_name, form, formset, **context):
    _attach_subline_values(formset, request=request)
    pannes_catalog = _build_pannes_catalog()
    return render(
        request,
        template_name,
        {
            "form": form,
            "formset": formset,
            "type_form": TypeMaintenanceForm(),
            "panne_form": PanneCatalogueForm(),
            "camions_catalog": _garage_camions_catalog(),
            "articles_stock_catalog": ArticleStock.objects.order_by("libelle"),
            "pannes_catalog_json": mark_safe(json.dumps(pannes_catalog)),
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("garage"),
            **context,
        },
    )


def _render_achat_form(request, template_name, form, **context):
    fournisseurs_catalog = [
        {"id": fournisseur.id, "label": str(fournisseur)}
        for fournisseur in Fournisseur.objects.filter(portefeuille=Fournisseur.PORTEFEUILLE_LOGISTIQUE).order_by("nom_fournisseur", "entreprise")
    ]
    prestataires_catalog = [
        {"label": str(prestataire)}
        for prestataire in Prestataire.objects.all()
    ]
    if context["maintenance"].prestataire and not any(
        item["label"] == context["maintenance"].prestataire for item in prestataires_catalog
    ):
        prestataires_catalog.insert(0, {"label": context["maintenance"].prestataire})
    return render(
        request,
        template_name,
        {
            "form": form,
            "invoice_formset": context["invoice_formset"],
            "fournisseur_form": FournisseurForm(),
            "prestataire_form": PrestataireForm(),
            "piece_rows": _attach_achat_piece_rows(context["maintenance"]),
            "invoice_rows": _maintenance_invoice_rows(context["maintenance"]),
            "fournisseurs_catalog": fournisseurs_catalog,
            "prestataires_catalog": prestataires_catalog,
            "is_admin_maintenance": is_admin_user(request.user),
            "can_edit_achat": context.get("can_edit_achat", False),
            **_maintenance_tabs_context("achat"),
            **context,
        },
    )


def _render_paiement_form(request, template_name, form, **context):
    preview_context = _build_validation_preview_context(context["maintenance"])
    return render(
        request,
        template_name,
        {
            "form": form,
            "invoice_rows": _maintenance_invoice_rows(context["maintenance"]),
            "is_admin_maintenance": is_admin_user(request.user),
            **_maintenance_tabs_context("paiements"),
            **preview_context,
            **context,
        },
    )


def _build_validation_preview_context(maintenance):
    chauffeur = (
        Chauffeur.objects.filter(camion=maintenance.camion)
        .values_list("nom", flat=True)
        .first()
    )
    validation_steps = [
        {
            "label": "Validation logistique",
            "status": (
                "Non requise pour stock interne"
                if maintenance.is_stock_only() and not maintenance.validation_logistique_at
                else "Valide par la logistique"
                if maintenance.validation_logistique_at
                else "En attente de validation logistique"
            ),
            "variant": (
                "neutral"
                if maintenance.is_stock_only() and not maintenance.validation_logistique_at
                else "ok"
                if maintenance.validation_logistique_at
                else "pending"
            ),
            "by": maintenance.validation_logistique_by.get_full_name() or maintenance.validation_logistique_by.username
            if maintenance.validation_logistique_by
            else "-",
            "at": maintenance.validation_logistique_at,
            "at_display": _format_step_date(maintenance.validation_logistique_at),
        },
        {
            "label": "Validation DGA",
            "status": (
                "Rejete par le DGA"
                if maintenance.statut == "rejetee_dga"
                else "Valide par le DGA"
                if maintenance.validation_dga_at
                else "En attente de validation DGA"
            ),
            "variant": (
                "danger"
                if maintenance.statut == "rejetee_dga"
                else "ok"
                if maintenance.validation_dga_at
                else "pending"
            ),
            "by": maintenance.validation_dga_by.get_full_name() or maintenance.validation_dga_by.username
            if maintenance.validation_dga_by
            else "-",
            "at": maintenance.validation_dga_at,
            "at_display": _format_step_date(maintenance.validation_dga_at),
        },
        {
            "label": "Validation DG",
            "status": (
                "Rejete par le DG"
                if maintenance.statut == "rejetee_dg"
                else "Valide par le DG"
                if maintenance.validation_dg_at
                else "En attente de validation DG"
            ),
            "variant": (
                "danger"
                if maintenance.statut == "rejetee_dg"
                else "ok"
                if maintenance.validation_dg_at
                else "pending"
            ),
            "by": maintenance.validation_dg_by.get_full_name() or maintenance.validation_dg_by.username
            if maintenance.validation_dg_by
            else "-",
            "at": maintenance.validation_dg_at,
            "at_display": _format_step_date(maintenance.validation_dg_at),
        },
        {
            "label": "Paiement",
            "status": (
                "Non applicable (stock interne)"
                if maintenance.statut == "validee_stock"
                else
                "Paiement effectue"
                if maintenance.statut == "payee"
                else "En attente de paiement"
                if maintenance.statut == "attente_paiement"
                else "Non disponible dans le circuit actuel"
            ),
            "variant": (
                "neutral"
                if maintenance.statut == "validee_stock"
                else
                "ok"
                if maintenance.statut == "payee"
                else "pending"
                if maintenance.statut == "attente_paiement"
                else "neutral"
            ),
            "by": "-",
            "at": maintenance.date_paiement,
            "at_display": _format_step_date(maintenance.date_paiement),
        },
    ]
    return {
        "maintenance": maintenance,
        "piece_rows": _attach_achat_piece_rows(maintenance),
        "invoice_rows": _maintenance_invoice_rows(maintenance),
        "montant_en_lettres": _amount_to_words(maintenance.total_facture),
        "montant_total_formatte": _format_amount(maintenance.total_facture),
        "montant_stock_information_formatte": _format_amount(maintenance.total_information_stock),
        "montant_global_information_formatte": _format_amount(maintenance.total_global_information),
        "chauffeur_nom": chauffeur or "-",
        "validation_steps": validation_steps,
    }


def _set_form_read_only(form):
    for field in form.fields.values():
        field.disabled = True


def _set_formset_read_only(formset):
    for form in formset.forms:
        for field in form.fields.values():
            field.disabled = True


def ajouter_maintenance_garage(request):
    if request.method == "POST":
        form = MaintenanceGarageForm(request.POST)
        formset = MaintenanceGarageLigneFormSet(request.POST)
        if form.is_valid() and formset.is_valid():
            with transaction.atomic():
                maintenance = form.save(commit=False)
                maintenance.statut = "attente_prix"
                maintenance.save()
                formset.instance = maintenance
                formset.save()
                _save_subline_items(request, formset)
                maintenance.refresh_total_facture()
                maintenance.statut = "attente_dga" if maintenance.is_stock_only() else "attente_prix"
                maintenance.save(update_fields=["statut"])
                maintenance_label = maintenance.reference or f"Diagnostic #{maintenance.id}"
                journaliser_action(
                    request.user,
                    "Maintenance",
                    "Creation de diagnostic",
                    maintenance_label,
                    f"{request.user.username} a cree le diagnostic {maintenance_label} pour le camion {maintenance.camion}.",
                )
            messages.success(
                request,
                (
                    f"Le diagnostic {maintenance.reference} a ete enregistre et transmis a la validation DGA."
                    if maintenance.is_stock_only()
                    else f"Le diagnostic {maintenance.reference} a ete enregistre et transmis a la saisie des prix."
                ),
            )
            return redirect("garage_maintenances")
    else:
        form = MaintenanceGarageForm(initial={"statut": "en_cours"})
        formset = MaintenanceGarageLigneFormSet()

    return _render_garage_form(
        request,
        "maintenance/ajouter_maintenance_garage.html",
        form,
        formset,
        submit_label="Enregistrer le diagnostic",
    )


def modifier_maintenance_garage(request, id):
    maintenance = get_object_or_404(Maintenance, id=id)
    if maintenance.statut in {"payee", "validee_stock"}:
        messages.info(request, "Cette fiche finalisee est disponible dans l'historique garage.")
        return redirect("/maintenance/garage/?scope=historique")
    user_role = get_user_role(request.user)
    is_admin = is_admin_user(request.user)
    can_edit_diagnostic = is_admin or (user_role == "maintenancier" and maintenance.statut in {"en_cours", "attente_prix"})
    if request.method == "POST":
        if not can_edit_diagnostic:
            messages.error(request, "Seuls le maintenancier et l'administrateur peuvent modifier le diagnostic.")
            return redirect("modifier_maintenance_garage", id=maintenance.id)
        form = MaintenanceGarageForm(request.POST, instance=maintenance)
        formset = MaintenanceGarageLigneFormSet(request.POST, instance=maintenance)
        if form.is_valid() and formset.is_valid():
            if not is_admin and any(
                ligne_form.cleaned_data.get("DELETE")
                for ligne_form in formset.forms
                if hasattr(ligne_form, "cleaned_data")
            ):
                messages.error(request, "Seul l'administrateur peut supprimer une ligne de panne.")
                return redirect("modifier_maintenance_garage", id=maintenance.id)
            with transaction.atomic():
                maintenance = form.save()
                formset.instance = maintenance
                formset.save()
                _save_subline_items(
                    request,
                    formset,
                    allow_create=can_edit_diagnostic,
                    allow_delete=is_admin,
                )
                maintenance.refresh_total_facture()
                maintenance.statut = "attente_dga" if maintenance.is_stock_only() else "attente_prix"
                maintenance.save(update_fields=["statut"])
                maintenance_label = maintenance.reference or f"Diagnostic #{maintenance.id}"
                journaliser_action(
                    request.user,
                    "Maintenance",
                    "Modification de diagnostic",
                    maintenance_label,
                    f"{request.user.username} a modifie le diagnostic {maintenance_label}.",
                )
            messages.success(
                request,
                (
                    f"Le diagnostic {maintenance.reference} a ete mis a jour et transmis a la validation DGA."
                    if maintenance.is_stock_only()
                    else f"Le diagnostic {maintenance.reference} a ete mis a jour et transmis a la saisie des prix."
                ),
            )
            return redirect("garage_maintenances")
    else:
        form = MaintenanceGarageForm(instance=maintenance)
        formset = MaintenanceGarageLigneFormSet(instance=maintenance)
        if not can_edit_diagnostic:
            _set_form_read_only(form)
            _set_formset_read_only(formset)

    return _render_garage_form(
        request,
        "maintenance/modifier_maintenance_garage.html",
        form,
        formset,
        maintenance=maintenance,
        read_only=not can_edit_diagnostic,
        allow_structure_changes=can_edit_diagnostic,
        can_delete_lines=is_admin,
        submit_label="Mettre a jour le diagnostic",
    )


def modifier_maintenance_achat(request, id):
    maintenance = get_object_or_404(Maintenance, id=id)
    if maintenance.statut in {"payee", "validee_stock"}:
        messages.info(request, "Cette fiche finalisee est disponible dans l'historique achat.")
        return redirect("/maintenance/achat/?scope=historique")
    user_role = get_user_role(request.user)
    is_admin = is_admin_user(request.user)
    can_edit_achat = is_admin or (
        user_role == "logistique"
        and maintenance.validation_dga_at is None
        and maintenance.statut in {"attente_prix", "attente_dga", "rejetee_dga"}
    )
    if request.method == "POST":
        if not can_edit_achat:
            messages.error(request, "Seuls la logistique et l'administrateur peuvent modifier les achats avant decision finale.")
            return redirect("achat_maintenances")
        post_data = request.POST.copy()
        if not is_admin:
            post_data["statut"] = maintenance.statut
        post_data["prestataire"] = (post_data.get("prestataire_search") or post_data.get("prestataire") or "").strip()
        form = MaintenanceAchatForm(post_data, request.FILES, instance=maintenance)
        invoice_formset = MaintenanceFactureFormSet(
            post_data,
            request.FILES,
            instance=maintenance,
            prefix="factures",
        )
        if form.is_valid() and invoice_formset.is_valid():
            with transaction.atomic():
                maintenance = form.save(commit=False)
                maintenance.statut = maintenance.statut or "attente_prix"
                maintenance.save()
                invoice_formset.instance = maintenance
                invoice_formset.save()
                _sync_maintenance_invoice_snapshot(maintenance)
                _save_achat_piece_prices(request, maintenance)
                _memoize_panne_prices_from_maintenance(maintenance)
                maintenance.refresh_total_facture()
                maintenance.validation_logistique_at = timezone.now()
                maintenance.validation_logistique_by = request.user
                maintenance.statut = "attente_dga"
                maintenance.save(
                    update_fields=[
                        "validation_logistique_at",
                        "validation_logistique_by",
                        "statut",
                    ]
                )
                maintenance_label = maintenance.reference or f"Diagnostic #{maintenance.id}"
                journaliser_action(
                    request.user,
                    "Maintenance",
                    "Valorisation de diagnostic",
                    maintenance_label,
                    f"{request.user.username} a mis a jour l'achat et les prix du diagnostic {maintenance_label}.",
                )
                messages.success(request, "Les prix ont ete enregistres.")
                return redirect("achat_maintenances")
        messages.error(
            request,
            "Impossible d'enregistrer les prix. Verifiez les champs achat puis reessayez.",
        )
    else:
        form = MaintenanceAchatForm(instance=maintenance)
        invoice_formset = MaintenanceFactureFormSet(
            instance=maintenance,
            prefix="factures",
            initial=(
                []
                if maintenance.factures_achat.exists() or not (maintenance.fournisseur_id or maintenance.numero_facture or maintenance.facture_fichier)
                else [{
                    "fournisseur": maintenance.fournisseur_id,
                    "numero_facture": maintenance.numero_facture,
                }]
            ),
        )
        if not can_edit_achat:
            _set_form_read_only(form)
            _set_formset_read_only(invoice_formset)

    return _render_achat_form(
        request,
        "maintenance/modifier_maintenance_achat.html",
        form,
        maintenance=maintenance,
        invoice_formset=invoice_formset,
        read_only=not can_edit_achat,
        can_edit_achat=can_edit_achat,
    )


def terminer_maintenance(request, id):
    if request.method != "POST":
        return redirect("garage_maintenances")

    maintenance = get_object_or_404(Maintenance, id=id)
    if maintenance.statut != "en_cours":
        messages.info(request, "Ce diagnostic est deja transmis.")
        return redirect("garage_maintenances")

    maintenance.statut = "attente_prix"
    maintenance.save(update_fields=["statut"])
    journaliser_action(
        request.user,
        "Maintenance",
        "Transmission diagnostic",
        maintenance.reference,
        f"{request.user.username} a termine et transmis le diagnostic {maintenance.reference}.",
    )
    messages.success(request, "Le diagnostic a ete transmis pour saisie des prix.")
    return redirect("garage_maintenances")


def valider_maintenance_logistique(request, id):
    if request.method != "POST":
        return redirect("garage_maintenances")

    maintenance = get_object_or_404(Maintenance, id=id)
    messages.info(request, "La validation logistique est maintenant automatique apres la saisie des prix.")
    return redirect("garage_maintenances")


def rejeter_maintenance_dga(request, id):
    if request.method != "POST":
        return redirect("garage_maintenances")

    maintenance = get_object_or_404(Maintenance, id=id)
    if get_user_role(request.user) != "dga":
        messages.error(request, "Seul le role DGA peut rejeter cette fiche.")
        return redirect("garage_maintenances")
    if maintenance.statut != "attente_dga":
        messages.error(request, "Cette fiche n'est pas en attente de validation DGA.")
        return redirect("garage_maintenances")

    maintenance.statut = "rejetee_dga"
    if not maintenance.date_fin:
        maintenance.date_fin = timezone.now()
    maintenance.save(update_fields=["statut", "date_fin"])
    journaliser_action(
        request.user,
        "Maintenance",
        "Rejet DGA",
        maintenance.reference,
        f"{request.user.username} a rejete la fiche {maintenance.reference} au niveau DGA.",
    )
    return redirect("garage_maintenances")


def valider_maintenance_dga(request, id):
    if request.method != "POST":
        return redirect("garage_maintenances")

    maintenance = get_object_or_404(Maintenance, id=id)
    if get_user_role(request.user) != "dga":
        messages.error(request, "Seul le role DGA peut faire cette validation.")
        return redirect("garage_maintenances")
    if maintenance.statut != "attente_dga":
        messages.error(request, "Cette fiche n'est pas en attente de validation DGA.")
        return redirect("garage_maintenances")
    if maintenance.is_validated_by_dga() and maintenance.statut != "attente_dga":
        messages.info(request, "Cette fiche est deja validee par le DGA.")
        return redirect("garage_maintenances")

    maintenance.validation_dga_at = timezone.now()
    maintenance.validation_dga_by = request.user
    maintenance.statut = "attente_dg"
    maintenance.save(update_fields=["validation_dga_at", "validation_dga_by", "statut"])
    journaliser_action(
        request.user,
        "Maintenance",
        "Validation DGA",
        maintenance.reference,
        f"{request.user.username} a valide la fiche {maintenance.reference} au niveau DGA.",
    )
    return redirect("garage_maintenances")


def rejeter_maintenance_dg(request, id):
    if request.method != "POST":
        return redirect("garage_maintenances")

    maintenance = get_object_or_404(Maintenance, id=id)
    if get_user_role(request.user) != "directeur":
        messages.error(request, "Seul le role DG peut rejeter cette fiche.")
        return redirect("garage_maintenances")
    if maintenance.statut != "attente_dg":
        messages.error(request, "Cette fiche n'est pas en attente de validation DG.")
        return redirect("garage_maintenances")

    maintenance.statut = "rejetee_dg"
    if not maintenance.date_fin:
        maintenance.date_fin = timezone.now()
    maintenance.save(update_fields=["statut", "date_fin"])
    journaliser_action(
        request.user,
        "Maintenance",
        "Rejet DG",
        maintenance.reference,
        f"{request.user.username} a rejete la fiche {maintenance.reference} au niveau DG.",
    )
    return redirect("garage_maintenances")


def valider_maintenance_dg(request, id):
    maintenance = get_object_or_404(Maintenance, id=id)
    if get_user_role(request.user) != "directeur":
        messages.error(request, "Seul le role DG peut faire cette validation.")
        return redirect("garage_maintenances")
    if maintenance.statut != "attente_dg":
        messages.error(request, "Cette fiche n'est pas en attente de validation DG.")
        return redirect("garage_maintenances")
    if maintenance.is_validated_by_dg() and maintenance.statut != "attente_dg":
        messages.info(request, "Cette fiche est deja validee par le DG.")
        return redirect("garage_maintenances")
    if request.method != "POST":
        return redirect("apercu_validation_maintenance", id=maintenance.id)

    decision_form = MaintenanceDecisionPaiementForm(request.POST)
    chosen_mode = ""
    if not maintenance.is_stock_only():
        if not decision_form.is_valid():
            messages.error(request, "Le mode de paiement choisi est invalide.")
            return redirect("apercu_validation_maintenance", id=maintenance.id)
        chosen_mode = decision_form.cleaned_data.get("mode_paiement") or ""
        if chosen_mode not in {Maintenance.MODE_CHEQUE, Maintenance.MODE_ESPECE}:
            messages.error(request, "Le DG doit choisir cheque ou espece.")
            return redirect("apercu_validation_maintenance", id=maintenance.id)

    try:
        with transaction.atomic():
            _issue_stock_for_maintenance(maintenance, request.user)
            maintenance.validation_dg_at = timezone.now()
            maintenance.validation_dg_by = request.user
            if not maintenance.is_stock_only():
                maintenance.mode_paiement = chosen_mode
            maintenance.statut = "validee_stock" if maintenance.is_stock_only() else "attente_paiement"
            if not maintenance.date_fin:
                maintenance.date_fin = timezone.now()
            update_fields = ["validation_dg_at", "validation_dg_by", "statut", "date_fin"]
            if not maintenance.is_stock_only():
                update_fields.append("mode_paiement")
            maintenance.save(update_fields=update_fields)
    except ValidationError as error:
        messages.error(request, str(error))
        return redirect("garage_maintenances")

    journaliser_action(
        request.user,
        "Maintenance",
        "Validation DG",
        maintenance.reference,
        f"{request.user.username} a valide la fiche {maintenance.reference} au niveau DG{f' en mode {chosen_mode}' if chosen_mode else ''}.",
    )
    return redirect("garage_maintenances")


def imprimer_maintenance(request, id):
    maintenance = get_object_or_404(
        Maintenance.objects.select_related("camion").prefetch_related(
            "lignes__type_maintenance",
            "lignes__sous_lignes",
        ),
        id=id,
    )
    context = _build_validation_preview_context(maintenance)
    chauffeur = Chauffeur.objects.filter(camion=maintenance.camion).first()
    for row in context["piece_rows"]:
        if row["pieces"]:
            for piece in row["pieces"]:
                piece.prix_unitaire_affiche = _format_amount(piece.prix_unitaire)
                piece.montant_affiche = _format_amount(piece.montant)
        else:
            row["ligne"].prix_unitaire_affiche = _format_amount(row["ligne"].prix_unitaire)
            row["ligne"].montant_affiche = _format_amount(row["ligne"].montant)
    return render(
        request,
        "maintenance/imprimer_maintenance.html",
        {
            **context,
            "chauffeur": chauffeur,
        },
    )


def supprimer_maintenance(request, id):
    if not is_admin_user(request.user):
        messages.error(request, "Seul l'administrateur peut supprimer une fiche maintenance.")
        return redirect("garage_maintenances")
    maintenance = get_object_or_404(Maintenance, id=id)
    camion = maintenance.camion
    maintenance_label = maintenance.reference or f"Diagnostic #{maintenance.id}"
    maintenance.delete()

    maintenance_active = camion.maintenances.filter(statut="en_cours").exists()
    if not maintenance_active and camion.etat == "au_garage":
        camion.etat = "disponible"
        camion.save(update_fields=["etat"])

    journaliser_action(
        request.user,
        "Maintenance",
        "Suppression de diagnostic",
        maintenance_label,
        f"{request.user.username} a supprime le diagnostic {maintenance_label}.",
    )
    return redirect("garage_maintenances")


def ajouter_type_maintenance_modal(request):
    if request.method != "POST":
        return JsonResponse(
            {"success": False, "errors": {"__all__": ["Requete invalide."]}},
            status=405,
        )

    form = TypeMaintenanceForm(request.POST)
    if form.is_valid():
        type_maintenance = form.save()
        return JsonResponse(
            {
                "success": True,
                "type_maintenance": {
                    "id": type_maintenance.id,
                    "label": type_maintenance.libelle,
                },
            }
        )

    errors = {
        field: [item["message"] for item in messages]
        for field, messages in form.errors.get_json_data().items()
    }
    return JsonResponse({"success": False, "errors": errors}, status=400)


def ajouter_panne_modal(request):
    if request.method != "POST":
        return JsonResponse(
            {"success": False, "errors": {"__all__": ["Requete invalide."]}},
            status=405,
        )

    form = PanneCatalogueForm(request.POST)
    if form.is_valid():
        panne = form.save()
        return JsonResponse(
            {
                "success": True,
                "panne": {
                    "id": panne.id,
                    "label": panne.libelle,
                    "type_maintenance_id": panne.type_maintenance_id,
                },
            }
        )

    errors = {
        field: [item["message"] for item in messages]
        for field, messages in form.errors.get_json_data().items()
    }
    return JsonResponse({"success": False, "errors": errors}, status=400)


def ajouter_fournisseur_modal(request):
    if request.method != "POST":
        return JsonResponse(
            {"success": False, "errors": {"__all__": ["Requete invalide."]}},
            status=405,
        )

    portefeuille = (request.POST.get("portefeuille") or "").strip() or _fournisseur_portefeuille_for_role(request.user)
    form = FournisseurForm(
        request.POST,
        portefeuille=portefeuille,
        entite_reference=_fournisseur_entite_for_user(request.user, portefeuille),
    )
    if form.is_valid():
        fournisseur = form.save()
        return JsonResponse(
            {
                "success": True,
                "fournisseur": {
                    "id": fournisseur.id,
                    "label": str(fournisseur),
                },
            }
        )

    errors = {
        field: [item["message"] for item in messages]
        for field, messages in form.errors.get_json_data().items()
    }
    return JsonResponse({"success": False, "errors": errors}, status=400)


def ajouter_prestataire_modal(request):
    if request.method != "POST":
        return JsonResponse(
            {"success": False, "errors": {"__all__": ["Requete invalide."]}},
            status=405,
        )

    form = PrestataireForm(request.POST)
    if form.is_valid():
        prestataire = form.save()
        return JsonResponse(
            {
                "success": True,
                "prestataire": {
                    "label": str(prestataire),
                },
            }
        )

    errors = {
        field: [item["message"] for item in messages]
        for field, messages in form.errors.get_json_data().items()
    }
    return JsonResponse({"success": False, "errors": errors}, status=400)


def modifier_maintenance_paiement(request, id):
    maintenance = get_object_or_404(Maintenance, id=id)
    can_edit_paiement = _maintenance_payment_allowed(request.user, maintenance)
    payment_role = get_user_role(request.user)
    caisse_metrics = _get_caisse_metrics(request.user) if maintenance.mode_paiement == Maintenance.MODE_ESPECE and payment_role == "caissiere" else None
    montant_a_payer = maintenance.total_facture or Decimal("0")
    solde_courant = caisse_metrics["solde"] if caisse_metrics else None
    if request.method == "POST":
        if not can_edit_paiement or maintenance.statut != "attente_paiement":
            messages.error(request, "Cette fiche n'est pas disponible pour le paiement.")
            return redirect("paiements_maintenances")
        form = MaintenancePaiementForm(request.POST, instance=maintenance)
        if form.is_valid():
            if caisse_metrics and solde_courant is not None and montant_a_payer > solde_courant:
                form.add_error(
                    None,
                    f"Paiement impossible : le montant de { _format_amount(montant_a_payer) } GNF depasse le solde disponible de { _format_amount(solde_courant) } GNF.",
                )
            else:
                maintenance = form.save(commit=False)
                maintenance.statut = "payee"
                maintenance.paiement_saisi_par = request.user
                maintenance.paiement_saisi_le = timezone.now()
                maintenance.save()
                journaliser_action(
                    request.user,
                    "Maintenance",
                    "Paiement maintenance",
                    maintenance.reference,
                    f"{request.user.username} a enregistre le paiement de la fiche {maintenance.reference}.",
                )
                messages.success(request, f"Le paiement de la fiche {maintenance.reference} a ete enregistre.")
                return redirect("paiements_maintenances")
        messages.error(request, "Impossible d'enregistrer le paiement. Verifiez les champs.")
    else:
        form = MaintenancePaiementForm(instance=maintenance)
        if not can_edit_paiement or maintenance.statut != "attente_paiement":
            _set_form_read_only(form)

    return _render_paiement_form(
        request,
        "maintenance/modifier_maintenance_paiement.html",
        form,
        maintenance=maintenance,
        can_edit_paiement=can_edit_paiement and maintenance.statut == "attente_paiement",
        historique=maintenance.statut == "payee",
        banques=Banque.objects.filter(actif=True).order_by("nom"),
        banque_form=BanqueForm(),
        type_piece_form=TypePieceIdentiteForm(),
        type_pieces=TypePieceIdentite.objects.order_by("libelle"),
        is_cheque_payment=maintenance.mode_paiement == Maintenance.MODE_CHEQUE,
        is_espece_payment=maintenance.mode_paiement == Maintenance.MODE_ESPECE,
        solde_courant_display=_format_amount(solde_courant) if solde_courant is not None else None,
        montant_a_payer_display=_format_amount(montant_a_payer),
        solde_apres_display=_format_amount((solde_courant - montant_a_payer) if solde_courant is not None else Decimal("0")) if solde_courant is not None else None,
        caisse_insuffisante=bool(solde_courant is not None and montant_a_payer > solde_courant),
    )


def apercu_validation_maintenance(request, id):
    maintenance = get_object_or_404(_maintenance_queryset(), id=id)
    user_role = get_user_role(request.user)
    if user_role == "dga" and maintenance.statut != "attente_dga":
        messages.error(request, "Cette fiche n'est pas en attente de validation DGA.")
        return redirect("garage_maintenances")
    if user_role == "directeur" and maintenance.statut != "attente_dg":
        messages.error(request, "Cette fiche n'est pas en attente de validation DG.")
        return redirect("garage_maintenances")
    return render(
        request,
        "maintenance/apercu_validation_maintenance.html",
        {
            **_build_validation_preview_context(maintenance),
            "user_role": user_role,
            "decision_form": (
                MaintenanceDecisionPaiementForm(initial={"mode_paiement": maintenance.mode_paiement})
                if user_role == "directeur" and maintenance.statut == "attente_dg" and not maintenance.is_stock_only()
                else None
            ),
            **_maintenance_tabs_context("garage"),
        },
    )


def export_garage_xls(request):
    historique = request.GET.get("scope") == "historique"
    queryset = _maintenance_queryset()
    queryset = queryset.filter(statut="payee") if historique else queryset.exclude(statut="payee")
    queryset, _ = _apply_maintenance_filters(request, queryset)
    return _export_maintenance_xls(queryset, "maintenance_garage")


def export_garage_pdf(request):
    historique = request.GET.get("scope") == "historique"
    queryset = _maintenance_queryset()
    queryset = queryset.filter(statut="payee") if historique else queryset.exclude(statut="payee")
    queryset, _ = _apply_maintenance_filters(request, queryset)
    return _export_maintenance_pdf(queryset, "maintenance_garage")


def export_achat_xls(request):
    historique = request.GET.get("scope") == "historique"
    queryset = _maintenance_queryset()
    queryset = queryset.exclude(statut="en_cours") if historique else queryset.filter(statut="en_cours")
    queryset, _ = _apply_maintenance_filters(request, queryset)
    return _export_maintenance_xls(queryset, "maintenance_achat")


def export_achat_pdf(request):
    historique = request.GET.get("scope") == "historique"
    queryset = _maintenance_queryset()
    queryset = queryset.exclude(statut="en_cours") if historique else queryset.filter(statut="en_cours")
    queryset, _ = _apply_maintenance_filters(request, queryset)
    return _export_maintenance_pdf(queryset, "maintenance_achat")


liste_maintenances = role_required("logistique", "maintenancier", "directeur")(garage_maintenances)
garage_maintenances = role_required("logistique", "maintenancier", "dga", "directeur", "invite", "controleur")(garage_maintenances)
achat_maintenances = role_required("logistique", "directeur", "controleur")(achat_maintenances)
paiements_maintenances = role_required("comptable", "comptable_sogefi", "comptable_avena", "caissiere", "caissiere_soni", "caissiere_avena", "directeur")(paiements_maintenances)
modifier_maintenance_paiement = role_required("comptable", "comptable_sogefi", "comptable_avena", "caissiere", "caissiere_soni", "caissiere_avena", "directeur")(modifier_maintenance_paiement)
appro_caisse = role_required("caissiere", "caissiere_soni", "caissiere_avena", "comptable", "comptable_sogefi", "comptable_avena", "directeur")(appro_caisse)
situation_caisse = role_required("caissiere", "caissiere_soni", "caissiere_avena", "comptable", "comptable_sogefi", "comptable_avena", "directeur")(situation_caisse)
export_situation_caisse_xls = role_required("caissiere", "caissiere_soni", "caissiere_avena", "comptable", "comptable_sogefi", "comptable_avena", "directeur")(export_situation_caisse_xls)
situation_dg = role_required("caissiere", "caissiere_soni", "caissiere_avena", "comptable", "comptable_sogefi", "comptable_avena", "directeur")(situation_dg)
stock_maintenances = role_required("logistique", "maintenancier", "dga", "directeur", "comptable", "invite", "controleur")(stock_maintenances)
ajouter_article_stock = role_required("logistique", "directeur")(ajouter_article_stock)
modifier_article_stock = role_required("logistique", "directeur")(modifier_article_stock)
ajouter_mouvement_stock = role_required("logistique", "directeur")(ajouter_mouvement_stock)
fournisseurs_maintenance = role_required("logistique", "directeur", "responsable_achat", "dga_sogefi", "dga_avena", "comptable_avena")(fournisseurs_maintenance)
ajouter_fournisseur = role_required("logistique", "directeur", "responsable_achat", "dga_sogefi", "dga_avena", "comptable_avena")(ajouter_fournisseur)
modifier_fournisseur = role_required("logistique", "directeur", "responsable_achat", "dga_sogefi", "dga_avena", "comptable_avena")(modifier_fournisseur)
ajouter_maintenance_garage = role_required("logistique", "maintenancier", "directeur")(ajouter_maintenance_garage)
modifier_maintenance_garage = role_required("logistique", "maintenancier")(modifier_maintenance_garage)
modifier_maintenance_achat = role_required("logistique", "directeur", "controleur")(modifier_maintenance_achat)
terminer_maintenance = role_required("logistique", "directeur")(terminer_maintenance)
valider_maintenance_logistique = role_required("logistique")(valider_maintenance_logistique)
rejeter_maintenance_dga = role_required("dga")(rejeter_maintenance_dga)
valider_maintenance_dga = role_required("dga")(valider_maintenance_dga)
rejeter_maintenance_dg = role_required("directeur")(rejeter_maintenance_dg)
valider_maintenance_dg = role_required("directeur")(valider_maintenance_dg)
apercu_validation_maintenance = role_required("dga", "directeur")(apercu_validation_maintenance)
imprimer_maintenance = role_required("logistique", "maintenancier", "dga", "directeur", "comptable", "comptable_avena", "caissiere", "caissiere_soni", "caissiere_avena", "invite", "controleur")(imprimer_maintenance)
supprimer_maintenance = role_required("logistique", "maintenancier", "dga", "directeur")(supprimer_maintenance)
ajouter_type_maintenance_modal = role_required("logistique", "maintenancier", "directeur")(ajouter_type_maintenance_modal)
types_maintenance = role_required("logistique", "maintenancier", "directeur")(types_maintenance)
ajouter_type_maintenance = role_required("logistique", "maintenancier", "directeur")(ajouter_type_maintenance)
modifier_type_maintenance = role_required("logistique", "maintenancier", "directeur")(modifier_type_maintenance)
supprimer_type_maintenance = role_required("logistique", "maintenancier", "directeur")(supprimer_type_maintenance)
gerer_pannes = role_required("logistique", "maintenancier", "directeur")(gerer_pannes)
ajouter_panne_catalogue = role_required("logistique", "maintenancier", "directeur")(ajouter_panne_catalogue)
modifier_panne_catalogue = role_required("logistique", "maintenancier", "directeur")(modifier_panne_catalogue)
supprimer_panne_catalogue = role_required("logistique", "maintenancier", "directeur")(supprimer_panne_catalogue)
ajouter_prix_panne = role_required("logistique", "maintenancier", "directeur")(ajouter_prix_panne)
supprimer_prix_panne = role_required("logistique", "maintenancier", "directeur")(supprimer_prix_panne)
ajouter_panne_modal = role_required("logistique", "maintenancier", "directeur")(ajouter_panne_modal)
ajouter_fournisseur_modal = role_required("logistique", "directeur", "responsable_achat", "dga_sogefi", "dga_avena", "comptable_avena")(ajouter_fournisseur_modal)
ajouter_prestataire_modal = role_required("logistique", "directeur")(ajouter_prestataire_modal)
supprimer_fournisseur = role_required("logistique", "directeur", "responsable_achat", "dga_sogefi", "dga_avena", "comptable_avena")(supprimer_fournisseur)
export_garage_xls = role_required("logistique", "maintenancier", "dga", "directeur")(export_garage_xls)
export_garage_pdf = role_required("logistique", "maintenancier", "dga", "directeur")(export_garage_pdf)
export_achat_xls = role_required("logistique", "directeur")(export_achat_xls)
export_achat_pdf = role_required("logistique", "directeur")(export_achat_pdf)
rapport_maintenances = role_required("comptable", "comptable_avena", "caissiere", "caissiere_soni", "caissiere_avena", "logistique", "maintenancier", "dga", "directeur", "invite", "controleur")(rapport_maintenances)
export_rapport_caisse_xls = role_required("caissiere", "caissiere_soni", "caissiere_avena", "directeur")(export_rapport_caisse_xls)
export_rapport_maintenances_xls = role_required("comptable", "comptable_avena", "caissiere", "caissiere_soni", "caissiere_avena", "logistique", "maintenancier", "dga", "directeur", "invite", "controleur")(export_rapport_maintenances_xls)
